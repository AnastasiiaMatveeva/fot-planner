# -*- coding: utf-8 -*-
"""Кризис-тесты правил решателя.

Каждый случай — маленький входной файл под одно правило, собранный из
шаблона демо (справочник должностей и правила замещения берутся оттуда, все
переменные листы очищаются). Решатель вызывается напрямую, без сервиса.
Проверка двух видов:

* сценарные — что именно должно получиться в этом случае;
* инварианты — правила, которые обязаны выполняться в любом удачном плане
  (полная зарплата, шаг ставки, 122 при окладе, П2556, П4, БЭП, 152, …).

Запуск:  python harness/crisis.py [--only 07,09] [--limit 60]
Выход:   таблица «случай — вердикт — что не так» и код возврата 1 при провалах.
"""
from __future__ import annotations

import argparse
import io
import os
import shutil
import sys
import time
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook  # noqa: E402

from fot_planner import optimizer as opt  # noqa: E402
from fot_planner.excel.load import load_context  # noqa: E402
from fot_planner.models import PaymentKind as K  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "harness" / "crisis_template.xlsx"
WORK = ROOT / "harness" / "_crisis"
YEAR = 2026
EPS = 1.0  # рубль
MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
               "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

# Справочник (лист лимиты_по_должностям демо): оклад, П2556, П4 на ставку
REF = {
    "Инженер": (40400, 110000, 149648.9),
    "Ведущий инженер": (47000, 140000, 149648.9),
    "Программист": (40400, 110000, 149648.9),
    "Аналитик": (40400, 110000, 149648.9),
    "Лаборант": (38300, 95000, 149648.9),
    "Старший научный сотрудник": (74100, 120000, 257643.19),
}
BEP = 112261.0

VARIABLE_SHEETS = (
    "сотрудники", "договоры", "фот_по_месяцам", "120_надбавка",
    "трудоемкость_по_договорам", "фиксация_фот_по_месяцам",
    "минимальные_остатки", "ручные_назначения", "ручные_запреты",
)


# ── сборка входа ─────────────────────────────────────────────────────────
def emp(code, position, wage, *, rate=1.0, etype="основное", cat="основной",
        dep="Отдел 12", start="01.01.2026", end="31.12.2026", allowed=None, forbidden=None):
    return [code, code, position, dep, rate, etype, cat, wage, start, end, allowed, forbidden]


def ctr(code, fot, *, start="01.01.2026", end="31.12.2026", account="20576У17400",
        goz="нет", salary="да", k120="нет", k122="да", k124="нет", k152="нет",
        order="нет", priority="нет", main="да", part="да", salary_deadline=None,
        allowance_deadline=None, dep=None, kind="НИР"):
    return {
        "код": code, "название": code, "номер": "1/26", "тип договора": kind, "счет": account,
        "ГОЗ": goz, "дата начала": start, "дата окончания": end, "фот": fot,
        "оклад разрешен": salary, "120 разрешена": k120, "122 разрешена": k122,
        "124 разрешена": k124, "152 разрешена": k152, "стимулирующая приказом разрешена": order,
        "Приоритет": priority, "основное место разрешено": main, "совместительство разрешено": part,
        "конечная дата выплат оклада": salary_deadline or end,
        "конечная дата выплат надбавок": allowance_deadline or end,
        "подразделение": dep,
    }


def build(name, employees, contracts, inflow, *, labor=(), secret=(), manual=(),
          forbid=(), deficit="нет", max_contracts=3, tolerance=0.05, bep=None):
    """inflow: {код: сумма в месяц | список из 12}."""
    WORK.mkdir(exist_ok=True)
    path = WORK / f"{name}.xlsx"
    wb = load_workbook(TEMPLATE)
    for sheet in VARIABLE_SHEETS:
        ws = wb[sheet]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    ws = wb["договоры"]
    hdr = [c.value for c in ws[1]]
    if "подразделение" not in hdr:
        ws.cell(1, len(hdr) + 1).value = "подразделение"
        hdr.append("подразделение")
    for c in contracts:
        ws.append([c.get(h) for h in hdr])
    for e in employees:
        wb["сотрудники"].append(e)
    for code, v in inflow.items():
        row = v if isinstance(v, (list, tuple)) else [v] * 12
        wb["фот_по_месяцам"].append([code] + list(row))
    # Шестой элемент строки трудоёмкости — план по месяцам {месяц: чел.-мес.}.
    ws_l = wb["трудоемкость_по_договорам"]
    if any(len(row) > 5 for row in labor):
        for name in MONTH_NAMES:
            ws_l.cell(1, ws_l.max_column + 1).value = name
    for row in labor:
        code, pm, position, avg, heads = row[:5]
        line = [code, YEAR, pm, position, None, None, None, avg, heads]
        if len(row) > 5:
            line += [row[5].get(m) for m in range(1, 13)]
        ws_l.append(line)
    for e_id, c_id, rate in secret:
        wb["120_надбавка"].append([e_id, c_id, rate])
    for e_id, c_id, m1, m2, kind, fixed in manual:
        wb["ручные_назначения"].append([e_id, c_id, YEAR, m1, m2, kind, fixed])
    for e_id, c_id, m1, m2, kind in forbid:
        wb["ручные_запреты"].append([e_id, c_id, YEAR, m1, m2, kind])
    if bep:
        wl = wb["лимиты_по_должностям"]
        for r in wl.iter_rows(min_row=2):
            if r[0].value in bep:
                r[8].value = bep[r[0].value]
    ws = wb["настройки"]
    ws.cell(2, 2).value = deficit
    ws.cell(2, 3).value = max_contracts
    ws.cell(2, 8).value = tolerance
    wb.save(path)
    return path


def solve(path, limit):
    ctx = load_context(path)
    return ctx, opt.solve(ctx, time_limit_sec=limit)


# ── разбор результата ────────────────────────────────────────────────────
class Plan:
    def __init__(self, ctx, res):
        self.ctx, self.res = ctx, res
        self.status = res.solver_status
        self.pay = defaultdict(float)          # (e, c, m, kind) → сумма
        self.by_em = defaultdict(float)        # (e, m) → всего
        self.kinds = defaultdict(set)          # (e, c, m) → виды
        for a in res.allocations:
            if a.amount <= 0.5:
                continue
            self.pay[(a.employee_id, a.contract_id, a.month, a.payment_kind)] += a.amount
            self.by_em[(a.employee_id, a.month)] += a.amount
            self.kinds[(a.employee_id, a.contract_id, a.month)].add(a.payment_kind)
        self.rate = {}                          # (e, c, m) → (ставка, основное)
        self.rate_pos = {}                      # (e, c, m) → должность ставки
        for r in res.open_rate_attributions:
            if r.open_rate > 0:
                self.rate[(r.employee_id, r.contract_id, r.month)] = (r.open_rate, r.is_main)
                self.rate_pos[(r.employee_id, r.contract_id, r.month)] = r.position
        self.pm = defaultdict(float)            # (e, c, m) → чел.-мес.
        self.pm_row = defaultdict(float)        # (c, должность) → чел.-мес. за год
        self.people_row = defaultdict(set)      # (c, должность, m) → люди
        for r in res.labor_pm_attributions:
            if r.person_months > 1e-6:
                self.pm[(r.employee_id, r.contract_id, r.month)] += r.person_months
                self.pm_row[(r.contract_id, r.position)] += r.person_months
                self.people_row[(r.contract_id, r.position, r.month)].add(r.employee_id)
        self.lpay = defaultdict(float)          # (c, должность) → выплаты за год
        for r in res.labor_payment_attributions:
            self.lpay[(r.contract_id, r.position)] += r.amount
        self.deficit = defaultdict(float)
        for d in res.deficits:
            self.deficit[(d.employee_id, d.month)] += d.amount
        self.emp = {e.id: e for e in ctx.employees}
        self.ctr = {c.id: c for c in ctx.contracts}

    def amount(self, e, c, m, kind):
        return self.pay.get((e, c, m, kind), 0.0)

    def total(self, e, m, kinds, c=None):
        return sum(v for (ee, cc, mm, k), v in self.pay.items()
                   if ee == e and mm == m and k in kinds and (c is None or cc == c))

    def rate_total(self, e, m):
        return sum(v for (ee, cc, mm), (v, _) in self.rate.items() if ee == e and mm == m)

    def months_with(self, e, kind, c=None):
        return sorted({mm for (ee, cc, mm, k) in self.pay if ee == e and k == kind and (c is None or cc == c)})


# ── инварианты: обязаны выполняться в любом удачном плане ────────────────
def invariants(p: Plan):
    out = []
    # Шаг ставки и прочее, что умеет проверка результата, стенд не повторяет:
    # он требует, чтобы её поймал сам сервис. Правило, живущее только в
    # стенде, в работе не работает.
    if p.res.audit_status == "FAIL":
        out += ["проверка результата, %s: %s" % (v.rule_id, v.message)
                for v in p.res.audit_violations]
    ctx = p.ctx
    months = range(1, 13)
    active = lambda e, m: opt.employee_active_in_month(e, YEAR, m)  # noqa: E731
    for e in ctx.employees:
        for m in months:
            if not active(e, m):
                continue
            paid = p.by_em.get((e.id, m), 0.0) + p.deficit.get((e.id, m), 0.0)
            if abs(paid - e.monthly_wage) > EPS:
                out.append(f"зарплата {e.id} м{m}: выплачено {paid:.0f} + дефицит ≠ {e.monthly_wage:.0f}")
            # ставки
            rates = [(c, v, main) for (ee, c, mm), (v, main) in p.rate.items() if ee == e.id and mm == m]
            mains = [c for c, v, main in rates if main]
            if rates:
                # Основное — одно на человека, а не на строку: у второй
                # строки (совместительство) основного быть не должно.
                person = " ".join(str(e.full_name or "").lower().split()) or e.id
                pm_main = [c for (ee, c, mm), (v, main) in p.rate.items()
                           if mm == m and main and
                           (" ".join(str(p.emp[ee].full_name or "").lower().split()) or ee) == person]
                if len(pm_main) != 1:
                    out.append(f"основное место {e.id} м{m}: у человека {len(pm_main)} вместо 1")
                if e.employment_type == "part_time" and mains:
                    out.append(f"строка совместительства {e.id} м{m} помечена основным")
                # Основное место не трогаем: ставка ровно штатная, должность своя.
                # Сверх неё только совместительство.
                for c, v, main in rates:
                    if main and abs(v - e.rate) > 1e-6:
                        out.append(f"основная ставка изменена {e.id} {c} м{m}: {v} вместо штатной {e.rate}")
                    pos = p.rate_pos.get((e.id, c, m))
                    if main and pos and pos.strip().lower() != e.position.strip().lower():
                        out.append(f"должность основного места изменена {e.id} {c} м{m}: {pos} вместо {e.position}")
                tot = sum(v for _, v, _ in rates)
                if tot + 1e-6 < e.rate:
                    out.append(f"штатная ставка {e.id} м{m}: открыто {tot} < {e.rate}")
                cap = {"student": 0.5, "graduate_student": 0.75}.get(e.employment_category, 1.5)
                if e.employment_type == "part_time":
                    cap = 0.5
                if tot > cap + 1e-6:
                    out.append(f"суммарная ставка {e.id} м{m}: {tot} > {cap}")
                part = sum(v for _, v, main in rates if not main)
                if part > 0.5 + 1e-6:
                    out.append(f"совместительство {e.id} м{m}: {part} > 0,5")
            # оклад = справочный × ставка на договоре
            for (ee, c, mm, k), v in list(p.pay.items()):
                if ee != e.id or mm != m or k is not K.SALARY:
                    continue
                r = p.rate.get((e.id, c, m))
                if r is None:
                    out.append(f"оклад без ставки {e.id} {c} м{m}")
                    continue
                ref = REF.get(e.position)
                if ref and abs(v - ref[0] * r[0]) > EPS:
                    out.append(f"оклад {e.id} {c} м{m}: {v:.0f} ≠ {ref[0]}×{r[0]}")
            # 122 только с окладом на том же договоре
            for (ee, c, mm), kinds in p.kinds.items():
                if ee != e.id or mm != m:
                    continue
                if K.K122 in kinds and K.SALARY not in kinds:
                    out.append(f"122 без оклада {e.id} {c} м{m}")
                if K.K152 in kinds and kinds & {K.K120, K.K122, K.K124, K.ORDER_INCENTIVE}:
                    out.append(f"152 вместе с другими надбавками {e.id} {c} м{m}")
                ctr = p.ctr[c]
                for k, flag in ((K.K120, "allow_secret"), (K.K122, "allow_allowance"),
                                (K.K124, "allow_incentive"), (K.K152, "allow_extra_work"),
                                (K.ORDER_INCENTIVE, "allow_order_incentive"), (K.SALARY, "allow_salary")):
                    if k in kinds and not getattr(ctr, flag) and not (k is K.K120):
                        out.append(f"{k.name} с договора, где он запрещён: {e.id} {c} м{m}")
                if ctr.priority_payment_mode and K.K152 in kinds and kinds & {K.SALARY, K.K120, K.K122, K.K124}:
                    out.append(f"приоритет: штатные и 152 вместе {e.id} {c} м{m}")
                if ctr.priority_payment_mode and kinds & {K.SALARY, K.K120, K.K122, K.K124}:
                    other = [cc for (e2, cc, m2), kk in p.kinds.items()
                             if e2 == e.id and m2 == m and cc != c and K.SALARY in kk]
                    if other:
                        out.append(f"приоритет: штатные на {c} при окладе на {other} {e.id} м{m}")
            # П2556 по договору и открытой ставке
            ref = REF.get(e.position)
            if ref:
                for (ee, c, mm), (rv, _) in p.rate.items():
                    if ee != e.id or mm != m:
                        continue
                    staff = p.total(e.id, m, {K.SALARY, K.K122}, c)
                    if staff > ref[1] * rv + EPS:
                        out.append(f"П2556 {e.id} {c} м{m}: оклад+122 {staff:.0f} > {ref[1]}×{rv}")
                order = p.total(e.id, m, {K.ORDER_INCENTIVE})
                if order > ref[1] * e.rate + EPS:
                    out.append(f"П2556 приказ {e.id} м{m}: {order:.0f} > {ref[1]}×{e.rate}")
                n_orders = sum(1 for (ee, c, mm, k) in p.pay if ee == e.id and mm == m and k is K.ORDER_INCENTIVE)
                if n_orders > 1:
                    out.append(f"приказ дважды в месяц {e.id} м{m}: {n_orders}")
                # П4: в месяце со 124 штатная часть ровно предел × суммарная ставка
                if p.total(e.id, m, {K.K124}) > 0:
                    staff = p.total(e.id, m, {K.SALARY, K.K122, K.K124})
                    lim = ref[2] * p.rate_total(e.id, m)
                    if staff > lim + EPS:
                        out.append(f"П4 {e.id} м{m}: {staff:.0f} > {lim:.0f}")
                    elif staff < lim - EPS:
                        out.append(f"П4 не ровно {e.id} м{m}: {staff:.0f} < {lim:.0f}")
            # чел.-мес. ≤ открытая ставка на договоре
            for (ee, c, mm), pm in p.pm.items():
                if ee != e.id or mm != m:
                    continue
                rv = p.rate.get((e.id, c, m), (0.0, False))[0]
                if pm > rv + 1e-6:
                    out.append(f"чел.-мес. {e.id} {c} м{m}: {pm} > ставка {rv}")
    # 120: только по листу, в окне договора секретности, ровно процент от оклада
    listed = {s.employee_id: s for s in ctx.secret_allowances}
    for e in ctx.employees:
        for m in months:
            k120 = p.total(e.id, m, {K.K120})
            s = listed.get(e.id)
            if s is None:
                if k120 > 0:
                    out.append(f"120 не по листу {e.id} м{m}")
                continue
            sc = p.ctr.get(s.secret_contract_id)
            from fot_planner.contract_calendar import contract_payment_window_includes_month
            inwin = sc is not None and contract_payment_window_includes_month(sc, YEAR, m, K.K120)
            sal = p.total(e.id, m, {K.SALARY})
            if inwin and abs(k120 - s.rate * sal) > EPS:
                out.append(f"120 {e.id} м{m}: {k120:.0f} ≠ {s.rate}×{sal:.0f}")
            if not inwin and k120 > 0:
                out.append(f"120 вне окна секретности {e.id} м{m}")
            for (ee, c, mm, k), v in p.pay.items():
                if ee == e.id and mm == m and k is K.K120:
                    ctr = p.ctr[c]
                    acc = "".join(ch for ch in ctr.account if ch.isdigit())
                    if not acc.startswith("23") and K.SALARY not in p.kinds[(e.id, c, m)]:
                        out.append(f"120 без оклада на договоре не 23: {e.id} {c} м{m}")
    # БЭП: оклад+122 по ГОЗ-договору ≤ Σ БЭП должности × ставка
    bep_by_pos = {}
    for row in (ctx.position_limit_tables.get("bep") or []):
        if row.limit:
            bep_by_pos[row.position.strip().lower()] = float(row.limit)
    for c in ctx.contracts:
        if not c.is_goz_defense_order:
            continue
        for m in months:
            staff = sum(v for (ee, cc, mm, k), v in p.pay.items()
                        if cc == c.id and mm == m and k in (K.SALARY, K.K122))
            cap = sum(bep_by_pos.get(p.emp[ee].position.strip().lower(), BEP) * v
                      for (ee, cc, mm), (v, _) in p.rate.items() if cc == c.id and mm == m)
            if staff > cap + EPS:
                out.append(f"БЭП {c.id} м{m}: {staff:.0f} > {cap:.0f}")
    # касса: остаток не отрицательный, расход не раньше поступления
    for b in p.res.contract_balances:
        if getattr(b, "closing_balance", 0) < -EPS:
            out.append(f"касса {b.contract_id} м{b.month}: остаток {b.closing_balance:.0f}")
    # сроки договора
    for (e, c, m, k), v in p.pay.items():
        ctr = p.ctr[c]
        from fot_planner.contract_calendar import contract_allows_payment_month
        if not contract_allows_payment_month(ctr, YEAR, m, k) and k is not K.K120:
            out.append(f"выплата вне срока {e} {c} м{m} {k.name}")
    return out


# ── случаи ───────────────────────────────────────────────────────────────
CASES = []


def case(num, title):
    def wrap(fn):
        CASES.append((num, title, fn))
        return fn
    return wrap


def expect(cond, msg, out):
    if not cond:
        out.append(msg)


@case("01", "Зарплата целиком: оклад по справочнику, остаток надбавкой 122, приказ не нужен")
def c01(limit):
    path = build("c01", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1200000, order="да")], {"C_A": 100000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(abs(p.amount("E1", "C_A", m, K.SALARY) - 40400) < EPS, f"оклад м{m} ≠ 40400", out)
        expect(abs(p.amount("E1", "C_A", m, K.K122) - 59600) < EPS, f"122 м{m} ≠ 59600", out)
        expect(p.amount("E1", "C_A", m, K.ORDER_INCENTIVE) < EPS, f"приказ м{m} без нужды", out)
    return p, out


@case("02", "П2556: оклад+122 не выше предела на ставку, остаток — приказом, приказ раз в месяц")
def c02(limit):
    path = build("c02", [emp("E1", "Инженер", 130000)],
                 [ctr("C_A", 1600000, order="да"), ctr("C_B", 600000, salary="нет", k122="нет", order="да")],
                 {"C_A": 130000, "C_B": 50000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        staff = p.total("E1", m, {K.SALARY, K.K122})
        expect(abs(staff - 110000) < EPS, f"м{m}: оклад+122 {staff:.0f} ≠ предел 110000", out)
        expect(abs(p.total("E1", m, {K.ORDER_INCENTIVE}) - 20000) < EPS, f"м{m}: приказ ≠ 20000", out)
    return p, out


@case("03", "П2556 без приказа и без 124/152 — решения нет; с дефицитом — дефицит 20 000")
def c03(limit):
    path = build("c03a", [emp("E1", "Инженер", 130000)], [ctr("C_A", 1600000)], {"C_A": 130000})
    ctx, res = solve(path, limit)
    out = []
    expect(res.solver_status != "OPTIMAL", f"без приказа статус {res.solver_status}, ждали INFEASIBLE", out)
    path = build("c03b", [emp("E1", "Инженер", 130000)], [ctr("C_A", 1600000)], {"C_A": 130000}, deficit="да")
    ctx, res = solve(path, limit)
    p = Plan(ctx, res)
    expect(p.status == "OPTIMAL", f"с дефицитом статус {p.status}", out)
    for m in range(1, 13):
        expect(abs(p.deficit.get(("E1", m), 0) - 20000) < EPS, f"м{m}: дефицит {p.deficit.get(('E1', m), 0):.0f} ≠ 20000", out)
    return p, out


@case("04", "122 только с окладом на том же договоре")
def c04(limit):
    path = build("c04", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 500000), ctr("C_B", 900000, salary="нет", order="да")],
                 {"C_A": 41000, "C_B": 70000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(p.amount("E1", "C_B", m, K.K122) < EPS, f"м{m}: 122 на C_B без оклада", out)
        expect(p.amount("E1", "C_B", m, K.ORDER_INCENTIVE) > 0, f"м{m}: остаток не добран приказом с C_B", out)
    return p, out


@case("05", "120: только по листу, в окне договора секретности, ровно 5 % оклада; чужим — нет")
def c05(limit):
    path = build("c05", [emp("E1", "Инженер", 100000), emp("E2", "Инженер", 100000)],
                 [ctr("C_A", 2500000, k120="да"),
                  ctr("C_SEC", 60000, start="01.03.2026", end="31.08.2026", account="23000У00001",
                      salary="нет", k122="нет", k120="да")],
                 {"C_A": 200000, "C_SEC": [0, 0, 5000, 5000, 5000, 5000, 5000, 5000, 0, 0, 0, 0]},
                 secret=[("E1", "C_SEC", 0.05)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    expect(p.months_with("E1", K.K120) == [3, 4, 5, 6, 7, 8], f"120 у E1 в месяцах {p.months_with('E1', K.K120)}", out)
    expect(p.months_with("E2", K.K120) == [], "120 у E2, которого нет в листе", out)
    for m in range(3, 9):
        expect(abs(p.total("E1", m, {K.K120}) - 2020) < EPS, f"м{m}: 120 {p.total('E1', m, {K.K120}):.0f} ≠ 2020", out)
    return p, out


@case("06", "«120 разрешена = нет» на договоре оклада: оклад уходит туда, где 120 разрешена; одному такому — нет решения")
def c06(limit):
    path = build("c06a", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1300000, k120="нет"), ctr("C_B", 1300000, k120="да")],
                 {"C_A": 110000, "C_B": 110000},
                 secret=[("E1", "C_A", 0.05)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(p.amount("E1", "C_A", m, K.K120) < EPS, f"м{m}: 120 с договора, где она запрещена", out)
        expect(abs(p.amount("E1", "C_B", m, K.K120) - 2020) < EPS, f"м{m}: 120 с C_B {p.amount('E1', 'C_B', m, K.K120):.0f} ≠ 2020", out)
        expect(p.amount("E1", "C_B", m, K.SALARY) > 0, f"м{m}: оклад не на C_B", out)
    path = build("c06b", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1300000, k120="нет")], {"C_A": 110000},
                 secret=[("E1", "C_A", 0.05)])
    ctx, res = solve(path, limit)
    expect(res.solver_status != "OPTIMAL", f"единственный договор с запретом 120: статус {res.solver_status}, ждали INFEASIBLE", out)
    return p, out


@case("07", "152 вытесняет остальные надбавки в месяце (П2556 заставляет уйти в 152)")
def c07(limit):
    path = build("c07", [emp("E1", "Инженер", 130000)],
                 [ctr("C_A", 1600000, k152="да")], {"C_A": 130000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(p.total("E1", m, {K.K152}) > 0, f"м{m}: без 152 зарплата не закрывается", out)
        expect(p.total("E1", m, {K.K122, K.K124, K.K120, K.ORDER_INCENTIVE}) < EPS, f"м{m}: 152 вместе с другими", out)
        expect(abs(p.total("E1", m, {K.K152}) - 89600) < EPS, f"м{m}: 152 ≠ 89600", out)
    return p, out


@case("08", "П4: в месяце со 124 оклад+122+124 режется ровно по пределу, остаток приказом")
def c08(limit):
    path = build("c08", [emp("E1", "Ведущий инженер", 160000)],
                 [ctr("C_A", 2000000, k124="да", order="да")], {"C_A": 160000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        k124 = p.total("E1", m, {K.K124})
        staff = p.total("E1", m, {K.SALARY, K.K122, K.K124})
        expect(k124 > 0, f"м{m}: 124 не назначена, хотя без неё П2556 не даёт закрыть", out)
        expect(abs(staff - 149648.9) < EPS, f"м{m}: оклад+122+124 {staff:.0f} ≠ 149 649", out)
        expect(abs(p.total("E1", m, {K.ORDER_INCENTIVE}) - (160000 - 149648.9)) < EPS, f"м{m}: приказ ≠ остаток", out)
    return p, out


@case("09", "П4 с совместительством: предел на суммарную ставку 1,5")
def c09(limit):
    path = build("c09", [emp("E1", "Ведущий инженер", 230000)],
                 [ctr("C_A", 2000000, k124="да", order="да", dep="Отдел 12"),
                  ctr("C_B", 700000, k124="да", order="да", dep="Лаборатория 3")],
                 {"C_A": 160000, "C_B": 70000},
                 labor=[("C_B", 6, "Ведущий инженер", 60000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        rt = p.rate_total("E1", m)
        staff = p.total("E1", m, {K.SALARY, K.K122, K.K124})
        if p.total("E1", m, {K.K124}) > 0:
            expect(abs(staff - 149648.9 * rt) < EPS, f"м{m}: штатная {staff:.0f} ≠ 149 649×{rt}", out)
    parts = [m for m in range(1, 13) if p.rate_total("E1", m) > 1.0]
    expect(len(parts) > 0, "совместительство на C_B не открыто", out)
    return p, out


@case("10", "БЭП по ГОЗ: средняя оклад+122 на ставку не выше 112 261, остаток приказом")
def c10(limit):
    path = build("c10", [emp("E1", "Ведущий инженер", 150000), emp("E2", "Ведущий инженер", 150000)],
                 [ctr("C_G", 3600000, goz="да", order="да")], {"C_G": 300000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        staff = sum(p.total(e, m, {K.SALARY, K.K122}, "C_G") for e in ("E1", "E2"))
        expect(staff <= 2 * BEP + EPS, f"м{m}: оклад+122 {staff:.0f} > 2×БЭП", out)
        expect(staff >= 2 * BEP - 1000, f"м{m}: БЭП не выбран ({staff:.0f}), приказ съел лишнее", out)
    return p, out


@case("11", "БЭП: одна строка выше БЭП допустима, если средняя в пределе")
def c11(limit):
    path = build("c11", [emp("E1", "Ведущий инженер", 130000), emp("E2", "Инженер", 80000)],
                 [ctr("C_G", 2600000, goz="да", order="да")], {"C_G": 210000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    above = [m for m in range(1, 13) if p.total("E1", m, {K.SALARY, K.K122}, "C_G") > BEP + EPS]
    expect(len(above) == 12, f"E1 выше БЭП только в {len(above)} месяцах, ждали 12 (средняя позволяет)", out)
    return p, out


@case("11b", "БЭП по должности: предел договора — сумма БЭП × ставка, а не наименьший для всех")
def c11b(limit):
    path = build("c11b", [emp("E1", "Ведущий инженер", 160000), emp("E2", "Инженер", 120000)],
                 [ctr("C_G", 3400000, goz="да", order="да")], {"C_G": 280000},
                 bep={"Ведущий инженер": 130000, "Инженер": 100000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        staff = sum(p.total(e, m, {K.SALARY, K.K122}, "C_G") for e in ("E1", "E2"))
        expect(staff <= 230000 + EPS, f"м{m}: оклад+122 {staff:.0f} > 130 000 + 100 000", out)
        expect(staff >= 230000 - 1000, f"м{m}: {staff:.0f} — предел по должностям не выбран (старый общий был бы 200 000)", out)
    return p, out


@case("12", "ГОЗ с «124 разрешена = да» во входе: решатель 124 не запрещает")
def c12(limit):
    path = build("c12", [emp("E1", "Ведущий инженер", 160000)],
                 [ctr("C_G", 2000000, goz="да", k124="да", order="да")], {"C_G": 160000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    k124 = sum(p.total("E1", m, {K.K124}) for m in range(1, 13))
    out.append(f"справка: 124 на ГОЗ выплачено {k124:.0f} — запрет живёт только во входе (графа договора)")
    return p, out


@case("13", "Приоритет: оклад на другом договоре → на приоритете только 152")
def c13(limit):
    path = build("c13", [emp("E1", "Инженер", 100000)],
                 [ctr("C_BASE", 500000), ctr("C_PRIO", 900000, priority="да", k152="да")],
                 {"C_BASE": 41000, "C_PRIO": 70000},
                 manual=[("E1", "C_BASE", 1, 12, "оклад", None)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(p.total("E1", m, {K.SALARY, K.K122, K.K120, K.K124}, "C_PRIO") < EPS, f"м{m}: штатные на приоритете при окладе на C_BASE", out)
        expect(p.total("E1", m, {K.K152}, "C_PRIO") > 0, f"м{m}: 152 с приоритета нет", out)
    return p, out


@case("14", "Приоритет: на договоре в месяце либо штатные, либо 152")
def c14(limit):
    path = build("c14", [emp("E1", "Инженер", 130000)],
                 [ctr("C_BASE", 700000), ctr("C_PRIO", 1600000, priority="да", k152="да")],
                 {"C_BASE": 60000, "C_PRIO": 130000})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        kinds = p.kinds[("E1", "C_PRIO", m)]
        expect(not (K.K152 in kinds and kinds & {K.SALARY, K.K122}), f"м{m}: штатные и 152 вместе", out)
    return p, out


@case("15", "Строки человека: основное + совместительство; совместительство без основного — ошибка входа")
def c15(limit):
    from fot_planner.validation import validate_context
    path = build("c15a", [emp("E1", "Инженер", 50000, rate=0.5, etype="совместительство")],
                 [ctr("C_A", 700000, order="да")], {"C_A": 60000})
    codes = [c.code for c in validate_context(load_context(path))]
    out = []
    expect("PART_TIME_WITHOUT_MAIN" in codes, f"совместительство без основного не отклонено: {codes}", out)
    # две строки одного человека: основное 1,0 инженером и 0,25 аналитиком
    e1 = emp("E1-1", "Инженер", 100000); e1[1] = "Иванов Иван Иванович"
    e2 = emp("E1-2", "Аналитик", 20000, rate=0.25, etype="совместительство"); e2[1] = "Иванов Иван Иванович"
    path = build("c15b", [e1, e2], [ctr("C_A", 1600000, order="да")], {"C_A": 130000})
    ctx, res = solve(path, limit)
    p = Plan(ctx, res)
    expect(p.status == "OPTIMAL", f"две строки: статус {p.status}", out)
    for m in range(1, 13):
        expect(p.rate.get(("E1-1", "C_A", m), (0, False)) == (1.0, True), f"м{m}: основное E1-1 не 1,0", out)
        r2 = p.rate.get(("E1-2", "C_A", m))
        expect(r2 is not None and abs(r2[0] - 0.25) < 1e-6 and not r2[1], f"м{m}: строка E1-2 {r2}", out)
    return p, out


@case("16", "Совместительство сверх основного: не больше 0,5, всего 1,5; чужое подразделение")
def c16(limit):
    path = build("c16", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1300000, dep="Отдел 12"), ctr("C_B", 500000, dep="Лаборатория 3")],
                 {"C_A": 110000, "C_B": 40000},
                 labor=[("C_B", 9, "Инженер", 30000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    part = sum(v for (e, c, m), (v, main) in p.rate.items() if e == "E1" and c == "C_B")
    expect(5.7 <= part <= 6.0 + 1e-6, f"совместительство на C_B за год {part} (потолок 0,5×12 = 6, план 9 недостижим)", out)
    for m in range(1, 13):
        expect(p.rate.get(("E1", "C_A", m), (0, False)) == (1.0, True), f"м{m}: основное на C_A не 1,0", out)
    return p, out


@case("17", "Два совместительства по 0,25 на разных договорах")
def c17(limit):
    path = build("c17", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1300000, dep="Отдел 12"), ctr("C_B", 200000, dep="Лаборатория 3"),
                  ctr("C_C", 200000, dep="Лаборатория 4")],
                 {"C_A": 110000, "C_B": 15000, "C_C": 15000},
                 labor=[("C_B", 3, "Инженер", 45000, 1), ("C_C", 3, "Инженер", 45000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for c in ("C_B", "C_C"):
        pm = p.pm_row.get((c, "Инженер"), 0)
        expect(2.85 <= pm <= 3.15, f"{c}: закрыто {pm} чел.-мес. вместо 3", out)
    both = [m for m in range(1, 13) if p.rate.get(("E1", "C_B", m)) and p.rate.get(("E1", "C_C", m))]
    out.append(f"справка: месяцев с двумя совместительствами сразу — {len(both)}")
    return p, out


@case("18", "Студент ≤ 0,5, аспирант ≤ 0,75, лаборант без дополнительных ставок")
def c18(limit):
    path = build("c18", [emp("S1", "Инженер", 40000, rate=0.5, cat="студент", dep="Отдел 12"),
                         emp("A1", "Инженер", 40000, rate=0.5, cat="аспирант", dep="Отдел 12"),
                         emp("L1", "Лаборант", 38300, rate=1.0, dep="Отдел 12")],
                 [ctr("C_A", 1600000, dep="Отдел 12"), ctr("C_B", 900000, dep="Лаборатория 3")],
                 {"C_A": 130000, "C_B": 70000},
                 labor=[("C_B", 12, "Инженер", 25000, 3), ("C_B", 6, "Лаборант", 20000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        expect(p.rate_total("S1", m) <= 0.5 + 1e-6, f"м{m}: студент {p.rate_total('S1', m)}", out)
        expect(p.rate_total("A1", m) <= 0.75 + 1e-6, f"м{m}: аспирант {p.rate_total('A1', m)}", out)
        expect(p.rate_total("L1", m) <= 1.0 + 1e-6, f"м{m}: лаборант {p.rate_total('L1', m)}", out)
    asp = sum(p.rate_total("A1", m) for m in range(1, 13))
    expect(asp > 6.0 + 1e-6, "аспирант не взял совместительство, хотя 0,25 разрешено", out)
    return p, out


@case("19", "Замещение направленное: ведущий инженер закрывает программиста, аналитик — нет")
def c19(limit):
    path = build("c19", [emp("E1", "Ведущий инженер", 100000), emp("E2", "Аналитик", 90000)],
                 [ctr("C_A", 2500000), ctr("C_P", 700000, dep="Лаборатория 3")],
                 {"C_A": 200000, "C_P": 60000},
                 labor=[("C_P", 6, "Программист", 40000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    expect(not any(c == "C_P" for (e, c, m, k) in p.pay if e == "E2"), "аналитик получил выплаты с C_P", out)
    pm = p.pm_row.get(("C_P", "Программист"), 0)
    expect(pm >= 5.7, f"ведущий инженер закрыл {pm} чел.-мес. из 6", out)
    # чужую строку закрывает любой своей ставкой, но должность остаётся своей
    for m in range(1, 13):
        if p.rate.get(("E1", "C_P", m)):
            expect((p.rate_pos.get(("E1", "C_P", m)) or "").lower() in ("", "ведущий инженер") or not p.rate[("E1", "C_P", m)][1],
                   f"м{m}: основное место подписано чужой должностью", out)
    return p, out


@case("20", "Замещение в обратную сторону не работает: программист не закроет ведущего инженера")
def c20(limit):
    path = build("c20", [emp("E1", "Программист", 90000)],
                 [ctr("C_A", 1300000), ctr("C_P", 700000, dep="Лаборатория 3")],
                 {"C_A": 110000, "C_P": 60000},
                 labor=[("C_P", 6, "Ведущий инженер", 60000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    expect(not any(c == "C_P" for (e, c, m, k) in p.pay), "программист посажен на строку ведущего инженера", out)
    expect(p.pm_row.get(("C_P", "Ведущий инженер"), 0) < 1e-6, "строка закрыта несовместимым", out)
    return p, out


@case("21", "Трудоёмкость: ±5 % по чел.-мес. и сумме, людей не больше заданного")
def c21(limit):
    path = build("c21", [emp("E1", "Инженер", 100000), emp("E2", "Инженер", 100000)],
                 [ctr("C_A", 2500000, dep="Отдел 12"), ctr("C_L", 700000, dep="Лаборатория 3")],
                 {"C_A": 210000, "C_L": 60000},
                 labor=[("C_L", 6, "Инженер", 50000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    pm = p.pm_row.get(("C_L", "Инженер"), 0)
    expect(5.7 - 1e-6 <= pm <= 6.3 + 1e-6, f"чел.-мес. {pm} вне 6 ± 5 %", out)
    amt = p.lpay.get(("C_L", "Инженер"), 0)
    expect(285000 - EPS <= amt <= 315000 + EPS, f"сумма по строке {amt:.0f} вне 300 000 ± 5 %", out)
    for m in range(1, 13):
        n = len(p.people_row.get(("C_L", "Инженер", m), ()))
        expect(n <= 1, f"м{m}: на строке {n} человека при пределе 1", out)
    return p, out


@case("22", "Трудоёмкость: деньги строки только от людей со ставкой на договоре")
def c22(limit):
    path = build("c22", [emp("E1", "Инженер", 100000), emp("E2", "Ведущий инженер", 120000)],
                 [ctr("C_A", 3000000, dep="Отдел 12"), ctr("C_L", 700000, dep="Лаборатория 3", k124="да")],
                 {"C_A": 240000, "C_L": 60000},
                 labor=[("C_L", 6, "Инженер", 50000, 1)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for r in res.labor_payment_attributions:
        if r.amount > EPS:
            expect(p.rate.get((r.employee_id, r.contract_id, r.month)) is not None,
                   f"деньги строки от {r.employee_id} м{r.month} без ставки на {r.contract_id}", out)
    return p, out


@case("23", "Касса: деньги не раньше поступления — без дефицита решения нет, с дефицитом 11 месяцев")
def c23(limit):
    inflow = [0] * 11 + [1200000]
    path = build("c23a", [emp("E1", "Инженер", 100000)], [ctr("C_A", 1200000)], {"C_A": inflow})
    ctx, res = solve(path, limit)
    out = []
    expect(res.solver_status != "OPTIMAL", f"без дефицита статус {res.solver_status}", out)
    path = build("c23b", [emp("E1", "Инженер", 100000)], [ctr("C_A", 1200000)], {"C_A": inflow}, deficit="да")
    ctx, res = solve(path, limit)
    out.append(f"справка: пустая касса и дефицит разрешён — статус {res.solver_status}: "
               "оклад открытой ставки дефицитом не покрывается, дефицит бывает только сверх оклада")
    # частичный дефицит: денег хватает на оклад, но не на всю зарплату
    path = build("c23c", [emp("E1", "Инженер", 100000)], [ctr("C_A", 1200000)],
                 {"C_A": [60000] * 11 + [540000]}, deficit="да")
    ctx, res = solve(path, limit)
    p = Plan(ctx, res)
    expect(p.status == "OPTIMAL", f"частичный дефицит: статус {p.status}", out)
    months = sorted(m for (e, m), v in p.deficit.items() if v > EPS)
    expect(len(months) >= 6, f"дефицит в месяцах {months}", out)
    expect(abs(p.by_em.get(("E1", 12), 0) - 100000) < EPS, "декабрь не выплачен целиком", out)
    return p, out


@case("24", "Сроки: договор до 30.06, оклад до 31.05 — после этих дат с него не платят")
def c24(limit):
    path = build("c24", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 600000, end="30.06.2026", salary_deadline="31.05.2026", allowance_deadline="30.06.2026"),
                  ctr("C_B", 800000, start="01.06.2026")],
                 {"C_A": [100000] * 6 + [0] * 6, "C_B": [0] * 5 + [110000] * 7})
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    expect(p.months_with("E1", K.SALARY, "C_A") == [1, 2, 3, 4, 5], f"оклад с C_A в {p.months_with('E1', K.SALARY, 'C_A')}", out)
    expect(max(p.months_with("E1", K.K122, "C_A") or [0]) <= 6, "надбавка с C_A после июня", out)
    expect(min(p.months_with("E1", K.SALARY, "C_B") or [13]) >= 6, "оклад с C_B до начала договора", out)
    return p, out


@case("25", "Ручные назначения и запреты")
def c25(limit):
    path = build("c25", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1300000), ctr("C_B", 1300000)], {"C_A": 110000, "C_B": 110000},
                 manual=[("E1", "C_B", 3, 6, "оклад", None)],
                 forbid=[("E1", "C_A", 3, 6, None)])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in (3, 4, 5, 6):
        expect(p.amount("E1", "C_B", m, K.SALARY) > 0, f"м{m}: оклад не на C_B", out)
        expect(p.total("E1", m, set(K), "C_A") < EPS, f"м{m}: выплата с запрещённого C_A", out)
    return p, out


@case("26", "Макс договоров оклада в год = 1")
def c26(limit):
    path = build("c26", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 700000), ctr("C_B", 700000), ctr("C_C", 700000)],
                 {"C_A": [110000] * 6 + [0] * 6, "C_B": [0] * 6 + [110000] * 6, "C_C": 110000},
                 max_contracts=1)
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    used = {c for (e, c, m, k) in p.pay if k is K.SALARY}
    expect(len(used) <= 1, f"оклад с {sorted(used)} при пределе 1", out)
    return p, out


@case("27", "Одна должность в одном подразделении: совместительство инженером только в чужом отделе")
def c27(limit):
    out, plans = [], []
    for dep, want in (("Отдел 12", 0), ("Лаборатория 3", 1)):
        path = build("c27_" + ("own" if want == 0 else "other"), [emp("E1", "Инженер", 100000, dep="Отдел 12")],
                     [ctr("C_A", 1300000, dep="Отдел 12"), ctr("C_B", 500000, dep=dep)],
                     {"C_A": 110000, "C_B": 40000}, labor=[("C_B", 6, "Инженер", 30000, 1)],
                     manual=[("E1", "C_A", 1, 12, "оклад", None)])
        ctx, res = solve(path, limit)
        p = Plan(ctx, res)
        plans.append(p)
        expect(p.status == "OPTIMAL", f"{dep}: статус {p.status}", out)
        part = sum(v for (e, c, m), (v, main) in p.rate.items() if c == "C_B")
        if want == 0:
            expect(part < 1e-6, f"{dep}: совместительство инженером в своём отделе {part}", out)
        else:
            expect(part > 5.7, f"{dep}: в чужом отделе совместительство {part} чел.-мес. вместо ~6", out)
    return plans[-1], out


@case("28", "Приказ ограничен П2556 отдельно и один в месяц при двух договорах")
def c28(limit):
    path = build("c28", [emp("E1", "Инженер", 200000)],
                 [ctr("C_A", 1600000, order="да"), ctr("C_B", 1600000, salary="нет", k122="нет", order="да")],
                 {"C_A": 110000, "C_B": 100000}, deficit="да")
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        n = sum(1 for (e, c, mm, k) in p.pay if mm == m and k is K.ORDER_INCENTIVE)
        expect(n <= 1, f"м{m}: приказов {n}", out)
        expect(p.total("E1", m, {K.ORDER_INCENTIVE}) <= 110000 + EPS, f"м{m}: приказ выше П2556", out)
    return p, out


@case("31", "Трудоёмкость по месяцам: этап июнь–сентябрь закрывается по 1,0 в своих месяцах и никогда вне их")
def c31(limit):
    path = build("c31", [emp("E1", "Инженер", 100000)],
                 [ctr("C_A", 1200000, order="да"), ctr("C_B", 500000, order="да")],
                 {"C_A": 100000, "C_B": 100000},
                 labor=[("C_B", 4, "Инженер", 100000, 1, {6: 1.0, 7: 1.0, 8: 1.0, 9: 1.0})])
    ctx, res = solve(path, limit)
    p, out = Plan(ctx, res), []
    expect(p.status == "OPTIMAL", f"статус {p.status}", out)
    for m in range(1, 13):
        got = p.pm.get(("E1", "C_B", m), 0.0)
        if 6 <= m <= 9:
            expect(abs(got - 1.0) < 0.05, f"м{m}: закрыто {got:.2f} ≠ 1,0", out)
        else:
            expect(got < 1e-6, f"м{m}: {got:.2f} чел.-мес. вне этапа", out)
    return p, out


# ── запуск ───────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="")
    ap.add_argument("--limit", type=int, default=60)
    args = ap.parse_args()
    only = {x.strip() for x in args.only.split(",") if x.strip()}
    rows, failed = [], 0
    for num, title, fn in CASES:
        if only and num not in only:
            continue
        t0 = time.perf_counter()
        try:
            plan, notes = fn(args.limit)
            inv = invariants(plan) if plan is not None and plan.status == "OPTIMAL" else []
        except Exception as e:  # noqa: BLE001
            plan, notes, inv = None, [f"ошибка стенда: {type(e).__name__}: {e}"], []
        dt = time.perf_counter() - t0
        info = [n for n in notes if n.startswith("справка")]
        bad = [n for n in notes if not n.startswith("справка")] + inv
        verdict = "ОК" if not bad else "ПРОВАЛ"
        failed += bool(bad)
        rows.append((num, verdict, title, dt, bad, info))
        print(f"[{num}] {verdict:6} {dt:5.0f} с  {title}")
        for b in bad[:12]:
            print("      ✗", b)
        if len(bad) > 12:
            print(f"      … ещё {len(bad) - 12}")
        for i in info:
            print("      ·", i)
        sys.stdout.flush()
    print(f"\nитого: {len(rows)} случаев, провалов {failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
