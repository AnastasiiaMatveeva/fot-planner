"""Шаблон входного Excel."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from fot_planner.excel.constants import (
    SHEET_CONTRACT_LABOR,
    SHEET_CONTRACTS,
    SHEET_EMPLOYEES,
    SHEET_FOT_MATRIX,
    SHEET_POSITION_LIMITS,
    SHEET_SECRET_ALLOWANCES,
    SHEET_SETTINGS,
)
from fot_planner.excel.workbook_format import format_workbook
from fot_planner.defaults.settings import (
    DEFAULT_GOZ_AVERAGE_SALARY_LIMIT,
    DEFAULT_SETTINGS_ROW,
)
from fot_planner.position_reference import (
    default_position_reference,
    normalize_position,
)
from fot_planner.salary_limits_2556 import default_position_salary_limits


def _default_position_limits_table() -> pd.DataFrame:
    reference_by_position = {
        normalize_position(row.position): row
        for row in default_position_reference()
    }
    limits_by_position = {
        normalize_position(row.position): row
        for row in default_position_salary_limits()
    }
    rows: list[dict[str, object]] = []
    for key in sorted(reference_by_position | limits_by_position):
        reference = reference_by_position.get(key)
        limits = limits_by_position.get(key)
        position = (
            limits.position
            if limits is not None
            else reference.position
            if reference is not None
            else key
        )
        rows.append(
            {
                "должность": position,
                "категория персонала": limits.personnel_category if limits else "",
                "страница": reference.salary_page if reference else "",
                "номер группы": reference.salary_group_number if reference else "",
                "номер уровня": reference.level if reference else "",
                "оклад": (
                    reference.reference_salary_for_rate
                    if reference and reference.reference_salary_for_rate is not None
                    else ""
                ),
                "П2556": (
                    limits.order_2556_limit
                    if limits and limits.order_2556_limit is not None
                    else ""
                ),
                "П4": (
                    limits.p4_limit
                    if limits and limits.p4_limit is not None
                    else ""
                ),
                "БЭП": DEFAULT_GOZ_AVERAGE_SALARY_LIMIT,
                "примечание": limits.note if limits and limits.note else "",
            }
        )
    return pd.DataFrame(rows)


def create_template(path: str | Path) -> None:
    path = Path(path)
    year = date.today().year

    employees = pd.DataFrame(
        [
            {
                "код строки": "E001-1",
                "фио": "Иванов Иван Иванович",
                "должность": "инженер",
                "подразделение": "лаборатория",
                "ставка": 1.0,
                "тип занятости": "основное",
                "категория занятости": "основной",
                "зарплата": 100000,
                "дата начала": f"{year}-01-01",
                "дата окончания": "",
                "разрешенные договоры": "",
                "запрещенные договоры": "",
            }
        ]
    )
    contract_labor = pd.DataFrame(
        columns=[
            "договор",
            "год",
            "трудоемкость",
            "должность",
            "страница",
            "номер группы",
            "номер уровня",
            "средняя стоимость выполнения работ в месяц",
        ]
    )
    contracts = pd.DataFrame(
        [
            {
                "код": "C001",
                "название": "НИОКР Альфа",
                "номер": "123/2025",
                "тип договора": "госзаказ",
                "счет": "2301",
                "ГОЗ": True,
                "дата начала": f"{year}-01-01",
                "дата окончания": f"{year}-12-31",
                "фот": 1_200_000,
                "оклад разрешен": True,
                "120 разрешена": False,
                "122 разрешена": True,
                "124 разрешена": False,
                "152 разрешена": False,
                "стимулирующая приказом разрешена": False,
                "проект Приоритет": "нет",
                "основное место разрешено": "да",
                "совместительство разрешено": "да",
                "конечная дата выплат оклада": f"{year}-12-31",
                "конечная дата выплат надбавок": f"{year}-12-31",
            }
        ]
    )
    secret_allowances = pd.DataFrame(
        columns=[
            "сотрудник",
            "договор секретности",
            "ставка 120",
        ]
    )
    settings = pd.DataFrame([{"год": year, **DEFAULT_SETTINGS_ROW}])
    fot_by_month = pd.DataFrame({"договор": ["C001"]})
    for m in range(1, 13):
        fot_by_month[str(m)] = [1_200_000 if m == 1 else ""]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        employees.to_excel(writer, sheet_name=SHEET_EMPLOYEES, index=False)
        contracts.to_excel(writer, sheet_name=SHEET_CONTRACTS, index=False)
        secret_allowances.to_excel(writer, sheet_name=SHEET_SECRET_ALLOWANCES, index=False)
        contract_labor.to_excel(writer, sheet_name=SHEET_CONTRACT_LABOR, index=False)
        fot_by_month.to_excel(writer, sheet_name=SHEET_FOT_MATRIX, index=False)
        settings.to_excel(writer, sheet_name=SHEET_SETTINGS, index=False)
        _default_position_limits_table().to_excel(
            writer,
            sheet_name=SHEET_POSITION_LIMITS,
            index=False,
        )

    wb = load_workbook(path)
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for sheet_name in (
        SHEET_CONTRACTS,
        SHEET_EMPLOYEES,
        SHEET_SECRET_ALLOWANCES,
    ):
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = yellow
    wb.save(path)
    format_workbook(path)
