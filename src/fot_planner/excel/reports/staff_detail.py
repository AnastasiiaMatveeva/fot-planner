"""Детализация штатного расписания в формате кадрового шаблона."""

from __future__ import annotations

import calendar
from collections import defaultdict
from datetime import date
from pathlib import Path

from fot_planner.contract_calendar import payment_deadline_date
from fot_planner.models import PaymentKind, PlanningContext, PlanningResult
from fot_planner.position_reference import normalize_position, personnel_category_for_position

SHEET_STAFF_DETAIL = "ШР_детализация"

STAFF_DETAIL_HEADERS = [
    "Таб.№",
    "Назначение",
    "Фамилия И.О., уч. ст., уч. зван.",
    "Код \nподр.",
    "Подразделение",
    "Должность",
    "Категория\nперсонала",
    "Код \nпар-ра",
    "Название параметра",
    "Ставка",
    "Номинальное\n значение \nпараметра",
    "Значение\n параметра \nпо ставке",
    "Сумма \nв\nруб.",
    "Начало \nдействия",
    "Окончание \nдействия",
    "Код\nшифра\nзатрат",
    "Шифр \nзатрат",
    "Лицевой \nсчет",
    "УИ\nПНИЭР",
    "Номер \nприказа\n ввода",
    "Дата \nприказа\nввода",
    "Номер \nприказа \nзакрытия",
    "Дата \nприказа\nзакрытия",
]

_PAYMENT_PARAMETERS: dict[PaymentKind, tuple[str, str]] = {
    PaymentKind.SALARY: ("1", "Оклад (должностной оклад)"),
    PaymentKind.K120: ("120", "За работу со сведениями, составляющими государственную тайну"),
    PaymentKind.K122: ("122", "За качество выполняемых работ"),
    PaymentKind.K124: ("124", "За интенсивность и высокие результаты работы"),
    PaymentKind.K152: ("152", "За выполнение дополнительной работы"),
    PaymentKind.ORDER_INCENTIVE: ("приказ", "Стимулирующая выплата приказом"),
}


def _month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def _contiguous_ranges(months: set[int]) -> list[tuple[int, int]]:
    if not months:
        return []
    ordered = sorted(months)
    ranges: list[tuple[int, int]] = []
    start = previous = ordered[0]
    for month in ordered[1:]:
        if month == previous + 1:
            previous = month
            continue
        ranges.append((start, previous))
        start = previous = month
    ranges.append((start, previous))
    return ranges


def _period_dates(ctx, employee, contract, kind, first_month: int, last_month: int):
    start = date(ctx.year, first_month, 1)
    for boundary in (employee.start_date, contract.start_date):
        if boundary is not None and boundary.year == ctx.year and boundary > start:
            start = boundary

    end = _month_end(ctx.year, last_month)
    known_ends = [payment_deadline_date(contract, kind)]
    if employee.end_date is not None:
        known_ends.append(employee.end_date)
    for boundary in known_ends:
        if boundary.year == ctx.year and boundary < end:
            end = boundary

    continues_after_year = last_month == 12 and all(
        boundary.year > ctx.year for boundary in known_ends
    )
    return start, None if continues_after_year else end


def build_staff_detail_rows(ctx: PlanningContext, result: PlanningResult) -> list[dict]:
    """Свернуть месячные выплаты в непрерывные кадровые интервалы."""
    employees = {employee.id: employee for employee in ctx.employees}
    contracts = {contract.id: contract for contract in ctx.contracts}
    personnel_categories = {
        normalize_position(row.position): row.personnel_category
        for row in ctx.position_salary_limits
    }
    rates = {
        (record.employee_id, record.contract_id, record.month): (
            record.open_rate,
            record.is_main,
            record.position,
        )
        for record in result.open_rate_attributions
        if record.open_rate > 0
    }

    grouped: dict[tuple, set[int]] = defaultdict(set)
    for allocation in result.allocations:
        if allocation.amount <= 0.005:
            continue
        employee = employees.get(allocation.employee_id)
        contract = contracts.get(allocation.contract_id)
        if employee is None or contract is None:
            continue

        if allocation.payment_kind not in _PAYMENT_PARAMETERS:
            continue

        rate, is_main, assigned_position = rates.get(
            (allocation.employee_id, allocation.contract_id, allocation.month),
            (employee.rate, True, employee.position),
        )
        assigned_position = assigned_position or employee.position
        if rate <= 0:
            rate = employee.rate if employee.rate > 0 else 1.0

        if allocation.payment_kind is PaymentKind.SALARY and employee.reference_salary_for_rate:
            nominal = employee.reference_salary_for_rate
        else:
            nominal = allocation.amount / rate
        value_by_rate = nominal * rate

        key = (
            allocation.employee_id,
            allocation.contract_id,
            allocation.payment_kind,
            bool(is_main),
            assigned_position,
            round(rate, 4),
            round(nominal, 2),
            round(value_by_rate, 2),
        )
        grouped[key].add(allocation.month)

    rows = []
    for key, allocated_months in grouped.items():
        employee_id, contract_id, kind, is_main, assigned_position, rate, nominal, value_by_rate = key
        employee = employees[employee_id]
        contract = contracts[contract_id]
        parameter_code, parameter_name = _PAYMENT_PARAMETERS[kind]
        for first_month, last_month in _contiguous_ranges(allocated_months):
            starts, ends = _period_dates(
                ctx, employee, contract, kind, first_month, last_month
            )
            row = {header: "" for header in STAFF_DETAIL_HEADERS}
            # Должность в форме кадров — с видом занятости и ставкой:
            # «Инженер, 1 ставк.», «Инженер, внутр. совмест., 0,25 ставк.».
            rate_text = ("%g" % rate).replace(".", ",")
            position_text = "%s, %s%s ставк." % (
                assigned_position, "" if is_main else "внутр. совмест., ", rate_text)
            row.update(
                {
                    "Таб.№": employee.id,
                    "Фамилия И.О., уч. ст., уч. зван.": employee.full_name,
                    "Подразделение": employee.department or "",
                    "Должность": position_text,
                    "Код\nшифра\nзатрат": contract.id,
                    "Шифр \nзатрат": contract.name or "",
                    "Лицевой \nсчет": contract.account or "",
                    "Категория\nперсонала": personnel_categories.get(
                        normalize_position(assigned_position)
                    )
                    or personnel_category_for_position(assigned_position)
                    or "",
                    "Код \nпар-ра": parameter_code,
                    "Название параметра": parameter_name,
                    "Ставка": rate,
                    "Номинальное\n значение \nпараметра": nominal,
                    "Значение\n параметра \nпо ставке": value_by_rate,
                    "Сумма \nв\nруб.": value_by_rate,
                    "Начало \nдействия": starts,
                    "Окончание \nдействия": ends or "",
                    "_sort": (
                        employee.full_name,
                        starts,
                        0 if is_main else 1,
                        {
                            PaymentKind.SALARY: 0,
                            PaymentKind.K120: 1,
                            PaymentKind.K122: 2,
                            PaymentKind.K124: 3,
                            PaymentKind.K152: 4,
                            PaymentKind.ORDER_INCENTIVE: 5,
                        }[kind],
                        contract_id,
                    ),
                }
            )
            rows.append(row)

    rows.sort(key=lambda row: row["_sort"])
    for row in rows:
        row.pop("_sort", None)
    return rows


def write_staff_detail_sheet(
    path: str | Path, ctx: PlanningContext, result: PlanningResult
) -> None:
    """Добавить в результат оформленный лист с 23 колонками кадрового шаблона."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    if SHEET_STAFF_DETAIL in wb.sheetnames:
        del wb[SHEET_STAFF_DETAIL]
    ws = wb.create_sheet(SHEET_STAFF_DETAIL)

    ws["A1"] = "Подразделение -"
    ws["A2"] = f"Дата построения {date.today():%d.%m.%Y}"
    ws["A3"] = "Лицевой счет - по всем ЛС"
    ws["A4"] = "Параметры назначений на дату,"

    pale_yellow = PatternFill(fill_type="solid", fgColor="FFFFC0")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name="Arial", size=10, bold=True)
    data_font = Font(name="Arial", size=10)

    for column, header in enumerate(STAFF_DETAIL_HEADERS, start=1):
        cell = ws.cell(row=7, column=column, value=header)
        cell.font = header_font
        cell.fill = pale_yellow
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        number_cell = ws.cell(row=8, column=column, value=str(column))
        number_cell.font = header_font
        number_cell.fill = pale_yellow
        number_cell.border = border
        number_cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = build_staff_detail_rows(ctx, result)
    for row_index, row in enumerate(rows, start=9):
        for column, header in enumerate(STAFF_DETAIL_HEADERS, start=1):
            cell = ws.cell(row=row_index, column=column, value=row[header])
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if column in (10, 11, 12, 13):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            elif column in (14, 15, 21, 23) and cell.value:
                cell.number_format = "dd.mm.yyyy"
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row_index].height = 30

    widths = [8, 14, 49, 7, 27, 28, 13, 9, 18, 9, 16, 14, 13, 14, 14, 9, 31, 13, 33, 14, 14, 12, 14]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.row_dimensions[7].height = 39
    ws.freeze_panes = "A9"
    ws.auto_filter.ref = f"A7:W{max(8, ws.max_row)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "F4B183"
    wb.save(path)


__all__ = [
    "SHEET_STAFF_DETAIL",
    "STAFF_DETAIL_HEADERS",
    "build_staff_detail_rows",
    "write_staff_detail_sheet",
]
