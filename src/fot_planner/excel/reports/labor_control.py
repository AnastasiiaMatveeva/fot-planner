"""Лист «Контроль трудоёмкости» — единый пользовательский отчёт по трудоёмкости."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fot_planner.fot_schedule import active_months_in_year
from fot_planner.labor_rules import labor_rows_for_contract, planned_labor_amount
from fot_planner.models import PlanningContext, PlanningResult, labor_row_id
from fot_planner.excel.reports.user_report import (
    GROUP_DONE_COMMENT,
    MONTH_SHORT,
    PAYMENT_WITHOUT_PM_MSG,
    _labor_avg_value,
    _labor_row_checks,
)

SHEET_LABOR_CONTROL = "Контроль трудоёмкости"

MONTHS = list(range(1, 13))
MONTH_LABELS = [MONTH_SHORT[m] for m in MONTHS] + ["Итого"]

# Colors
FILL_HEADER = PatternFill(fill_type="solid", fgColor="1F4E78")
FILL_CONTRACT = PatternFill(fill_type="solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill(fill_type="solid", fgColor="BDD7EE")
FILL_PLAN = PatternFill(fill_type="solid", fgColor="DDEBF7")
FILL_FACT = PatternFill(fill_type="solid", fgColor="E2EFDA")
FILL_INACTIVE = PatternFill(fill_type="solid", fgColor="F2F2F2")
FILL_OK = PatternFill(fill_type="solid", fgColor="C6EFCE")
FILL_WARN = PatternFill(fill_type="solid", fgColor="FFEB9C")
FILL_ERR = PatternFill(fill_type="solid", fgColor="FFC7CE")
FILL_GROUP_OK = PatternFill(fill_type="solid", fgColor="C6EFCE")
FILL_GROUP_WARN = PatternFill(fill_type="solid", fgColor="FFEB9C")

FONT_HEADER = Font(color="FFFFFF", bold=True)
FONT_BOLD = Font(bold=True)
FONT_CONTRACT = Font(bold=True, size=12)

THIN = Side(style="thin", color="BFBFBF")
BORDER_TOP = Border(top=Side(style="medium", color="1F4E78"))

EMPLOYEE_METRICS = (
    ("pm", "Закрыто чел.-мес."),
    ("salary", "Оклад с договора"),
    ("allowance", "Надбавка с договора"),
    ("incentive", "Стимулирующая с договора"),
    ("total", "Всего денег на строку"),
)


@dataclass
class RowBlock:
    label: str
    group: str
    is_group_total: bool = False
    checks: dict = field(default_factory=dict)
    plan_pm_by_month: dict[int, float] = field(default_factory=dict)
    fact_pm_by_month: dict[int, float] = field(default_factory=dict)
    plan_amt_by_month: dict[int, float] = field(default_factory=dict)
    fact_amt_by_month: dict[int, float] = field(default_factory=dict)
    status: str = ""
    comment: str = ""
    pm_ind: str = "ОК"
    amt_ind: str = "ОК"
    avg_ind: str = "ОК"
    group_compensated: bool = False
    has_attribution_error: bool = False


def _distribute_evenly(total: float, active: list[int]) -> dict[int, float]:
    out = {m: 0.0 for m in MONTHS}
    if not active or total <= 0:
        return out
    per = total / len(active)
    for m in active:
        out[m] = per
    return out


def _month_series(values: dict[int, float], *, is_avg: bool = False) -> list:
    row: list = []
    total_pm = sum(values.values()) if not is_avg else 0.0
    total_amt = 0.0
    for m in MONTHS:
        v = values.get(m, 0.0)
        if is_avg:
            row.append(round(v, 2) if v and abs(v) > 0.005 else "")
        else:
            row.append(round(v, 2) if abs(v) > 0.005 else "")
            total_amt += v if not is_avg else 0
    if is_avg:
        return row
    total = sum(values.get(m, 0.0) for m in MONTHS)
    row.append(round(total, 2) if abs(total) > 0.005 else "")
    return row


def _month_avg_series(amounts: dict[int, float], pms: dict[int, float]) -> list:
    row: list = []
    total_amt = 0.0
    total_pm = 0.0
    for m in MONTHS:
        pm = pms.get(m, 0.0)
        amt = amounts.get(m, 0.0)
        if pm > 0.005:
            row.append(round(amt / pm, 2))
            total_amt += amt
            total_pm += pm
        else:
            row.append("")
    if total_pm > 0.005:
        row.append(round(total_amt / total_pm, 2))
    else:
        row.append("")
    return row


def _indicator(ok: bool, deviation: float, base: float, tol: float) -> tuple[str, str]:
    if base <= 0 and abs(deviation) < 0.01:
        return "ОК", "ok"
    if ok and abs(deviation) < 0.02:
        return "ОК", "ok"
    if ok:
        return "Предупреждение", "warn"
    return "Ошибка", "err"


def _resolve_status(
    checks: dict,
    *,
    group_compensated: bool = False,
    has_attribution_error: bool = False,
) -> tuple[str, str]:
    if has_attribution_error:
        return "Ошибка распределения", "Выплата с договора не полностью отнесена на трудоёмкость"
    if checks["row_ok"]:
        return "Выполнено", "Выполнено"
    if group_compensated:
        return "Выполнено по группе", GROUP_DONE_COMMENT
    if checks["pm_ok"] and checks["amount_ok"] and not checks["avg_ok"]:
        return "Отклонение по средней", "Фактическая средняя отличается от плановой"
    if not checks["pm_ok"] and checks["amount_ok"]:
        return "Отклонение по чел.-мес.", "Деньги отнесены, но чел.-мес. не закрыты"
    if checks["pm_ok"] and not checks["amount_ok"]:
        return "Отклонение по сумме", "Чел.-мес. закрыты, но ФОТ по строке не сошёлся"
    return "Отклонение по сумме", "Отклонение по строке трудоёмкости"


def _contract_has_attribution_errors(ctx: PlanningContext, result: PlanningResult, contract_id: str) -> bool:
    from fot_planner.labor_rules import contract_has_labor_plan

    if not contract_has_labor_plan(ctx, contract_id):
        return False
    attributed: dict[tuple[str, str, int, str], float] = defaultdict(float)
    for rec in result.labor_payment_attributions:
        key = (rec.employee_id, rec.contract_id, rec.month, rec.payment_kind)
        attributed[key] += rec.amount
    for a in result.allocations:
        if a.contract_id != contract_id or a.amount <= 0.005:
            continue
        key = (a.employee_id, a.contract_id, a.month, a.payment_kind)
        if abs(attributed.get(key, 0.0) - a.amount) > 0.02:
            return True
    return False


def _collect_labor_data(ctx: PlanningContext, result: PlanningResult) -> dict[str, dict]:
    tol = ctx.salary_stability.goz_labor_tolerance

    pm_by: dict[tuple[str, str, int], float] = defaultdict(float)
    for rec in result.labor_pm_attributions:
        pm_by[(rec.labor_row_id, rec.employee_id, rec.month)] += rec.person_months

    pay_by: dict[tuple[str, str, int, str], float] = defaultdict(float)
    for rec in result.labor_payment_attributions:
        pay_by[(rec.labor_row_id, rec.employee_id, rec.month, rec.payment_kind)] += rec.amount

    contract_data: dict[str, dict] = {}
    all_specs: list[tuple[str, RowBlock, str]] = []
    group_members: dict[tuple[str, str], list[int]] = defaultdict(list)

    for contract in ctx.contracts:
        active = active_months_in_year(contract, ctx.year)
        line_blocks: list[RowBlock] = []
        attr_err = _contract_has_attribution_errors(ctx, result, contract.id)

        for _lp_idx, lp in labor_rows_for_contract(ctx, contract.id):
            row_id = labor_row_id(lp)
            group = lp.equivalence_group or ""
            plan_pm = lp.person_months
            plan_amount = planned_labor_amount(lp)
            plan_pm_m = _distribute_evenly(plan_pm, active)
            plan_amt_m = _distribute_evenly(plan_amount, active)

            fact_pm_m = {m: 0.0 for m in MONTHS}
            fact_amt_m = {m: 0.0 for m in MONTHS}
            for (rid, _e, month), pm in pm_by.items():
                if rid == row_id:
                    fact_pm_m[month] += pm
            for (rid, _e, month, _kind), amt in pay_by.items():
                if rid == row_id:
                    fact_amt_m[month] += amt

            fact_pm = sum(fact_pm_m.values())
            fact_amt = sum(fact_amt_m.values())
            checks = _labor_row_checks(plan_pm, plan_amount, fact_pm, fact_amt, tol)
            label = lp.position or group or "—"
            block = RowBlock(
                label=label,
                group=group,
                checks=checks,
                plan_pm_by_month=plan_pm_m,
                fact_pm_by_month=fact_pm_m,
                plan_amt_by_month=plan_amt_m,
                fact_amt_by_month=fact_amt_m,
                has_attribution_error=attr_err,
            )
            line_blocks.append(block)
            idx = len(all_specs)
            all_specs.append((contract.id, block, group))
            if group:
                group_members[(contract.id, group)].append(idx)

        group_ok: dict[tuple[str, str], bool] = {}
        group_all_ok: dict[tuple[str, str], bool] = {}
        for key, indices in group_members.items():
            if key[0] != contract.id or len(indices) < 2:
                continue
            plan_pm = sum(all_specs[i][1].checks["plan_pm"] for i in indices)
            plan_amt = sum(all_specs[i][1].checks["plan_amt"] for i in indices)
            fact_pm = sum(all_specs[i][1].checks["fact_pm"] for i in indices)
            fact_amt = sum(all_specs[i][1].checks["fact_amt"] for i in indices)
            gc = _labor_row_checks(plan_pm, plan_amt, fact_pm, fact_amt, tol)
            group_ok[key] = gc["row_ok"]
            group_all_ok[key] = all(all_specs[i][1].checks["row_ok"] for i in indices)

        for block in line_blocks:
            gkey = (contract.id, block.group) if block.group else None
            compensated = bool(
                gkey
                and group_ok.get(gkey, False)
                and not group_all_ok.get(gkey, True)
            )
            block.group_compensated = compensated
            block.status, block.comment = _resolve_status(
                block.checks,
                group_compensated=compensated,
                has_attribution_error=block.has_attribution_error,
            )
            block.pm_ind, _ = _indicator(
                block.checks["pm_ok"], block.checks["pm_dev"], block.checks["plan_pm"], tol
            )
            block.amt_ind, _ = _indicator(
                block.checks["amount_ok"], block.checks["amt_dev"], block.checks["plan_amt"], tol
            )
            avg_dev = (block.checks["fact_avg"] or 0) - (block.checks["plan_avg"] or 0)
            block.avg_ind, _ = _indicator(
                block.checks["avg_ok"], avg_dev, block.checks["plan_avg"] or 0, tol
            )

        group_blocks: list[RowBlock] = []
        for (cid, group), indices in sorted(group_members.items()):
            if cid != contract.id or len(indices) < 2:
                continue
            plan_pm_m = {m: 0.0 for m in MONTHS}
            fact_pm_m = {m: 0.0 for m in MONTHS}
            plan_amt_m = {m: 0.0 for m in MONTHS}
            fact_amt_m = {m: 0.0 for m in MONTHS}
            plan_pm = plan_amt = fact_pm = fact_amt = 0.0
            for i in indices:
                b = all_specs[i][1]
                plan_pm += b.checks["plan_pm"]
                plan_amt += b.checks["plan_amt"]
                fact_pm += b.checks["fact_pm"]
                fact_amt += b.checks["fact_amt"]
                for m in MONTHS:
                    plan_pm_m[m] += b.plan_pm_by_month[m]
                    fact_pm_m[m] += b.fact_pm_by_month[m]
                    plan_amt_m[m] += b.plan_amt_by_month[m]
                    fact_amt_m[m] += b.fact_amt_by_month[m]
            checks = _labor_row_checks(plan_pm, plan_amt, fact_pm, fact_amt, tol)
            g_compensated = group_ok.get((cid, group), False) and not group_all_ok.get((cid, group), True)
            status, comment = _resolve_status(checks)
            if g_compensated:
                comment = GROUP_DONE_COMMENT
            gb = RowBlock(
                label=f"ИТОГО группа {group}",
                group=group,
                is_group_total=True,
                checks=checks,
                plan_pm_by_month=plan_pm_m,
                fact_pm_by_month=fact_pm_m,
                plan_amt_by_month=plan_amt_m,
                fact_amt_by_month=fact_amt_m,
                status=status,
                comment=comment,
                group_compensated=g_compensated,
                has_attribution_error=attr_err,
            )
            gb.pm_ind, _ = _indicator(checks["pm_ok"], checks["pm_dev"], checks["plan_pm"], tol)
            gb.amt_ind, _ = _indicator(checks["amount_ok"], checks["amt_dev"], checks["plan_amt"], tol)
            avg_dev = (checks["fact_avg"] or 0) - (checks["plan_avg"] or 0)
            gb.avg_ind, _ = _indicator(checks["avg_ok"], avg_dev, checks["plan_avg"] or 0, tol)
            group_blocks.append(gb)

        if line_blocks:
            contract_data[contract.id] = {
                "line_blocks": line_blocks,
                "group_blocks": group_blocks,
                "group_members": {
                    k: v for k, v in group_members.items() if k[0] == contract.id and len(v) >= 2
                },
            }

    return contract_data


def _collect_employee_rows(
    ctx: PlanningContext, result: PlanningResult, contract_id: str
) -> list[dict]:
    employees = {e.id: e for e in ctx.employees}
    keys: set[tuple[str, str, str, str]] = set()

    for rec in result.labor_pm_attributions:
        if rec.contract_id == contract_id:
            keys.add((rec.employee_id, rec.labor_row_id, rec.position or "", rec.equivalence_group or ""))
    for rec in result.labor_payment_attributions:
        if rec.contract_id == contract_id:
            keys.add((rec.employee_id, rec.labor_row_id, rec.position or "", rec.equivalence_group or ""))

    rows: list[dict] = []
    for e_id, row_id, row_pos, row_group in sorted(keys):
        emp = employees.get(e_id)
        group = row_group or (emp.equivalence_group if emp else "")
        name = emp.full_name if emp else e_id
        position = emp.position if emp else ""
        pm_m = {
            m: sum(
                rec.person_months
                for rec in result.labor_pm_attributions
                if rec.contract_id == contract_id
                and rec.employee_id == e_id
                and rec.labor_row_id == row_id
                and rec.month == m
            )
            for m in MONTHS
        }
        sal_m = {
            m: sum(
                rec.amount
                for rec in result.labor_payment_attributions
                if rec.contract_id == contract_id
                and rec.employee_id == e_id
                and rec.labor_row_id == row_id
                and rec.month == m
                and rec.payment_kind == "salary"
            )
            for m in MONTHS
        }
        alw_m = {
            m: sum(
                rec.amount
                for rec in result.labor_payment_attributions
                if rec.contract_id == contract_id
                and rec.employee_id == e_id
                and rec.labor_row_id == row_id
                and rec.month == m
                and rec.payment_kind == "allowance"
            )
            for m in MONTHS
        }
        inc_m = {
            m: sum(
                rec.amount
                for rec in result.labor_payment_attributions
                if rec.contract_id == contract_id
                and rec.employee_id == e_id
                and rec.labor_row_id == row_id
                and rec.month == m
                and rec.payment_kind == "incentive"
            )
            for m in MONTHS
        }
        tot_m = {m: sal_m[m] + alw_m[m] + inc_m[m] for m in MONTHS}
        metrics = {
            "pm": pm_m,
            "salary": sal_m,
            "allowance": alw_m,
            "incentive": inc_m,
            "total": tot_m,
        }
        for key, label in EMPLOYEE_METRICS:
            rows.append(
                {
                    "contract_id": contract_id,
                    "row_label": row_pos or row_id,
                    "group": group,
                    "employee": name,
                    "position": position,
                    "metric": label,
                    "values": metrics[key],
                    "status_by_month": {
                        m: (
                            PAYMENT_WITHOUT_PM_MSG
                            if tot_m[m] > 0.005 and pm_m[m] <= 0.005 and label == "Всего денег на строку"
                            else ""
                        )
                        for m in MONTHS
                    },
                }
            )
    return rows


def _fill_indicator(cell, level: str) -> None:
    if level == "ok":
        cell.fill = FILL_OK
    elif level == "warn":
        cell.fill = FILL_WARN
    else:
        cell.fill = FILL_ERR


def _fill_deviation(cell, deviation: float, ok: bool, tol_base: float, tol: float) -> None:
    if abs(deviation) < 0.02:
        cell.fill = FILL_OK
    elif ok:
        cell.fill = FILL_WARN
    elif tol_base > 0 and abs(deviation) <= tol * tol_base:
        cell.fill = FILL_WARN
    else:
        cell.fill = FILL_ERR


def _write_summary_row(ws, row_idx: int, block: RowBlock, tol: float, *, inactive: set[int]) -> None:
    c = block.checks
    avg_dev = (c["fact_avg"] or 0) - (c["plan_avg"] or 0) if c["plan_avg"] else 0
    values = [
        block.label,
        round(c["plan_pm"], 2),
        round(c["fact_pm"], 2),
        round(c["pm_dev"], 2),
        round(c["plan_amt"], 0),
        round(c["fact_amt"], 0),
        round(c["amt_dev"], 0),
        round(c["plan_avg"], 2) if c["plan_avg"] else "",
        round(c["fact_avg"], 2) if c["fact_avg"] else "",
        round(avg_dev, 2) if c["plan_avg"] else "",
        block.pm_ind,
        block.amt_ind,
        block.avg_ind,
        block.status,
        block.comment,
    ]
    font = FONT_BOLD if block.is_group_total else Font()
    fill = FILL_GROUP_OK if block.is_group_total and block.checks["row_ok"] else None
    if block.is_group_total and block.group_compensated:
        fill = FILL_GROUP_WARN
    if block.is_group_total and not block.checks["row_ok"] and not block.group_compensated:
        fill = FILL_ERR if block.status != "Выполнено" else FILL_GROUP_OK
    if block.group_compensated and not block.is_group_total:
        fill = FILL_WARN

    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill and col <= 14:
            cell.fill = fill
        if col == 2:
            cell.fill = FILL_PLAN
        elif col == 3:
            cell.fill = FILL_FACT
        elif col == 4:
            _fill_deviation(cell, c["pm_dev"], c["pm_ok"], c["plan_pm"], tol)
        elif col == 5:
            cell.fill = FILL_PLAN
        elif col == 6:
            cell.fill = FILL_FACT
        elif col == 7:
            _fill_deviation(cell, c["amt_dev"], c["amount_ok"], c["plan_amt"], tol)
        elif col == 8:
            cell.fill = FILL_PLAN
        elif col == 9:
            cell.fill = FILL_FACT
        elif col == 10:
            _fill_deviation(cell, avg_dev, c["avg_ok"], c["plan_avg"] or 0, tol)
        elif col == 11:
            _fill_indicator(cell, "ok" if block.pm_ind == "ОК" else ("warn" if block.pm_ind == "Предупреждение" else "err"))
        elif col == 12:
            _fill_indicator(cell, "ok" if block.amt_ind == "ОК" else ("warn" if block.amt_ind == "Предупреждение" else "err"))
        elif col == 13:
            _fill_indicator(cell, "ok" if block.avg_ind == "ОК" else ("warn" if block.avg_ind == "Предупреждение" else "err"))
        if col in (2, 3, 4) and isinstance(val, float):
            cell.number_format = "0.00"
        if col in (5, 6, 7, 8, 9, 10) and isinstance(val, float):
            cell.number_format = "#,##0"

    if block.is_group_total:
        for col in range(1, 16):
            ws.cell(row=row_idx, column=col).border = BORDER_TOP


def _write_monthly_matrix(ws, start_row: int, block: RowBlock, active: set[int]) -> int:
    headers = ["Показатель"] + MONTH_LABELS
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.alignment = Alignment(horizontal="center")

    matrix_rows = [
        ("Чел.-мес. план", _month_series(block.plan_pm_by_month), "plan"),
        ("Чел.-мес. факт", _month_series(block.fact_pm_by_month), "fact"),
        ("Сумма план", _month_series(block.plan_amt_by_month), "plan"),
        ("Сумма факт", _month_series(block.fact_amt_by_month), "fact"),
        (
            "Средняя факт",
            _month_avg_series(block.fact_amt_by_month, block.fact_pm_by_month),
            "fact",
        ),
    ]
    r = start_row + 1
    for label, data, kind in matrix_rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        for i, val in enumerate(data):
            col = i + 2
            month_num = i + 1 if i < 12 else None
            cell = ws.cell(row=r, column=col, value=val if val != "" else None)
            if month_num and month_num not in active and i < 12:
                cell.fill = FILL_INACTIVE
            elif kind == "plan" and val != "":
                cell.fill = FILL_PLAN
            elif kind == "fact" and val != "":
                cell.fill = FILL_FACT
            if isinstance(val, float):
                cell.number_format = "#,##0" if "Сумма" in label or "Средняя" in label else "0.00"
        r += 1
    return r + 1


def _write_employee_section(ws, start_row: int, ctx: PlanningContext, result: PlanningResult, contract_id: str, active: set[int]) -> int:
    ws.cell(row=start_row, column=1, value="Что сформировало трудоёмкость").font = Font(bold=True, size=11)
    start_row += 1
    headers = ["Договор", "Строка", "Группа", "Сотрудник", "Должность", "Показатель"] + MONTH_LABELS
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
    start_row += 1
    for er in _collect_employee_rows(ctx, result, contract_id):
        series = _month_series(er["values"])
        row_vals = [
            contract_id,
            er.get("row_label", ""),
            er["group"],
            er["employee"],
            er["position"],
            er["metric"],
        ] + series
        for col, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=start_row, column=col, value=val if val != "" else None)
            if col >= 7 and (col - 6) <= 12:
                m = col - 6
                if m not in active:
                    cell.fill = FILL_INACTIVE
                elif er.get("status_by_month", {}).get(m) == PAYMENT_WITHOUT_PM_MSG:
                    cell.fill = FILL_ERR
            if isinstance(val, float):
                if er["metric"] == "Закрыто чел.-мес.":
                    cell.number_format = "0.00"
                else:
                    cell.number_format = "#,##0"
        start_row += 1
    return start_row + 1


def write_labor_control_sheet(path, ctx: PlanningContext, result: PlanningResult) -> None:
    wb = load_workbook(path)
    if SHEET_LABOR_CONTROL in wb.sheetnames:
        del wb[SHEET_LABOR_CONTROL]
    for old in ("Трудоёмкость по строкам", "Расшифровка трудоёмкости"):
        if old in wb.sheetnames:
            del wb[old]

    ws = wb.create_sheet(SHEET_LABOR_CONTROL, 4)
    ws.sheet_properties.tabColor = "70AD47"
    ws.sheet_view.showGridLines = False

    tol = ctx.salary_stability.goz_labor_tolerance
    contracts = {c.id: c for c in ctx.contracts}
    contract_data = _collect_labor_data(ctx, result)

    summary_headers = [
        "Строка / группа",
        "План ч/м",
        "Факт ч/м",
        "Δ ч/м",
        "План сумма",
        "Факт сумма",
        "Δ сумма",
        "Средняя план",
        "Средняя факт",
        "Δ средняя",
        "Ч/м",
        "ФОТ",
        "Средняя",
        "Статус",
        "Комментарий",
    ]
    col_widths = {
        1: 28,
        2: 12,
        3: 12,
        4: 12,
        5: 14,
        6: 14,
        7: 12,
        8: 14,
        9: 14,
        10: 12,
        11: 8,
        12: 8,
        13: 10,
        14: 22,
        15: 48,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for col in range(16, 29):
        ws.column_dimensions[get_column_letter(col)].width = 13

    row = 1
    if not contract_data:
        ws.cell(row=1, column=1, value="Нет строк трудоёмкости для контроля")
        wb.save(path)
        return

    for contract_id, data in contract_data.items():
        contract = contracts[contract_id]
        active = set(active_months_in_year(contract, ctx.year))
        line_blocks = data["line_blocks"]
        group_blocks = data["group_blocks"]
        group_members = data["group_members"]

        title = f"ДОГОВОР: {contract_id} — {contract.name}"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = FONT_CONTRACT
        cell.fill = FILL_CONTRACT
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22
        row += 1

        for col, h in enumerate(summary_headers, start=1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_HEADER
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[row].height = 26
        row += 1

        for block in line_blocks:
            _write_summary_row(ws, row, block, tol, inactive=set(MONTHS) - active)
            row += 1
            row = _write_monthly_matrix(ws, row, block, active)
            row += 1

        if group_blocks:
            for (cid, group), indices in sorted(group_members.items()):
                ws.cell(row=row, column=1, value=f"Баланс группы: {group}").font = FONT_BOLD
                row += 1
                compact_headers = summary_headers[:10] + ["Статус", "Комментарий"]
                for col, h in enumerate(compact_headers, start=1):
                    c = ws.cell(row=row, column=col, value=h)
                    c.font = FONT_BOLD
                    c.fill = FILL_TOTAL
                row += 1
                group_line_blocks = [
                    b for b in line_blocks if b.group == group and not b.is_group_total
                ]
                for block in group_line_blocks:
                    _write_summary_row(ws, row, block, tol, inactive=set(MONTHS) - active)
                    row += 1
                for gb in group_blocks:
                    if gb.group == group:
                        _write_summary_row(ws, row, gb, tol, inactive=set(MONTHS) - active)
                        row += 1
                row += 1

        row = _write_employee_section(ws, row, ctx, result, contract_id, active)
        row += 1

    ws.freeze_panes = "B2"
    if ws.max_row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(15)}1"
    wb.save(path)


__all__ = ["SHEET_LABOR_CONTROL", "write_labor_control_sheet"]
