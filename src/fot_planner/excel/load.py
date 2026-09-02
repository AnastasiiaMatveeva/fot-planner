from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path

import pandas as pd

from fot_planner.excel.constants import (
    SHEET_CONTRACT_LABOR,
    SHEET_CONTRACTS,
    SHEET_EMPLOYEES,
    SHEET_FOT_LOCK_MATRIX,
    SHEET_FOT_MATRIX,
    SHEET_MANUAL_ASSIGNMENTS,
    SHEET_MANUAL_PROHIBITIONS,
    SHEET_MIN_BALANCE_MATRIX,
    SHEET_PLAN,
    SHEET_POSITION_LIMITS,
    SHEET_SUBSTITUTIONS,
    SHEET_SECRET_ALLOWANCES,
    SHEET_SETTINGS,
)
from fot_planner.excel.parsing import (
    _bool,
    _canonicalize_columns,
    _optional_bool,
    _optional_float,
    _optional_int,
    _parse_date,
    _resolve_monthly_wage,
    _split_list,
    parse_payment_kind,
)
from fot_planner.fot_schedule import default_fot_inflow_at_start
from fot_planner.open_rate_rules import (
    normalize_employment_category,
    normalize_employment_type,
)
from fot_planner.models import (
    AllocationRecord,
    Contract,
    ContractLaborPlan,
    ContractMonthlyBudget,
    Employee,
    ManualAssignment,
    ManualProhibition,
    OptimizationWeights,
    PaymentKind,
    PlanningContext,
    PositionReference,
    SalaryStabilityRules,
    SecretAllowance,
    labor_row_id,
)
from fot_planner.position_reference import (
    PositionReferenceRow,
    build_position_index,
    default_position_reference,
    default_position_synonyms,
    normalize_position,
    resolve_position,
    salary_equivalence_group,
)
from fot_planner.salary_limits_2556 import (
    PositionSalaryLimit,
    default_position_salary_limits,
    position_limit_tables_from_salary_limits,
)

def load_context(path: str | Path, plan_path: str | Path | None = None) -> PlanningContext:
    path = Path(path)
    xl = pd.ExcelFile(path)

    year = date.today().year
    settings_df: pd.DataFrame | None = None
    if SHEET_SETTINGS in xl.sheet_names:
        settings_df = _canonicalize_columns(
            pd.read_excel(xl, SHEET_SETTINGS),
            SHEET_SETTINGS,
        )
        if not settings_df.empty and "year" in settings_df.columns:
            year = int(settings_df.iloc[0]["year"])

    position_limits_df = _load_position_limits_sheet(xl)
    defaults_used: list[tuple[str, str, object]] = []
    if position_limits_df is not None:
        position_reference_rows = _position_reference_from_limits_sheet(
            position_limits_df, defaults_used
        )
        position_salary_limit_defaults = _position_salary_limits_from_limits_sheet(
            position_limits_df, defaults_used
        )
    else:
        # Листа нет вовсе — весь справочник встроенный. Это не подстановка
        # отдельных величин, а работа на умолчаниях целиком, и сказать об этом
        # надо одной фразой, а не тридцатью тремя.
        position_reference_rows = list(default_position_reference())
        position_salary_limit_defaults = default_position_salary_limits()
        defaults_used.append(("", "весь справочник должностей", len(position_reference_rows)))
    position_synonyms = default_position_synonyms()
    position_index = build_position_index(position_reference_rows, position_synonyms)
    position_reference = [
        PositionReference(
            position=row.position,
            equivalence_group=row.equivalence_group,
            level=row.level,
            reference_salary_for_rate=row.reference_salary_for_rate,
            salary_page=row.salary_page,
            salary_group_number=row.salary_group_number,
        )
        for row in position_reference_rows
    ]
    substitution_rules = _load_substitutions(xl, position_index)
    employees = _load_employees(
        _canonicalize_columns(pd.read_excel(xl, SHEET_EMPLOYEES), SHEET_EMPLOYEES),
        position_index,
        substitution_rules,
    )
    contracts = _load_contracts(
        _canonicalize_columns(pd.read_excel(xl, SHEET_CONTRACTS), SHEET_CONTRACTS),
    )
    secret_allowances: list[SecretAllowance] = []
    if SHEET_SECRET_ALLOWANCES in xl.sheet_names:
        secret_allowances = _load_secret_allowances(
            _canonicalize_columns(
                pd.read_excel(xl, SHEET_SECRET_ALLOWANCES),
                SHEET_SECRET_ALLOWANCES,
            )
        )
    position_salary_limits = position_salary_limit_defaults
    position_limit_tables = position_limit_tables_from_salary_limits(position_salary_limits)

    if SHEET_FOT_MATRIX in xl.sheet_names:
        amounts_df = _canonicalize_columns(
            pd.read_excel(xl, SHEET_FOT_MATRIX),
            SHEET_FOT_MATRIX,
        )
        lock_df = None
        if SHEET_FOT_LOCK_MATRIX in xl.sheet_names:
            lock_df = _canonicalize_columns(
                pd.read_excel(xl, SHEET_FOT_LOCK_MATRIX),
                SHEET_FOT_LOCK_MATRIX,
            )
        fill_locks = _read_fot_fill_locks(path, SHEET_FOT_MATRIX)
        for c in contracts:
            c.monthly_budgets = []
        _apply_matrix_budgets(
            contracts,
            amounts_df,
            lock_df,
            year,
            fill_locks=fill_locks,
        )
    else:
        default_fot_inflow_at_start(contracts, year)

    if SHEET_MIN_BALANCE_MATRIX in xl.sheet_names:
        _apply_matrix_min_balances(
            contracts,
            _canonicalize_columns(
                pd.read_excel(xl, SHEET_MIN_BALANCE_MATRIX),
                SHEET_MIN_BALANCE_MATRIX,
            ),
            year,
        )

    labor_plans: list[ContractLaborPlan] = []
    if SHEET_CONTRACT_LABOR in xl.sheet_names:
        labor_plans = _load_contract_labor(
            _canonicalize_columns(
                pd.read_excel(xl, SHEET_CONTRACT_LABOR),
                SHEET_CONTRACT_LABOR,
            ),
            year,
            position_index,
        )

    manual_assignments: list[ManualAssignment] = []
    manual_prohibitions: list[ManualProhibition] = []
    if SHEET_MANUAL_ASSIGNMENTS in xl.sheet_names:
        manual_assignments = _load_manual_assignments(
            _canonicalize_columns(
                pd.read_excel(xl, SHEET_MANUAL_ASSIGNMENTS),
                SHEET_MANUAL_ASSIGNMENTS,
            )
        )
    if SHEET_MANUAL_PROHIBITIONS in xl.sheet_names:
        manual_prohibitions = _load_manual_prohibitions(
            _canonicalize_columns(
                pd.read_excel(xl, SHEET_MANUAL_PROHIBITIONS),
                SHEET_MANUAL_PROHIBITIONS,
            )
        )

    baseline_plan: list[AllocationRecord] | None = None
    plan_source = plan_path or path
    plan_source = Path(plan_source)
    if plan_source.exists() and plan_source.suffix.lower() in (".xlsx", ".xlsm"):
        with pd.ExcelFile(plan_source) as plan_xl:
            if SHEET_PLAN in plan_xl.sheet_names:
                baseline_plan = _load_plan_overrides(
                    _canonicalize_columns(pd.read_excel(plan_xl, SHEET_PLAN), SHEET_PLAN),
                    year,
                )

    weights = OptimizationWeights()
    salary_stability = SalaryStabilityRules()
    allow_deficit = False
    if settings_df is not None:
        weights = _load_weights(settings_df)
        salary_stability = _load_salary_stability(settings_df)
        if not settings_df.empty:
            row = settings_df.iloc[0]
            if "allow_deficit" in row and pd.notna(row["allow_deficit"]):
                allow_deficit = _bool(row["allow_deficit"])
    position_bep_limit = _bep_limit_from_position_limits(position_salary_limits)
    if position_bep_limit is not None:
        salary_stability.goz_average_salary_limit = position_bep_limit

    ctx = PlanningContext(
        year=year,
        employees=employees,
        contracts=contracts,
        position_reference=position_reference,
        position_salary_limits=position_salary_limits,
        position_limit_tables=position_limit_tables,
        manual_assignments=manual_assignments,
        manual_prohibitions=manual_prohibitions,
        weights=weights,
        salary_stability=salary_stability,
        labor_plans=labor_plans,
        substitution_rules=substitution_rules,
        defaults_used=defaults_used,
        secret_allowances=secret_allowances,
        baseline_plan=baseline_plan,
        allow_deficit=allow_deficit,
    )
    xl.close()
    return ctx


def _clean_optional_text(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _load_position_limits_sheet(xl: pd.ExcelFile) -> pd.DataFrame | None:
    if SHEET_POSITION_LIMITS not in xl.sheet_names:
        return None
    return _canonicalize_columns(
        pd.read_excel(xl, SHEET_POSITION_LIMITS),
        SHEET_POSITION_LIMITS,
    )


def _reference_group_from_fields(
    *,
    explicit_group: str | None,
    salary_page: str | None,
    salary_group_number: int | None,
    level: int | None,
    default: PositionReferenceRow | None = None,
    fallback: str | None = None,
) -> str:
    if explicit_group:
        return explicit_group
    if salary_page or salary_group_number is not None or level is not None:
        return salary_equivalence_group(
            salary_page,
            salary_group_number,
            level,
        )
    if default is not None and default.equivalence_group:
        return default.equivalence_group
    return fallback or "без окладной группы"


def _position_reference_from_limits_sheet(
    df: pd.DataFrame, defaults_used: list | None = None
) -> list[PositionReferenceRow]:
    defaults = {
        normalize_position(row.position): row
        for row in default_position_reference()
    }
    log = defaults_used if defaults_used is not None else []
    rows: list[PositionReferenceRow] = []
    for _, r in df.iterrows():
        position = _clean_optional_text(r.get("position"))
        if not position:
            continue
        default = defaults.get(normalize_position(position))

        def taken(column, sheet_value, default_value, title):
            """Значение и отметка, если оно пришло из встроенного списка."""
            if column in r.index:
                return sheet_value
            if default_value is not None:
                log.append((position, title, default_value))
            return default_value

        salary_page = taken(
            "salary_page", _clean_optional_text(r.get("salary_page")),
            default.salary_page if default else None, "страница")
        salary_group_number = taken(
            "salary_group_number", _optional_int(r.get("salary_group_number")),
            default.salary_group_number if default else None, "номер группы")
        level = taken(
            "position_level", _optional_int(r.get("position_level")),
            default.level if default else None, "номер уровня")
        reference_salary_for_rate = taken(
            "reference_salary_for_rate",
            _optional_float(r.get("reference_salary_for_rate")),
            default.reference_salary_for_rate if default else None, "оклад")

        # Отдельный случай: колонки в листе есть, но все три пусты. Тогда
        # окладная группа берется из встроенного списка внутри
        # _reference_group_from_fields — и это самая незаметная подстановка.
        if (not _clean_optional_text(r.get("equivalence_group"))
                and salary_page is None and salary_group_number is None
                and level is None
                and default is not None and default.equivalence_group
                and not default.equivalence_group.startswith("должность: ")):
            log.append((position, "окладная группа", default.equivalence_group))

        group = _reference_group_from_fields(
            explicit_group=_clean_optional_text(r.get("equivalence_group")),
            salary_page=salary_page,
            salary_group_number=salary_group_number,
            level=level,
            default=default,
            fallback=position,
        )
        rows.append(
            PositionReferenceRow(
                position=position,
                equivalence_group=group,
                level=level,
                reference_salary_for_rate=reference_salary_for_rate,
                salary_page=salary_page,
                salary_group_number=salary_group_number,
            )
        )
    return rows


def _position_salary_limits_from_limits_sheet(
    df: pd.DataFrame, defaults_used: list | None = None
) -> list[PositionSalaryLimit]:
    defaults = {
        normalize_position(row.position): row
        for row in default_position_salary_limits()
    }
    log = defaults_used if defaults_used is not None else []
    rows: list[PositionSalaryLimit] = []
    for _, r in df.iterrows():
        position = _clean_optional_text(r.get("position"))
        if not position:
            continue
        default = defaults.get(normalize_position(position))
        for column, title in (("personnel_category", "категория персонала"),
                              ("order_2556_limit", "П2556"),
                              ("p4_limit", "П4"), ("bep_limit", "БЭП")):
            if column not in r.index:
                value = getattr(default, column, None) if default else None
                if value not in (None, ""):
                    log.append((position, title, value))
        rows.append(
            PositionSalaryLimit(
                position=position,
                personnel_category=(
                    (
                        _clean_optional_text(r.get("personnel_category"))
                        if "personnel_category" in r.index
                        else (default.personnel_category if default else "")
                    )
                    or ""
                ),
                order_2556_limit=(
                    _optional_float(r.get("order_2556_limit"))
                    if "order_2556_limit" in r.index
                    else (default.order_2556_limit if default else None)
                ),
                p4_limit=(
                    _optional_float(r.get("p4_limit"))
                    if "p4_limit" in r.index
                    else (default.p4_limit if default else None)
                ),
                bep_limit=(
                    _optional_float(r.get("bep_limit"))
                    if "bep_limit" in r.index
                    else (default.bep_limit if default else None)
                ),
                note=(
                    _clean_optional_text(r.get("note"))
                    if "note" in r.index
                    else (default.note if default else None)
                ),
            )
        )
    return rows


def _bep_limit_from_position_limits(
    position_salary_limits: list[PositionSalaryLimit],
) -> float | None:
    limits = [
        float(row.bep_limit)
        for row in position_salary_limits
        if row.bep_limit is not None and row.bep_limit > 0
    ]
    if not limits:
        return None
    return min(limits)


def _split_substitutes(value: object) -> list[str]:
    """Замещающие должности из одной ячейки.

    Отдел кадров перечисляет их через запятую — «Программист, Инженер 2 кат.,
    Лаборант», — а общий _split_list режет только по «;». Точку с запятой тоже
    принимаем: в чужих выгрузках встречается и она.
    """
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    text = str(value).replace(";", ",")
    return [part.strip() for part in text.split(",") if part.strip()]


def _load_substitutions(
    xl: pd.ExcelFile, position_index: dict | None = None
) -> dict[str, frozenset[str]]:
    """Правила замещения: должность сотрудника → на какие должности его можно
    поставить дополнительно.

    Правила направленные, и направление именно такое: строка «Директор |
    Научный сотрудник» означает, что директор может взять себе ставку научного
    сотрудника, а не что работу директора закроет научный сотрудник. Так
    названы и графы исходного файла: «исходная должность» — «должность,
    которая может быть». По этой же причине в строке «Инженер» стоит «Инженер
    (в другом подразделении)»: двух ставок по одной должности в одном
    подразделении у человека быть не может, поэтому вторую оформляют по
    соседней должности из этого списка.

    Симметричная окладная группа такого не выражает — она либо пускает обоих,
    либо никого. Перечисляются должности в одной ячейке через запятую, как в
    выгрузке отдела кадров. Лист необязательный: без него остаются прежние
    правила — точное совпадение должности и окладная группа.
    """
    def canonical(name: str) -> str:
        """Название должности так, как оно записано в справочнике.

        Должности сотрудников и так проходят через словарь синонимов —
        «Вед. инженер» становится «ведущий инженер». Правила замещения шли
        мимо него, и то же самое название в правиле не совпадало ни с чем:
        правило молча не работало. Приводим одинаково.
        """
        key = normalize_position(name)
        if position_index:
            row = position_index.get(key)
            if row is not None:
                return normalize_position(row.position)
        return key

    if SHEET_SUBSTITUTIONS not in xl.sheet_names:
        return {}
    df = _canonicalize_columns(
        pd.read_excel(xl, SHEET_SUBSTITUTIONS), SHEET_SUBSTITUTIONS
    )
    rules: dict[str, set[str]] = {}
    for _, r in df.iterrows():
        position = canonical(_clean_optional_text(r.get("position")) or "")
        if not position:
            continue
        for name in _split_substitutes(r.get("substitutes")):
            key = canonical(name)
            # Должность, замещающая сама себя, ничего не добавляет: точное
            # совпадение и так разрешено.
            if key and key != position:
                rules.setdefault(position, set()).add(key)
    return {k: frozenset(v) for k, v in rules.items() if v}


def _load_employees(
    df: pd.DataFrame,
    position_index: dict[str, PositionReferenceRow],
    substitution_rules: dict[str, frozenset[str]] | None = None,
) -> list[Employee]:
    rows: list[Employee] = []
    covers = substitution_rules or {}
    for _, r in df.iterrows():
        position = str(r["position"]).strip()
        ref = resolve_position(position, position_index)
        rows.append(
            Employee(
                id=str(r["id"]).strip(),
                full_name=str(r.get("full_name", r.get("fio", ""))).strip(),
                position=position,
                department=str(r.get("department", "")).strip(),
                rate=float(r["rate"]),
                monthly_wage=_resolve_monthly_wage(r),
                start_date=_parse_date(r.get("start_date")),
                end_date=_parse_date(r.get("end_date")),
                allowed_contracts=_split_list(r.get("allowed_contracts")),
                forbidden_contracts=_split_list(r.get("forbidden_contracts")),
                equivalence_group=ref.equivalence_group if ref else None,
                extra_positions=frozenset(
                    covers.get(normalize_position(position), frozenset())
                ),
                position_level=ref.level if ref else None,
                reference_salary_for_rate=ref.reference_salary_for_rate if ref else None,
                employment_type=normalize_employment_type(r.get("employment_type")),
                employment_category=normalize_employment_category(
                    r.get("employment_category")
                ),
            )
        )
    return rows


def _has_value(value: object) -> bool:
    return value is not None and not pd.isna(value) and str(value).strip() != ""


def _optional_deadline(row: pd.Series, column: str) -> date | None:
    if column not in row.index or not _has_value(row.get(column)):
        return None
    return _parse_date(row.get(column))


def _parse_contract_payment_deadlines(row: pd.Series) -> tuple[date | None, date | None]:
    return (
        _optional_deadline(row, "salary_payment_deadline"),
        _optional_deadline(row, "allowances_payment_deadline"),
    )


def _apply_payment_deadlines(contract: Contract) -> None:
    """Пустые даты в Excel → дата окончания договора (в срок)."""
    if contract.salary_payment_deadline is None:
        contract.salary_payment_deadline = contract.end_date
    if contract.allowances_payment_deadline is None:
        contract.allowances_payment_deadline = contract.end_date


def _load_contracts(df: pd.DataFrame) -> list[Contract]:
    rows: list[Contract] = []
    for _, r in df.iterrows():
        ctype = _clean_optional_text(r.get("contract_type")) or ""
        salary_payment_deadline, allowances_payment_deadline = _parse_contract_payment_deadlines(r)
        rows.append(
            Contract(
                id=str(r["id"]).strip(),
                name=str(r.get("name", "")).strip(),
                number=str(r.get("number", "")).strip(),
                contract_type=ctype,
                start_date=_parse_date(r["start_date"]),
                end_date=_parse_date(r["end_date"]),
                total_fot=float(r["total_fot"]),
                account=_clean_optional_text(r.get("account")) or "",
                is_goz_defense_order=_bool(r.get("is_goz_defense_order"), False),
                allow_salary=_bool(r.get("allow_salary"), True),
                allow_secret=_bool(r.get("allow_secret"), False),
                allow_allowance=_bool(r.get("allow_allowance"), True),
                allow_incentive=_bool(r.get("allow_incentive"), False),
                allow_extra_work=_bool(r.get("allow_extra_work"), False),
                allow_order_incentive=_bool(
                    r.get("allow_order_incentive"),
                    False,
                ),
                priority_payment_mode=_bool(
                    r.get("priority_payment_mode"),
                    False,
                ),
                salary_payment_deadline=salary_payment_deadline,
                allowances_payment_deadline=allowances_payment_deadline,
                allow_main_employment=_bool(r.get("allow_main_employment"), True),
                allow_part_time=_bool(r.get("allow_part_time"), True),
            )
        )
        _apply_payment_deadlines(rows[-1])
    return rows


def _normalize_secret_allowance_rate(value: object) -> float:
    if isinstance(value, str):
        text = value.strip()
        if text.endswith("%"):
            return float(text[:-1].replace(",", ".")) / 100
    rate = _optional_float(value)
    if rate is None:
        return 0.05
    if rate > 1:
        return rate / 100
    return rate


def _load_secret_allowances(df: pd.DataFrame) -> list[SecretAllowance]:
    rows: list[SecretAllowance] = []
    for _, r in df.iterrows():
        if not any(
            _has_value(r.get(col))
            for col in ("employee_id", "secret_contract_id", "rate")
        ):
            continue
        employee_id = str(r.get("employee_id", "")).strip()
        secret_contract_id = str(r.get("secret_contract_id", "")).strip()
        if not employee_id:
            raise ValueError(f"В листе {SHEET_SECRET_ALLOWANCES} не указан сотрудник")
        if not secret_contract_id:
            raise ValueError(
                f"В листе {SHEET_SECRET_ALLOWANCES} не указан договор секретности для сотрудника {employee_id!r}"
            )
        rows.append(
            SecretAllowance(
                employee_id=employee_id,
                secret_contract_id=secret_contract_id,
                rate=_normalize_secret_allowance_rate(r.get("rate")),
            )
        )
    return rows


_MONTH_ALIASES: dict[str, int] = {
    "1": 1,
    "2": 2,
    "3": 3,
    "4": 4,
    "5": 5,
    "6": 6,
    "7": 7,
    "8": 8,
    "9": 9,
    "10": 10,
    "11": 11,
    "12": 12,
    "янв": 1,
    "январь": 1,
    "фев": 2,
    "февраль": 2,
    "мар": 3,
    "март": 3,
    "апр": 4,
    "апрель": 4,
    "май": 5,
    "июн": 6,
    "июнь": 6,
    "июл": 7,
    "июль": 7,
    "авг": 8,
    "август": 8,
    "сен": 9,
    "сентябрь": 9,
    "окт": 10,
    "октябрь": 10,
    "ноя": 11,
    "ноябрь": 11,
    "дек": 12,
    "декабрь": 12,
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


def _month_from_column(name) -> int | None:
    if name is None or (isinstance(name, float) and pd.isna(name)):
        return None
    key = str(name).strip().lower().replace(".", "")
    if key in _MONTH_ALIASES:
        return _MONTH_ALIASES[key]
    try:
        m = int(float(key))
        if 1 <= m <= 12:
            return m
    except ValueError:
        pass
    return None


def _matrix_month_columns(df: pd.DataFrame) -> dict[int, str]:
    """Первая колонка — id проекта, остальные — месяцы 1–12."""
    mapping: dict[int, str] = {}
    for col in df.columns[1:]:
        month = _month_from_column(col)
        if month is not None:
            mapping[month] = col
    return mapping


_WHITE_FILL_RGB = frozenset({"FFFFFF", "FFFFFFFF", "00FFFFFF"})


def _parse_matrix_cell_value(val) -> tuple[float | None, bool]:
    """Сумма из ячейки; lock по суффиксу * или ! или (фикс)."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None, False
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return float(val), False

    locked = False
    s = str(val).strip()
    for marker in ("(фикс)", "(lock)", "(l)"):
        if s.lower().endswith(marker):
            locked = True
            s = s[: -len(marker)].strip()
            break
    if s.endswith(("*", "!")):
        locked = True
        s = s[:-1].strip()
    s = s.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if not s:
        return None, locked
    try:
        return float(s), locked
    except ValueError:
        return None, locked


def _cell_has_lock_fill(cell) -> bool:
    """Заливка ячейки (жёлтая и т.п.) = фиксация суммы на листе «фот_по_месяцам»."""
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return False
    if fill.fill_type != "solid":
        return True
    color = fill.fgColor
    if color is None:
        return False
    if color.type == "rgb" and color.rgb:
        rgb = color.rgb.upper()
        if len(rgb) == 8:
            rgb = rgb[2:]
        return rgb not in _WHITE_FILL_RGB
    if color.type == "indexed" and color.indexed is not None:
        return color.indexed not in (9, 64, 65)
    if color.type == "theme":
        return True
    return False


def _read_fot_fill_locks(path: Path, sheet_name: str) -> dict[tuple[str, int], bool]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    try:
        wb = load_workbook(path, data_only=False, read_only=False)
    except Exception:
        return {}
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {}

    ws = wb[sheet_name]
    month_by_col: dict[int, int] = {}
    for col_idx in range(2, ws.max_column + 1):
        month = _month_from_column(ws.cell(1, col_idx).value)
        if month is not None:
            month_by_col[col_idx] = month

    locks: dict[tuple[str, int], bool] = {}
    for row_idx in range(2, ws.max_row + 1):
        raw_id = ws.cell(row_idx, 1).value
        if raw_id is None:
            continue
        cid = str(raw_id).strip()
        if not cid:
            continue
        for col_idx, month in month_by_col.items():
            if _cell_has_lock_fill(ws.cell(row_idx, col_idx)):
                locks[(cid, month)] = True
    wb.close()
    return locks


def _apply_matrix_budgets(
    contracts: list[Contract],
    amounts: pd.DataFrame,
    locks: pd.DataFrame | None,
    year: int,
    *,
    fill_locks: dict[tuple[str, int], bool] | None = None,
) -> None:
    """Матрица: строки = проекты, столбцы = месяцы. Фиксация: заливка, * в ячейке или отдельный лист фиксации."""
    by_id = {c.id: c for c in contracts}
    id_col = amounts.columns[0]
    month_cols = _matrix_month_columns(amounts)
    if not month_cols:
        return

    lock_month_cols: dict[int, str] = {}
    if locks is not None and not locks.empty:
        lock_month_cols = _matrix_month_columns(locks)

    for _, row in amounts.iterrows():
        cid = str(row[id_col]).strip()
        if not cid or cid not in by_id:
            continue
        lock_row = None
        if locks is not None and not locks.empty:
            match = locks[locks.iloc[:, 0].astype(str).str.strip() == cid]
            lock_row = match.iloc[0] if len(match) else None

        for month, col in month_cols.items():
            val = row[col]
            amount, locked = _parse_matrix_cell_value(val)
            if amount is None:
                continue
            if fill_locks and fill_locks.get((cid, month)):
                locked = True
            if lock_row is not None and month in lock_month_cols:
                lock_val = lock_row[lock_month_cols[month]]
                locked = locked or _bool(lock_val)
            _upsert_month_inflow(by_id[cid], year, month, amount, lock=locked)


def _upsert_month_inflow(
    contract: Contract,
    year: int,
    month: int,
    inflow_amount: float,
    *,
    lock: bool = False,
) -> None:
    for mb in contract.monthly_budgets:
        if mb.month == month:
            mb.inflow_amount = inflow_amount
            mb.lock = lock
            return
    contract.monthly_budgets.append(
        ContractMonthlyBudget(
            contract_id=contract.id,
            year=year,
            month=month,
            inflow_amount=inflow_amount,
            lock=lock,
        )
    )


def _upsert_month_min_balance(
    contract: Contract, year: int, month: int, min_balance: float
) -> None:
    for mb in contract.monthly_budgets:
        if mb.month == month:
            mb.min_balance = min_balance
            return
    contract.monthly_budgets.append(
        ContractMonthlyBudget(
            contract_id=contract.id,
            year=year,
            month=month,
            inflow_amount=0.0,
            min_balance=min_balance,
        )
    )


def _apply_matrix_min_balances(contracts: list[Contract], df: pd.DataFrame, year: int) -> None:
    """Мин. остаток по месяцам: строки = проекты, столбцы = месяцы (как «фот_по_месяцам»)."""
    by_id = {c.id: c for c in contracts}
    id_col = df.columns[0]
    month_cols = _matrix_month_columns(df)
    if not month_cols:
        return

    for _, row in df.iterrows():
        cid = str(row[id_col]).strip()
        if not cid or cid not in by_id:
            continue
        for month, col in month_cols.items():
            val = row[col]
            if val is None or (isinstance(val, float) and pd.isna(val)) or str(val).strip() == "":
                continue
            _upsert_month_min_balance(by_id[cid], year, month, float(val))


def _load_manual_assignments(df: pd.DataFrame) -> list[ManualAssignment]:
    rows: list[ManualAssignment] = []
    for _, r in df.iterrows():
        fixed = r.get("fixed_amount")
        rows.append(
            ManualAssignment(
                employee_id=str(r["employee_id"]).strip(),
                contract_id=str(r["contract_id"]).strip(),
                year=int(r.get("year", date.today().year)),
                month_from=int(r["month_from"]),
                month_to=int(r["month_to"]),
                payment_kind=parse_payment_kind(r.get("payment_kind")),
                fixed_amount=float(fixed) if pd.notna(fixed) else None,
            )
        )
    return rows


def _load_manual_prohibitions(df: pd.DataFrame) -> list[ManualProhibition]:
    rows: list[ManualProhibition] = []
    for _, r in df.iterrows():
        pk = r.get("payment_kind")
        rows.append(
            ManualProhibition(
                employee_id=str(r["employee_id"]).strip(),
                contract_id=str(r["contract_id"]).strip(),
                year=int(r.get("year", date.today().year)),
                month_from=int(r.get("month_from", 1)),
                month_to=int(r.get("month_to", 12)),
                payment_kind=parse_payment_kind(pk) if pd.notna(pk) else None,
            )
        )
    return rows


def _load_plan_overrides(df: pd.DataFrame, default_year: int) -> list[AllocationRecord]:
    """Читает лист «План выплат»: строки с фиксацией считаются ручными для пересчёта."""
    records: list[AllocationRecord] = []
    for _, r in df.iterrows():
        locked = _bool(r.get("lock"), False)
        if not locked:
            continue
        month = _month_from_column(r.get("month"))
        if month is None:
            raise ValueError(
                f"Не удалось прочитать месяц в листе {SHEET_PLAN!r}: {r.get('month')!r}"
            )
        records.append(
            AllocationRecord(
                employee_id=str(r["employee_id"]).strip(),
                contract_id=str(r["contract_id"]).strip(),
                year=int(r.get("year", default_year)),
                month=month,
                payment_kind=parse_payment_kind(r["payment_kind"]),
                amount=float(r["amount"]),
                is_manual=True,
                source="excel_lock",
            )
        )
    return records


def _load_contract_labor(
    df: pd.DataFrame,
    default_year: int,
    position_index: dict[str, PositionReferenceRow],
) -> list[ContractLaborPlan]:
    rows: list[ContractLaborPlan] = []
    for _, r in df.iterrows():
        cid = str(r["contract_id"]).strip()
        if not cid:
            continue
        pm = float(r["person_months"])
        if pm <= 0:
            continue
        position = _clean_optional_text(r.get("position"))
        ref = resolve_position(position, position_index) if position else None
        explicit_group = _clean_optional_text(r.get("equivalence_group"))
        salary_page = (
            _clean_optional_text(r.get("salary_page"))
            if "salary_page" in r.index
            else (ref.salary_page if ref else None)
        )
        salary_group_number = (
            _optional_int(r.get("salary_group_number"))
            if "salary_group_number" in r.index
            else (ref.salary_group_number if ref else None)
        )
        level = (
            _optional_int(r.get("position_level"))
            if "position_level" in r.index
            else (ref.level if ref else None)
        )
        equivalence_group = _reference_group_from_fields(
            explicit_group=explicit_group,
            salary_page=salary_page,
            salary_group_number=salary_group_number,
            level=level,
            default=ref,
            fallback=position,
        )
        avg_raw = r.get("avg_monthly_labor_cost")
        avg_cost = float(avg_raw) if pd.notna(avg_raw) and str(avg_raw).strip() else None
        rows.append(
            ContractLaborPlan(
                contract_id=cid,
                year=int(r.get("year", default_year)),
                person_months=pm,
                position=position,
                equivalence_group=equivalence_group,
                position_level=level,
                avg_monthly_labor_cost=avg_cost,
            )
        )
    return rows


def _load_salary_stability(settings: pd.DataFrame) -> SalaryStabilityRules:
    rules = SalaryStabilityRules()
    if settings.empty:
        return rules
    row = settings.iloc[0]
    if "max_salary_contracts_per_year" in row and pd.notna(row["max_salary_contracts_per_year"]):
        rules.max_contracts_per_year = int(row["max_salary_contracts_per_year"])
    if "goz_labor_tolerance" in row and pd.notna(row["goz_labor_tolerance"]):
        rules.goz_labor_tolerance = float(row["goz_labor_tolerance"])
    return rules


def _load_weights(settings: pd.DataFrame) -> OptimizationWeights:
    w = OptimizationWeights()
    mapping = {
        "weight_salary_switch": "salary_contract_switch",
        "weight_admin_complexity": "admin_complexity",
        "weight_plan_deviation": "plan_deviation",
        "weight_uniform_spend_deviation": "uniform_spend_deviation",
        "weight_labor_deviation": "labor_deviation",
        "weight_order_incentive_use": "order_incentive_use",
    }
    if settings.empty:
        return w
    row = settings.iloc[0]
    for col, attr in mapping.items():
        if col in row and pd.notna(row[col]):
            setattr(w, attr, float(row[col]))

    return w

