"""Оформление пользовательского выходного Excel."""

from __future__ import annotations

from pathlib import Path

from fot_planner.excel.reports.user_report import (
    MONTH_COLS,
    SHEET_CONTRACT_FOT,
    SHEET_DEFICIT_MONTH,
    SHEET_DEFICITS,
    SHEET_ISSUES,
    SHEET_LABOR_CONTROL,
    SHEET_PLAN,
    SHEET_SUMMARY,
)

MATRIX_SHEETS = frozenset(
    {
        SHEET_CONTRACT_FOT,
        SHEET_DEFICIT_MONTH,
    }
)

SHEET_STAFF_DETAIL = "ШР_детализация"

FINAL_SHEET_ORDER = (
    SHEET_SUMMARY,
    SHEET_ISSUES,
    SHEET_PLAN,
    SHEET_CONTRACT_FOT,
    SHEET_LABOR_CONTROL,
    SHEET_STAFF_DETAIL,
    SHEET_DEFICITS,
    SHEET_DEFICIT_MONTH,
)

HIDE_IF_ONLY_MESSAGE = frozenset({SHEET_DEFICITS})

TAB_COLORS = {
    SHEET_SUMMARY: "70AD47",
    SHEET_ISSUES: "FFC000",
    SHEET_CONTRACT_FOT: "9EADCC",
    SHEET_LABOR_CONTROL: "70AD47",
    SHEET_PLAN: "A9D18E",
    SHEET_STAFF_DETAIL: "D9EAD3",
    SHEET_DEFICITS: "FFC7CE",
    SHEET_DEFICIT_MONTH: "FFC7CE",
}

MONEY_HINTS = (
    "сумма",
    "выплат",
    "остаток",
    "поступ",
    "дефицит",
    "план",
    "факт",
    "отклон",
    "потреб",
    "накоп",
    "миним",
    "доступ",
    "значение",
    "итого",
)

PM_HINTS = ("чел", "мес", "труд")


def _is_money_header(header: str) -> bool:
    h = header.lower()
    return any(t in h for t in MONEY_HINTS)


def _is_pm_header(header: str) -> bool:
    h = header.lower()
    return any(t in h for t in PM_HINTS) and "сумма" not in h


def _apply_tab_color(ws) -> None:
    color = TAB_COLORS.get(ws.title)
    if color:
        ws.sheet_properties.tabColor = color


def _sheet_has_only_message(ws) -> bool:
    values: list[str] = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            text = str(value).strip()
            if text:
                values.append(text)
    return ws.max_column <= 1 and len(values) <= 1


def _deficit_month_has_values(ws) -> bool:
    for row in ws.iter_rows(min_row=2):
        label = str(row[0].value or "").lower() if row else ""
        if "дефицит" not in label:
            continue
        for cell in row[1:]:
            if isinstance(cell.value, (int, float)) and abs(cell.value) > 0.005:
                return True
    return False


def _reorder_sheets(wb) -> None:
    ordered = [wb[title] for title in FINAL_SHEET_ORDER if title in wb.sheetnames]
    ordered_titles = {ws.title for ws in ordered}
    ordered.extend(ws for ws in wb.worksheets if ws.title not in ordered_titles)
    wb._sheets = ordered


def finalize_user_workbook(path: Path) -> None:
    """Финальная раскладка вкладок после добавления всех отчётных листов."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    _reorder_sheets(wb)

    for ws in wb.worksheets:
        _apply_tab_color(ws)
        if ws.title in HIDE_IF_ONLY_MESSAGE and _sheet_has_only_message(ws):
            ws.sheet_state = "hidden"
        elif ws.title == SHEET_DEFICIT_MONTH and not _deficit_month_has_values(ws):
            ws.sheet_state = "hidden"
        else:
            ws.sheet_state = "visible"

    if not any(ws.sheet_state == "visible" for ws in wb.worksheets):
        wb.worksheets[0].sheet_state = "visible"

    for idx, ws in enumerate(wb.worksheets):
        if ws.sheet_state == "visible":
            wb.active = idx
            break

    wb.save(path)


def format_user_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    ok_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    warn_fill = PatternFill(fill_type="solid", fgColor="FFEB9C")
    err_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

        if ws.title in MATRIX_SHEETS:
            ws.freeze_panes = "D2" if ws.max_column >= 4 else "B2"

        status_col = None
        level_col = None
        for idx, cell in enumerate(ws[1], start=1):
            val = str(cell.value or "").lower()
            if val == "статус":
                status_col = idx
            if val == "уровень":
                level_col = idx

        for col_idx, column in enumerate(ws.columns, start=1):
            letter = get_column_letter(col_idx)
            header = str(column[0].value or "")
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 32)

            if _is_money_header(header) or header in MONTH_COLS:
                for cell in column[1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"
            elif _is_pm_header(header):
                for cell in column[1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.00"

        if status_col or level_col:
            col = status_col or level_col
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                val = str(cell.value or "").lower()
                if val in ("выполнено", "ок", "да") or val.startswith("выполнено по группе"):
                    cell.fill = ok_fill
                elif val.startswith("отклонение") or val in ("предупреждение", "warning"):
                    cell.fill = warn_fill
                elif val in ("ошибка", "error", "дефицит", "нет") or "ошибка" in val:
                    cell.fill = err_fill

        _apply_tab_color(ws)
        if ws.title == SHEET_LABOR_CONTROL:
            ws.freeze_panes = "B2"

    wb.save(path)
