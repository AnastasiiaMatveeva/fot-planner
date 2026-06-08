"""Оформление входного и результирующего Excel (openpyxl)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fot_planner.excel.constants import SHEET_FOT_MATRIX
from fot_planner.excel.load import _month_from_column

def format_workbook(path: Path, *, result: bool = False) -> None:
    """Оформление входного или результирующего Excel: шапка, ширина, числа."""
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    money_headers = (
        "лимита ФОТ",
        "(план)",
        "сумма",
        "остаток",
        "поступление",
        "потрачено",
        "перенос",
        "резерв",
        "фот",
        "оклад",
        "надбав",
        "план",
        "факт",
        "отклон",
        "накоплен",
        "amount",
        "balance",
        "inflow",
        "spent",
    )

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        for column in ws.columns:
            header = str(column[0].value or "")
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 28)
            if any(token in header.lower() for token in money_headers):
                for cell in column[1:]:
                    cell.number_format = "#,##0"

    if result:
        format_result_workbook_sheets(wb)
    else:
        if SHEET_FOT_MATRIX in wb.sheetnames:
            ws = wb[SHEET_FOT_MATRIX]
            ws.freeze_panes = "B2"
            for col in range(2, ws.max_column + 1):
                if _month_from_column(ws.cell(1, col).value) == 12 and ws.max_row >= 2:
                    ws.cell(2, col).fill = PatternFill(fill_type="solid", fgColor="FFFF00")
                    break

    wb.save(path)


def format_result_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    format_result_workbook_sheets(wb)
    wb.save(path)


def format_result_workbook_sheets(wb) -> None:
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    if "проекты_помесячно" in wb.sheetnames:
        ws = wb["проекты_помесячно"]
        ws.freeze_panes = "D2"
        ws.sheet_properties.tabColor = "70AD47"
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 14
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")

    if "переносы_матрица" in wb.sheetnames:
        ws = wb["переносы_матрица"]
        ws.freeze_panes = "D2"
        ws.sheet_properties.tabColor = "FFC000"
        if ws.max_row > 1 and ws.max_column > 3:
            start = get_column_letter(4)
            end = get_column_letter(ws.max_column)
            ws.conditional_formatting.add(
                f"{start}2:{end}{ws.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="FFFFFF",
                    end_type="max",
                    end_color="F4B183",
                ),
            )


