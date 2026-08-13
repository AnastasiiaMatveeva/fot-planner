# -*- coding: utf-8 -*-
"""Результат fot-planner (xlsx) → компактный JSON для встраивания в прототип."""
import openpyxl, json, sys, io, os

MON = ["Январь","Февраль","Март","Апрель","Май","Июнь","Июль","Август",
       "Сентябрь","Октябрь","Ноябрь","Декабрь"]
MIDX = {m: i for i, m in enumerate(MON)}

def sheet_rows(ws):
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        row = {hdr[c-1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        if any(v not in (None, "") for v in row.values()):
            yield row

def convert(inp_path, res_path, out_path):
    win = openpyxl.load_workbook(inp_path, data_only=True)
    wres = openpyxl.load_workbook(res_path, data_only=True)

    data = {"employees": [], "contracts": [], "inflow": {}, "plan": [],
            "cash": {}, "summary": [], "warnings": [], "limits": []}

    for row in sheet_rows(win["сотрудники"]):
        data["employees"].append({
            "code": row["код строки"], "fio": row["фио"], "pos": row["должность"],
            "dep": row["подразделение"], "rate": row["ставка"], "sal": row["зарплата"],
            "from": str(row["дата начала"])[:10]})

    for row in sheet_rows(win["договоры"]):
        data["contracts"].append({
            "code": row["код"], "name": row["название"], "num": row["номер"],
            "type": row["тип договора"], "goz": row["ГОЗ"], "fot": row["фот"],
            "kinds": {k: row.get(col) for k, col in [
                ("oklad", "оклад разрешен"), ("k120", "120 разрешена"),
                ("k122", "122 разрешена"), ("k124", "124 разрешена"),
                ("k152", "152 разрешена"), ("prk", "стимулирующая приказом разрешена")]}})

    ws = win["фот_по_месяцам"]
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if not code:
            continue
        data["inflow"][code] = [ws.cell(r, c).value or 0 for c in range(2, 14)]

    kind_map = {
        "оклад": "oklad", "120 надбавка (гостайна)": "k120",
        "122 надбавка (качество)": "k122", "124 надбавка (интенсивность)": "k124",
        "152 надбавка (доп. работа)": "k152", "стимулирующая приказом": "prk"}
    for row in sheet_rows(wres["План выплат"]):
        kind_raw = row.get("вид выплаты") or ""
        kind = kind_map.get(kind_raw)
        if kind is None:
            for k, v in kind_map.items():
                if kind_raw.startswith(k.split(" ")[0]):
                    kind = v
                    break
        data["plan"].append({
            "emp": row["код строки"], "m": MIDX.get(row["месяц"], -1),
            "ctr": row["договор"], "kind": kind or kind_raw,
            "sum": round(float(row["сумма"]), 2)})

    for row in sheet_rows(wres["ФОТ по договорам"]):
        if row.get("раздел") != "Поступления и выплаты":
            continue
        code, metric = row["договор"], row["показатель"]
        if metric in ("Поступление", "Выплаты", "Остаток на конец"):
            data["cash"].setdefault(code, {})[metric] = [
                row.get(m) or 0 for m in ["Янв","Фев","Мар","Апр","Май","Июн",
                                          "Июл","Авг","Сен","Окт","Ноя","Дек"]]

    for row in sheet_rows(wres["Итог расчета"]):
        data["summary"].append([row.get("показатель"), row.get("значение"),
                                row.get("статус"), row.get("комментарий")])

    if "Проблемы и предупреждения" in wres.sheetnames:
        for row in sheet_rows(wres["Проблемы и предупреждения"]):
            data["warnings"].append([row.get("уровень"), row.get("раздел"),
                row.get("объект"), row.get("месяц"), row.get("описание проблемы"),
                row.get("отклонение / сумма"), row.get("рекомендация")])

    with io.open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    print("employees:", len(data["employees"]), "| contracts:", len(data["contracts"]),
          "| plan rows:", len(data["plan"]), "| bytes:", os.path.getsize(out_path))

if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2], sys.argv[3])
