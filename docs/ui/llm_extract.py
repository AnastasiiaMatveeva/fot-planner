# -*- coding: utf-8 -*-
"""Разбор документа об обновлении сумм языковой моделью.

Зачем отдельный слой: якорный разбор в serve.py требует, чтобы в документе была
таблица с заголовками «должность», «оклад», «П2556». Реальные выписки из приказов
так не выглядят: заголовок бывает в два ряда, объединен, назван «Наименование
должности» или «Должностной оклад, руб.», часть величин вынесена в текст над
таблицей. Модель читает документ целиком и возвращает плоский список величин
с основанием и датой начала действия.

Модель ничего не записывает: результат идет в дифф, который экономист
подтверждает вручную. Без ключа модуль сообщает, что недоступен, и вызывающая
сторона откатывается на якорный разбор.

Провайдер выбирается по тому, какой ключ задан в окружении:
    ANTHROPIC_API_KEY  -> Claude, разбор со схемой (structured outputs)
    DEEPSEEK_API_KEY   -> DeepSeek, разбор в режиме JSON
Схему в обоих случаях проверяем у себя: модель может вернуть что угодно.
"""
from __future__ import annotations

import json
import os
import ssl
import urllib.error
import urllib.request

FIELDS = ("оклад", "П2556", "П4")
MAX_CHARS = 60000          # хватает на выписку в несколько листов
DEEPSEEK_URL = "https://api.deepseek.com/chat/completions"


def _ssl_context():
    """Связка корневых сертификатов для запроса.

    В хранилище Windows на рабочих машинах набор корней урезан: у api.deepseek.com
    сертификат Amazon, которого там нет, и проверка падает с
    CERTIFICATE_VERIFY_FAILED. Берем полную связку certifi, если она установлена,
    иначе остаемся на системной.
    """
    try:
        import certifi
        return ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        return ssl.create_default_context()


def provider():
    """(имя провайдера, модель, причина недоступности)."""
    forced = os.environ.get("FOT_LLM_PROVIDER", "").strip().lower()

    if forced != "deepseek" and os.environ.get("ANTHROPIC_API_KEY"):
        try:
            import anthropic  # noqa: F401
        except ImportError:
            return None, None, "задан ANTHROPIC_API_KEY, но не установлен пакет anthropic"
        return "anthropic", os.environ.get("FOT_LLM_MODEL", "claude-opus-5"), ""

    if os.environ.get("DEEPSEEK_API_KEY"):
        return "deepseek", os.environ.get("FOT_LLM_MODEL", "deepseek-chat"), ""

    return None, None, "не задан ни ANTHROPIC_API_KEY, ни DEEPSEEK_API_KEY"


def available():
    name, _model, why = provider()
    return (name is not None), why


# ── документ в текст ────────────────────────────────────────────
def workbook_to_text(path):
    """Все листы книги в виде текстовой сетки с координатами ячеек.

    Координаты нужны, чтобы объединенные и многоэтажные заголовки не теряли
    привязку к своим колонкам, а модель могла сослаться на место в документе.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("### Лист: %s (строк %d, колонок %d)" % (ws.title, ws.max_row, ws.max_column))
        for r in range(1, min(ws.max_row, 400) + 1):
            cells = []
            for c in range(1, min(ws.max_column, 40) + 1):
                v = ws.cell(r, c).value
                if v is None or str(v).strip() == "":
                    continue
                cells.append("%s%d=%s" % (chr(64 + c) if c <= 26 else "?", r, v))
            if cells:
                out.append(" | ".join(cells))
        if ws.merged_cells.ranges:
            out.append("объединенные диапазоны: " +
                       ", ".join(str(x) for x in list(ws.merged_cells.ranges)[:40]))
    return "\n".join(out)[:MAX_CHARS]


SYSTEM = """Ты разбираешь документ российской организации об оплате труда: выписку
из приказа, письмо военного представительства, положение об оплате труда или
приложение к ним. Нужно извлечь величины, относящиеся к должностям.

Извлекай только три вида величин:
- «оклад» — должностной оклад за одну полную ставку в месяц, рублей;
- «П2556» — предельный размер оклада по приказу № 2556, рублей;
- «П4» — предельный размер по приказу № 4, обычно задан категории персонала
  (НТП, НР), а не должности.

Правила:
1. Должность записывай так, как она названа в документе, без номера строки.
2. Значение — число в рублях. Если в документе тысячи рублей, приведи к рублям.
3. Не выдумывай величин, которых в документе нет: пустая ячейка, прочерк и
   «не установлен» означают отсутствие значения — такую строку пропусти.
4. Заголовки бывают в несколько строк и объединенными: соотнеси значение с
   правильной колонкой по координатам ячеек.
5. Дату начала действия и основание (номер и дата приказа или письма) возьми
   из шапки документа, если они там есть.

Если документ не об оплате труда должностей, верни пустой список rows и объясни
это в notes."""

JSON_SHAPE = """Ответь одним объектом JSON без пояснений вокруг:
{
  "document_kind": "что это за документ, одной фразой",
  "basis": "номер и дата приказа или письма, либо null",
  "effective_from": "ДД.ММ.ГГГГ, либо null",
  "rows": [{"position": "должность как в документе",
            "field": "оклад | П2556 | П4",
            "value": 12345,
            "cell": "лист и ячейка, откуда взято, либо null"}],
  "notes": "что осталось непонятным, либо null"
}"""


# ── проверка ответа ─────────────────────────────────────────────
def _clean(raw):
    """Отбросить строки, не подходящие под схему: модель могла ошибиться."""
    rows, dropped = [], 0
    for r in raw.get("rows") or []:
        if not isinstance(r, dict):
            dropped += 1
            continue
        pos = str(r.get("position") or "").strip()
        field = str(r.get("field") or "").strip()
        try:
            value = float(str(r.get("value")).replace(" ", "").replace(",", "."))
        except (TypeError, ValueError):
            dropped += 1
            continue
        if not pos or field not in FIELDS or value <= 0:
            dropped += 1
            continue
        rows.append({"pos": pos, "field": field, "value": value,
                     "cell": r.get("cell") or None})
    return rows, dropped


def _result(raw, model):
    rows, dropped = _clean(raw)
    notes = raw.get("notes") or ""
    if dropped:
        notes = (notes + " " if notes else "") + \
                "Отброшено строк, не подошедших под схему: %d." % dropped
    return {"ok": True, "model": model,
            "kind": raw.get("document_kind") or "",
            "basis": raw.get("basis") or None,
            "effective_from": raw.get("effective_from") or None,
            "notes": notes.strip() or None,
            "rows": rows}


# ── провайдеры ──────────────────────────────────────────────────
def _call_anthropic(text, filename, model):
    from typing import Literal, Optional

    import anthropic
    from pydantic import BaseModel, Field

    class Row(BaseModel):
        position: str = Field(description="должность или категория персонала, как в документе")
        field: Literal["оклад", "П2556", "П4"]
        value: float = Field(description="значение в рублях")
        cell: Optional[str] = Field(default=None, description="лист и ячейка, откуда взято")

    class Update(BaseModel):
        document_kind: str
        basis: Optional[str] = None
        effective_from: Optional[str] = None
        rows: list[Row]
        notes: Optional[str] = None

    client = anthropic.Anthropic()
    msg = client.messages.parse(
        model=model, max_tokens=8000, system=SYSTEM,
        messages=[{"role": "user", "content": "Документ «%s»:\n\n%s" % (filename, text)}],
        output_format=Update,
    )
    if msg.parsed_output is None:
        return None
    return msg.parsed_output.model_dump()


def _call_deepseek(text, filename, model):
    body = json.dumps({
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM + "\n\n" + JSON_SHAPE},
            {"role": "user", "content": "Документ «%s»:\n\n%s" % (filename, text)},
        ],
        "response_format": {"type": "json_object"},
        "temperature": 0,
        "max_tokens": 8000,
    }, ensure_ascii=False).encode("utf-8")

    req = urllib.request.Request(
        DEEPSEEK_URL, data=body, method="POST",
        headers={"Content-Type": "application/json",
                 "Authorization": "Bearer " + os.environ["DEEPSEEK_API_KEY"]})
    with urllib.request.urlopen(req, timeout=180, context=_ssl_context()) as resp:
        payload = json.loads(resp.read().decode("utf-8"))
    return json.loads(payload["choices"][0]["message"]["content"])


def parse(path, filename=""):
    """Разобрать документ. Всегда возвращает dict; ok=False — откат на якорный разбор."""
    name, model, why = provider()
    if name is None:
        return {"ok": False, "unavailable": True, "error": why}

    try:
        text = workbook_to_text(path)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": "Файл не прочитан: %s" % e}
    if not text.strip():
        return {"ok": False, "error": "Документ пуст"}

    filename = filename or os.path.basename(path)
    try:
        if name == "anthropic":
            raw = _call_anthropic(text, filename, model)
        else:
            raw = _call_deepseek(text, filename, model)
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", "replace")[:300]
        return {"ok": False, "error": "%s ответил %s: %s" % (name, e.code, detail)}
    except Exception as e:  # noqa: BLE001 — сеть, ключ, лимиты, битый JSON
        return {"ok": False, "error": "%s: %s" % (name, e)}
    if not raw:
        return {"ok": False, "error": "Модель вернула пустой разбор"}

    return _result(raw, "%s %s" % (name, model))
