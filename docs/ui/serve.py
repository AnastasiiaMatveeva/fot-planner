# -*- coding: utf-8 -*-
"""Локальный сервер демонстрации.

Возможности:
  GET  /prototype.html      — прототип
  POST /api/upload          — приём xlsx/xlsm, якорный разбор, паспорт и журнал
  POST /api/chat            — правка паспорта командой из чата
  POST /api/solve           — запуск настоящего fot-planner на текущем паспорте

Запуск из корня репозитория:
    .venv\\Scripts\\python.exe docs\\ui\\serve.py
Затем открыть http://127.0.0.1:8760/prototype.html

Без сервера прототип тоже работает: используется встроенный результат расчёта.
"""
import http.server
import json
import os
import subprocess
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
EXE = os.path.join(ROOT, ".venv", "Scripts", "fot-planner.exe")
TEMPLATE = os.path.join(HERE, "demo_input.xlsx")
UPLOAD = os.path.join(HERE, "_uploaded.xlsx")
BUILT = os.path.join(HERE, "_built_input.xlsx")
RES = os.path.join(HERE, "_result.xlsx")
PORT = 8760

sys.path.insert(0, HERE)
import extract          # noqa: E402
import llm_extract      # noqa: E402
import result2json      # noqa: E402

# Корень проекта — привычное место для .env; docs/ui — рядом с сервером, который
# его читает. Имя без точки Проводник создает охотнее, поэтому принимаем оба.
ENV_PATHS = [os.path.join(d, n) for d in (ROOT, HERE) for n in (".env", "env")]
_ENV_SOURCE = None


def _load_env_file():
    """Ключи модели из .env. Все варианты пути в .gitignore, в репозиторий не попадают.

    Формат: KEY=значение, по строке на ключ. Уже заданные переменные окружения
    имеют приоритет и не перетираются.
    """
    global _ENV_SOURCE
    path = next((p for p in ENV_PATHS if os.path.isfile(p)), None)
    if path is None:
        return []
    _ENV_SOURCE = os.path.relpath(path, ROOT)
    loaded = []
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and v and not os.environ.get(k):
                os.environ[k] = v
                loaded.append(k)
    return loaded


_ENV_LOADED = _load_env_file()

STATE = {"passport": None, "source": None}


# ── справочник должностей: хранится листом входного файла ────────
# Лист «лимиты_по_должностям» читает сам решатель (excel/load.py), поэтому
# правка из интерфейса попадает в расчет без промежуточных копий.
REFUP = os.path.join(HERE, "_reference_update.xlsx")
REFSHEET = "лимиты_по_должностям"
_num = extract._num
REF_EDITABLE = {"оклад": 6, "П2556": 7, "П4": 8}


def _ref_ws(write=False):
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    return wb, wb[REFSHEET]


def read_reference():
    """Справочник должностей из листа входного файла."""
    _wb, ws = _ref_ws()
    rows = []
    for r in range(2, ws.max_row + 1):
        pos = ws.cell(r, 1).value
        if not pos:
            continue
        rows.append({
            "pos": str(pos), "cat": ws.cell(r, 2).value or "",
            "sal": _num(ws.cell(r, 6).value), "p2556": _num(ws.cell(r, 7).value),
            "p4": _num(ws.cell(r, 8).value), "note": ws.cell(r, 10).value or "",
        })
    return {"ok": True, "rows": rows, "file": os.path.basename(TEMPLATE),
            "sheet": REFSHEET}


def write_reference(edits):
    """Записать правки сумм в лист. edits: [{pos, field, value}]."""
    if not edits:
        return {"ok": True, "applied": []}
    wb, ws = _ref_ws(write=True)
    at = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v:
            at[str(v).strip().lower()] = r
    applied, skipped = [], []
    for e in edits:
        r = at.get(str(e.get("pos", "")).strip().lower())
        col = REF_EDITABLE.get(e.get("field"))
        if not r or not col:
            skipped.append(e.get("pos"))
            continue
        new = _num(e.get("value"))
        old = _num(ws.cell(r, col).value)
        if old == new:
            continue
        ws.cell(r, col).value = new
        applied.append({"pos": ws.cell(r, 1).value, "field": e["field"],
                        "old": old, "new": new})
    if applied:
        wb.save(TEMPLATE)
        BASE_SUM["cache"] = None
    return {"ok": True, "applied": applied, "skipped": skipped}


def _changes_from_rows(rows):
    """Список величин из документа → расхождения с действующим справочником."""
    sys.path.insert(0, os.path.join(ROOT, "src"))
    from fot_planner.position_reference import normalize_position
    idx = _position_index(read_reference()["rows"])
    key = {"оклад": "sal", "П2556": "p2556", "П4": "p4"}
    changes, unknown, seen = [], [], set()
    for row in rows:
        pos = str(row.get("pos") or "").strip()
        f = key.get(row.get("field"))
        new = row.get("value")
        if not pos or not f or new is None:
            continue
        hit = idx.get(normalize_position(pos))
        if not hit:
            if pos not in unknown:
                unknown.append(pos)
            continue
        seen.add(row["field"])
        old = hit[f]
        if old != float(new):
            changes.append({"pos": hit["pos"], "field": row["field"],
                            "old": old, "new": float(new)})
    return changes, unknown, seen


_POS_INDEX = {"built_for": None, "map": {}}


def _position_index(rows):
    """Разрешение названия должности: нормализация плюс синонимы сервиса.

    Модель возвращает должность так, как она написана в документе: «Вед. инженер»,
    «Инженер I категории», «МНС». Сервис уже знает эти написания — берем его
    справочник синонимов, чтобы не заводить второй.
    """
    names = tuple(r["pos"] for r in rows)
    if _POS_INDEX["built_for"] == names:
        return _POS_INDEX["map"]
    sys.path.insert(0, os.path.join(ROOT, "src"))
    from fot_planner.position_reference import (
        default_position_synonyms,
        normalize_position,
    )
    idx = {normalize_position(r["pos"]): r for r in rows}
    for raw, canonical in default_position_synonyms().items():
        target = idx.get(normalize_position(canonical))
        if target is not None:
            idx.setdefault(normalize_position(raw), target)
    _POS_INDEX["built_for"] = names
    _POS_INDEX["map"] = idx
    return idx


def diff_reference(data, filename, use_llm=True):
    """Сверить присланный документ с действующим справочником.

    Модель читает документ произвольного вида; якорный разбор работает, только
    когда в книге есть таблица с ожидаемыми заголовками. Пробуем модель, при
    ее недоступности или пустом ответе откатываемся на якорный.
    """
    with open(REFUP, "wb") as f:
        f.write(data)

    if use_llm:
        r = llm_extract.parse(REFUP, filename)
        if r.get("ok") and r.get("rows"):
            changes, unknown, seen = _changes_from_rows(r["rows"])
            return {"ok": True, "file": filename, "sheet": r.get("kind") or "",
                    "by": "модель " + r["model"], "basis": r.get("basis"),
                    "effective_from": r.get("effective_from"), "notes": r.get("notes"),
                    "changes": changes, "unknown": unknown[:20],
                    "fields": [f for f in ("оклад", "П2556", "П4") if f in seen],
                    "scanned": len(r["rows"])}
        llm_note = r.get("error") if not r.get("ok") else "модель не нашла величин"
    else:
        llm_note = None

    out = _diff_by_anchors(filename)
    if out.get("ok"):
        out["by"] = "разбор по заголовкам"
        if llm_note:
            out["fallback"] = llm_note
    elif llm_note:
        out["fallback"] = llm_note
    return out


def _diff_by_anchors(filename):
    """Запасной разбор: ищем таблицу по заголовкам «должность», «оклад», «П2556», «П4»."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(REFUP, data_only=True)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": "Файл не прочитан: %s" % e}

    want = {"оклад": "sal", "п2556": "p2556", "п4": "p4"}
    best = None
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 12) + 1):
            hdr = {}
            for c in range(1, ws.max_column + 1):
                t = str(ws.cell(r, c).value or "").strip().lower()
                if t.startswith("должност"):
                    hdr["pos"] = c
                for k, f in want.items():
                    if t.replace(" ", "").startswith(k):
                        hdr[f] = c
            if "pos" in hdr and len(hdr) > 1:
                best = (ws, r, hdr)
                break
        if best:
            break
    if not best:
        return {"ok": False,
                "error": "Не нашел таблицу: нужна колонка «должность» и хотя бы одна из «оклад», «П2556», «П4»"}

    ws, hrow, hdr = best
    sys.path.insert(0, os.path.join(ROOT, "src"))
    from fot_planner.position_reference import normalize_position
    idx = _position_index(read_reference()["rows"])
    changes, unknown, seen = [], [], set()
    for r in range(hrow + 1, ws.max_row + 1):
        pos = ws.cell(r, hdr["pos"]).value
        if not pos:
            continue
        hit = idx.get(normalize_position(pos))
        if hit is None:
            unknown.append(str(pos))
            continue
        for field, col in hdr.items():
            if field == "pos":
                continue
            new = _num(ws.cell(r, col).value)
            if new is None:
                continue
            seen.add(field)   # величина считается найденной, только если есть значения
            old = hit[field]
            if old != new:
                changes.append({"pos": hit["pos"], "field": field,
                                "old": old, "new": new})
    label = {"sal": "оклад", "p2556": "П2556", "p4": "П4"}
    for c in changes:
        c["field"] = label[c["field"]]
    return {"ok": True, "file": filename, "sheet": ws.title,
            "changes": changes, "unknown": unknown[:20],
            "fields": [label[f] for f in hdr if f in seen],
            "scanned": ws.max_row - hrow}


def do_upload(data, filename):
    with open(UPLOAD, "wb") as f:
        f.write(data)
    try:
        out = extract.extract(UPLOAD)
    except Exception as e:  # noqa: BLE001 — сообщение вместо падения
        return {"ok": False, "error": "Файл не прочитан: %s" % e}
    STATE["passport"] = out["passport"]
    STATE["source"] = filename
    return {"ok": True, "file": filename, "log": out["log"],
            "questions": out["questions"], "ready": out["ready"],
            "passport": out["passport"]}


def do_chat(text):
    if not STATE["passport"]:
        return {"ok": False, "reply": "Сначала загрузите документ"}
    reply, changed, run = extract.apply_chat(STATE["passport"], text)
    return {"ok": True, "reply": reply, "changed": changed, "solve": run,
            "passport": STATE["passport"]}


SETTING_KEYS = {
    "Год планирования": "год",
    "Разрешить дефицит": "разрешить дефицит",
    "Максимум договоров оклада в год": "макс договоров оклада в год",
    "Штраф смены договора оклада": "штраф смены договора оклада",
    "Штраф административной сложности выплат": "штраф административной сложности выплат",
    "Штраф неравномерного освоения": "штраф отклонения от равномерного освоения",
    "Штраф использования приказа": "штраф использования приказа",
    "Допуск трудоёмкости": "допуск трудоёмкости",
}


def apply_settings(path, settings):
    """Записать параметры целевой функции в лист «настройки» входного файла."""
    if not settings:
        return []
    import openpyxl
    wb = openpyxl.load_workbook(path)
    if "настройки" not in wb.sheetnames:
        return []
    ws = wb["настройки"]
    hdr = {str(ws.cell(1, c).value or "").strip(): c
           for c in range(1, ws.max_column + 1)}
    applied = []
    for label, raw in settings.items():
        col_name = SETTING_KEYS.get(label)
        if not col_name or col_name not in hdr:
            continue
        txt = str(raw).replace(" ", " ").replace(" ", "").replace(",", ".")
        try:
            val = float(txt)
            val = int(val) if val.is_integer() else val
        except ValueError:
            val = raw
        cur = ws.cell(2, hdr[col_name]).value
        if str(cur) != str(val):
            ws.cell(2, hdr[col_name]).value = val
            applied.append("%s: %s -> %s" % (col_name, cur, val))
    if applied:
        wb.save(path)
    return applied


def do_solve(settings=None):
    src = TEMPLATE
    if STATE["passport"] and STATE["passport"].get("employees"):
        try:
            src = extract.passport_to_input(STATE["passport"], TEMPLATE, BUILT)
        except Exception as e:  # noqa: BLE001
            return {"ok": False, "error": "Не собрал вход из паспорта: %s" % e}
    if src == TEMPLATE and settings:
        import shutil
        shutil.copy(TEMPLATE, BUILT)
        src = BUILT
    applied = apply_settings(src, settings)
    t0 = time.time()
    p = subprocess.run([EXE, "solve", "-i", src, "-o", RES],
                       capture_output=True, text=True, timeout=900, cwd=ROOT)
    dt = round(time.time() - t0, 1)
    if p.returncode != 0:
        return {"ok": False, "sec": dt,
                "error": (p.stderr or p.stdout or "").strip()[-800:]}
    out = os.path.join(HERE, "_result.json")
    result2json.convert(src, RES, out)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    return {"ok": True, "sec": dt, "data": data,
            "source": os.path.basename(src), "settings_applied": applied}


def parse_multipart(raw, ctype):
    """Минимальный разбор multipart/form-data: первая часть с именем файла."""
    if 'boundary=' not in ctype:
        return None, None
    boundary = ctype.split('boundary=', 1)[1].strip().strip(chr(34))
    sep = ('--' + boundary).encode('utf-8')
    crlf = bytes([13, 10])
    crlf2 = crlf + crlf
    for part in raw.split(sep):
        head_end = part.find(crlf2)
        if head_end < 0:
            continue
        head = part[:head_end].decode('utf-8', 'replace')
        if 'filename=' not in head:
            continue
        fname = head.split('filename=', 1)[1]
        fname = fname.split(chr(13), 1)[0].split(chr(10), 1)[0].strip().strip(chr(34))
        body = part[head_end + len(crlf2):]
        if body.endswith(crlf):
            body = body[:-2]
        return (fname or 'документ.xlsx'), body
    return None, None


BASE_SUM = {"cache": None}


def _summary_of(res_path):
    import openpyxl
    out = {}
    wb = openpyxl.load_workbook(res_path, read_only=True, data_only=True)
    ws = wb["Итог расчета"]
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k:
            out[str(k)] = ws.cell(r, 2).value
    rest = 0.0
    if "ФОТ по договорам" in wb.sheetnames:
        f = wb["ФОТ по договорам"]
        hdr = [f.cell(1, c).value for c in range(1, f.max_column + 1)]
        dec = hdr.index("Дек") + 1 if "Дек" in hdr else None
        for r in range(2, f.max_row + 1):
            if (f.cell(r, 1).value == "Поступления и выплаты"
                    and f.cell(r, 4).value == "Остаток на конец" and dec):
                rest += float(f.cell(r, dec).value or 0)
    out["_rest"] = round(rest, 2)
    return out


def _base_summary():
    """Показатели базового плана: считаются один раз и кешируются."""
    if BASE_SUM["cache"] is None:
        base_res = os.path.join(HERE, "_base_result.xlsx")
        if not os.path.exists(base_res):
            subprocess.run([EXE, "solve", "-i", TEMPLATE, "-o", base_res],
                           capture_output=True, text=True, timeout=900, cwd=ROOT)
        BASE_SUM["cache"] = _summary_of(base_res) if os.path.exists(base_res) else {}
    return BASE_SUM["cache"]


def _infeasible_reason(inp_path):
    """Сопоставить потребность и фонды: назвать вероятную причину отказа."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(inp_path, data_only=True)
        emp, ctr = wb["сотрудники"], wb["договоры"]
        ehdr = {str(emp.cell(1, c).value or "").strip().lower(): c
                for c in range(1, emp.max_column + 1)}
        chdr = {str(ctr.cell(1, c).value or "").strip().lower(): c
                for c in range(1, ctr.max_column + 1)}

        def months(cell_from, cell_to):
            def mnum(v, default):
                t = str(v or "")
                if len(t) >= 10 and t[2] == "." and t[6:10] == "2026":
                    return int(t[3:5])
                if t[:4] == "2026":
                    return int(t[5:7])
                return default
            a = mnum(cell_from, 1)
            b = mnum(cell_to, 12)
            return max(0, b - a + 1)

        need = 0.0
        for r in range(2, emp.max_row + 1):
            if not emp.cell(r, 1).value:
                continue
            sal = float(emp.cell(r, ehdr["зарплата"]).value or 0)
            need += sal * months(emp.cell(r, ehdr.get("дата начала", 9)).value,
                                 emp.cell(r, ehdr.get("дата окончания", 10)).value)
        fot = 0.0
        for r in range(2, ctr.max_row + 1):
            if ctr.cell(r, 1).value:
                fot += float(ctr.cell(r, chdr["фот"]).value or 0)
    except Exception:  # noqa: BLE001 — диагностика не должна ронять ответ
        return ("Решатель установил, что при этих условиях допустимого решения "
                "не существует даже с разрешённым дефицитом.")

    def money(v):
        return "{:,.0f}".format(float(v or 0)).replace(",", " ")

    base_txt = ("Решатель установил, что при этих условиях допустимого решения "
                "не существует даже с разрешённым дефицитом. ")
    if fot > need * 1.02:
        return (base_txt + "Годовая потребность по штатному расписанию — " + money(need) +
                " ₽, фонды договоров — " + money(fot) + " ₽. Избыток " +
                money(fot - need) + " ₽: модель в текущей редакции не всегда находит "
                "решение, когда фонд заметно превышает потребность. Для такого сценария "
                "фонды нужно пересобрать под новую потребность либо разрешить "
                "неосвоенный остаток явно.")
    if need > fot * 1.02:
        return (base_txt + "Потребность " + money(need) + " ₽ превышает фонды " +
                money(fot) + " ₽ на " + money(need - fot) +
                " ₽ — денег договоров не хватает даже при полном освоении.")
    return (base_txt + "Потребность и фонды сопоставимы (" + money(need) + " и " +
            money(fot) + " ₽): причина в ограничениях по видам выплат, "
            "предельным размерам или кассе. Разбор — в разделе «Анализ невыполнимости».")


def do_scenario(text):
    """Разобрать требование, изменить вход, посчитать, сравнить с базой."""
    passport = STATE["passport"]
    if not passport:
        out = extract.extract(TEMPLATE)
        passport = out["passport"]
    parsed = extract.parse_scenario(text, passport)
    if parsed.get("error"):
        return {"ok": False, "error": parsed["error"], "hint": parsed.get("hint", "")}

    inp = os.path.join(HERE, "_scenario.xlsx")
    res = os.path.join(HERE, "_scenario_result.xlsx")
    extract.apply_scenario(TEMPLATE, parsed["ops"], inp)
    t0 = time.time()
    p = subprocess.run([EXE, "solve", "-i", inp, "-o", res],
                       capture_output=True, text=True, timeout=900, cwd=ROOT)
    dt = round(time.time() - t0, 1)

    base = _base_summary()
    if p.returncode != 0 or not os.path.exists(res):
        return {"ok": True, "sec": dt, "status": "INFEASIBLE",
                "title": parsed["title"], "what": parsed["what"],
                "rows": [["Статус решателя", "OPTIMAL", "INFEASIBLE",
                          "задача невыполнима", ""]],
                "note": _infeasible_reason(inp)}

    sm = _summary_of(res)

    def money(v):
        return "{:,.0f}".format(float(v or 0)).replace(",", " ")

    def row(name, key, is_money=True):
        b = float(base.get(key) or 0)
        n = float(sm.get(key) or 0)
        d = n - b
        if is_money:
            bt, nt = money(b) + " ₽", money(n) + " ₽"
            dt_ = "0" if abs(d) < 1 else (("+" if d > 0 else "−") + money(abs(d)) + " ₽")
        else:
            bt, nt = str(int(b)), str(int(n))
            dt_ = "0" if d == 0 else (("+" if d > 0 else "−") + str(int(abs(d))))
        return [name, bt, nt, dt_, ""]

    rows = [
        row("Выплачено за год", "Всего выплат"),
        row("Общий дефицит", "Общий дефицит"),
        row("Сотрудников с дефицитом", "Сотрудников с дефицитом", False),
        row("Остаток на конец года", "_rest"),
        row("Смен договора оклада", "Смен договора оклада", False),
    ]
    deficit = float(sm.get("Общий дефицит") or 0)
    demp = int(sm.get("Сотрудников с дефицитом") or 0)
    if deficit > 1:
        note = ("Решатель вернул частичный план: недоплата " + money(deficit) +
                " ₽ распределена по " + str(demp) +
                " сотрудникам и показана помесячно в отчёте «Дефициты».")
    else:
        note = ("Задача решена полностью: дефицита нет, все выплаты укладываются "
                "в ограничения договоров и предельные размеры.")
    return {"ok": True, "sec": dt, "status": str(sm.get("Статус решателя", "")),
            "title": parsed["title"], "what": parsed["what"],
            "rows": rows, "note": note}


WORK_FILES = ["_uploaded.xlsx", "_built_input.xlsx", "_result.xlsx", "_result.json",
              "_scenario.xlsx", "_scenario_result.xlsx", "_base_result.xlsx",
              "_reference_update.xlsx"]


def do_reset():
    """Сброс демонстрации: паспорт, кэш базового расчёта и рабочие файлы."""
    STATE["passport"] = None
    STATE["source"] = None
    BASE_SUM["cache"] = None
    removed = 0
    for name in WORK_FILES:
        path = os.path.join(HERE, name)
        if os.path.exists(path):
            try:
                os.remove(path)
                removed += 1
            except OSError:
                pass
    return {"ok": True, "removed": removed}


class H(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=HERE, **kw)

    def _send(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def end_headers(self):
        # демонстрация правится по ходу показа — страница не должна кешироваться
        self.send_header("Cache-Control", "no-store, must-revalidate")
        super().end_headers()

    def do_GET(self):
        if self.path == "/api/ping":
            self._send({"ok": True})
            return
        if self.path == "/api/reference":
            self._send(read_reference())
            return
        super().do_GET()

    def do_POST(self):
        try:
            if self.path == "/api/reset":
                self._send(do_reset())
                return
            if self.path == "/api/reference":
                n = int(self.headers.get("Content-Length") or 0)
                payload = json.loads(self.rfile.read(n) or b"{}") if n else {}
                self._send(write_reference(payload.get("edits") or []))
                return
            if self.path == "/api/reference/diff":
                ctype = self.headers.get("Content-Type", "")
                if "multipart/form-data" not in ctype:
                    self._send({"ok": False, "error": "Ожидался multipart"}, 400)
                    return
                n = int(self.headers.get("Content-Length") or 0)
                name, data = parse_multipart(self.rfile.read(n), ctype)
                if data is None:
                    self._send({"ok": False, "error": "Файл не передан"}, 400)
                    return
                self._send(diff_reference(data, name))
                return
            if self.path == "/api/solve":
                n = int(self.headers.get("Content-Length") or 0)
                payload = json.loads(self.rfile.read(n) or b"{}") if n else {}
                self._send(do_solve(payload.get("settings")))
            elif self.path == "/api/scenario":
                n = int(self.headers.get("Content-Length") or 0)
                payload = json.loads(self.rfile.read(n) or b"{}")
                self._send(do_scenario(payload.get("text", "")))
            elif self.path == "/api/chat":
                n = int(self.headers.get("Content-Length") or 0)
                payload = json.loads(self.rfile.read(n) or b"{}")
                self._send(do_chat(payload.get("text", "")))
            elif self.path == "/api/upload":
                ctype = self.headers.get("Content-Type", "")
                if "multipart/form-data" not in ctype:
                    self._send({"ok": False, "error": "Ожидался multipart"}, 400)
                    return
                n = int(self.headers.get("Content-Length") or 0)
                raw = self.rfile.read(n)
                name, data = parse_multipart(raw, ctype)
                if data is None:
                    self._send({"ok": False, "error": "Файл не передан"}, 400)
                    return
                self._send(do_upload(data, name))
            else:
                self.send_error(404)
        except Exception as e:  # noqa: BLE001 — ответ вместо падения сервера
            self._send({"ok": False, "error": str(e)}, 500)

    def log_message(self, fmt, *args):
        sys.stderr.write("%s %s\n" % (self.log_date_time_string(), fmt % args))


if __name__ == "__main__":
    if not os.path.exists(EXE):
        print("Не найден", EXE, "— установите пакет: pip install -e .")
        sys.exit(1)
    print("Прототип: http://127.0.0.1:%d/prototype.html" % PORT)
    print("Загрузка документов и чат правок активны")
    if _ENV_LOADED:
        print("Ключи из %s: %s" % (_ENV_SOURCE, ", ".join(_ENV_LOADED)))
    _p, _m, _why = llm_extract.provider()
    print("Разбор нормативных документов:",
          ("модель %s %s" % (_p, _m)) if _p else ("по заголовкам — %s" % _why))
    http.server.ThreadingHTTPServer(("127.0.0.1", PORT), H).serve_forever()
