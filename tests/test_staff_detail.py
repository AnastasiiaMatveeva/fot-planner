from datetime import date

from openpyxl import Workbook, load_workbook

from fot_planner.excel.reports.staff_detail import (
    SHEET_STAFF_DETAIL,
    STAFF_DETAIL_HEADERS,
    build_staff_detail_rows,
    write_staff_detail_sheet,
)
from fot_planner.models import (
    AllocationRecord,
    Contract,
    Employee,
    OpenRateAttribution,
    PaymentKind,
    PlanningContext,
    PlanningResult,
)


def _staff_scenario():
    year = 2026
    employee = Employee(
        id="E001",
        full_name="Мышенский Дмитрий Игоревич",
        position="Инженер",
        department="",
        rate=1.0,
        monthly_wage=66_500,
        reference_salary_for_rate=40_400,
    )
    contracts = [
        Contract(
            id="MAIN",
            name="Основное место",
            number="1",
            contract_type="grant",
            start_date=date(2025, 1, 1),
            end_date=date(2027, 12, 31),
            total_fot=1,
        ),
        Contract(
            id="PART",
            name="Совместительство",
            number="2",
            contract_type="grant",
            start_date=date(2026, 4, 1),
            end_date=date(2026, 8, 31),
            total_fot=1,
        ),
    ]

    allocations = []
    rates = []
    for month in range(1, 13):
        allocations.append(AllocationRecord("E001", "MAIN", year, month, PaymentKind.SALARY, 40_400))
        rates.append(OpenRateAttribution("E001", "MAIN", year, month, 1.0, True))
    for month in range(4, 7):
        allocations.append(AllocationRecord("E001", "MAIN", year, month, PaymentKind.K122, 16_000))
    for month in range(4, 9):
        allocations.append(AllocationRecord("E001", "PART", year, month, PaymentKind.SALARY, 10_100))
        rates.append(
            OpenRateAttribution(
                "E001",
                "PART",
                year,
                month,
                0.25,
                False,
                position="Программист",
                equivalence_group="инженерно-технические специалисты",
            )
        )

    ctx = PlanningContext(year=year, employees=[employee], contracts=contracts)
    result = PlanningResult(
        year=year,
        allocations=allocations,
        deficits=[],
        conflicts=[],
        contract_balances=[],
        solver_status="OPTIMAL",
        objective_value=0,
        solve_time_sec=0,
        open_rate_attributions=rates,
    )
    return ctx, result


def test_staff_detail_matches_hr_structure(tmp_path):
    ctx, result = _staff_scenario()
    rows = build_staff_detail_rows(ctx, result)

    assert len(rows) == 3
    assert [row["Название параметра"] for row in rows] == [
        "Оклад (должностной оклад)",
        "За качество выполняемых работ",
        "Оклад (должностной оклад)",
    ]
    assert [row["Код \nпар-ра"] for row in rows] == ["1", "122", "1"]
    assert all(row["Категория\nперсонала"] == "НТП" for row in rows)
    assert rows[0]["Номинальное\n значение \nпараметра"] == 40_400
    assert rows[0]["Значение\n параметра \nпо ставке"] == 40_400
    assert rows[1]["Значение\n параметра \nпо ставке"] == 16_000
    assert rows[0]["Должность"] == "Инженер"
    assert rows[2]["Должность"] == "Программист"
    assert rows[2]["Ставка"] == 0.25
    assert rows[2]["Значение\n параметра \nпо ставке"] == 10_100
    assert rows[2]["Сумма \nв\nруб."] == 10_100
    assert rows[1]["Начало \nдействия"] == date(2026, 4, 1)
    assert rows[1]["Окончание \nдействия"] == date(2026, 6, 30)
    assert all(rows[0][header] == "" for header in STAFF_DETAIL_HEADERS[-8:])

    path = tmp_path / "result.xlsx"
    Workbook().save(path)
    write_staff_detail_sheet(path, ctx, result)
    ws = load_workbook(path)[SHEET_STAFF_DETAIL]
    assert ws.max_column == 23
    assert [ws.cell(7, column).value for column in range(1, 24)] == STAFF_DETAIL_HEADERS
    assert [ws.cell(8, column).value for column in range(1, 24)] == [str(i) for i in range(1, 24)]
    assert ws["C9"].value == "Мышенский Дмитрий Игоревич"
    assert ws["M11"].value == 10_100
