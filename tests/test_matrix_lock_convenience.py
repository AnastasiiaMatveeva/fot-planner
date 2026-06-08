from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from fot_planner.excel import SHEET_FOT_MATRIX, create_template, load_context


def test_matrix_lock_by_asterisk(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    create_template(path)

    fot_matrix = pd.DataFrame({"договор": ["C001"]})
    for m in range(1, 13):
        fot_matrix[str(m)] = [""]
    fot_matrix["12"] = "2000000*"

    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        fot_matrix.to_excel(w, sheet_name=SHEET_FOT_MATRIX, index=False)

    wb = load_workbook(path)
    ws = wb[SHEET_FOT_MATRIX]
    for col in range(2, ws.max_column + 1):
        ws.cell(2, col).fill = PatternFill()
    wb.save(path)

    ctx = load_context(path)
    dec = next(mb for mb in ctx.contracts[0].monthly_budgets if mb.month == 12)
    assert dec.lock is True
    assert dec.inflow_amount == 2_000_000


def test_matrix_lock_by_fill(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    create_template(path)

    fot_matrix = pd.DataFrame({"договор": ["C001"]})
    for m in range(1, 13):
        fot_matrix[str(m)] = [""]
    fot_matrix["12"] = 2_000_000

    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        fot_matrix.to_excel(w, sheet_name=SHEET_FOT_MATRIX, index=False)

    wb = load_workbook(path)
    ws = wb[SHEET_FOT_MATRIX]
    yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
    for col in range(2, ws.max_column + 1):
        if str(ws.cell(1, col).value).strip() in ("12", "12.0"):
            ws.cell(2, col).fill = yellow
            break
    wb.save(path)

    ctx = load_context(path)
    dec = next(mb for mb in ctx.contracts[0].monthly_budgets if mb.month == 12)
    assert dec.lock is True
    assert dec.inflow_amount == 2_000_000


def test_default_inflow_at_start_not_locked(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    create_template(path)
    ctx = load_context(path)
    jan = next(mb for mb in ctx.contracts[0].monthly_budgets if mb.month == 1)
    dec = next(mb for mb in ctx.contracts[0].monthly_budgets if mb.month == 12)
    assert jan.lock is False
    assert jan.inflow_amount == pytest.approx(1_200_000, rel=0.01)
    assert dec.inflow_amount == pytest.approx(0.0, abs=0.01)
