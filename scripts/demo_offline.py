# -*- coding: utf-8 -*-
"""Демо-набор без сервиса: якорный разбор → реестр в памяти → вход → решатель.

Тот же путь, что проходит документ в приложении, но без модели и без базы:
четыре файла из data/demo читаются якорным разбором (docs/ui/extract.py),
раскладываются по реестру в памяти тем же кодом, что в сервисе
(intake._store_passport), собираются во входной файл сборщиком приложения
(build_input.build) и решаются тем же исполняемым файлом. Правила замещения
берутся из реестра сервиса (должности.xlsx), справочник — из шаблона.

Печатает: статус, виды выплат по договорам, кто где и на какой ставке,
закрытие трудоёмкости, проверки условий (rules.check) и разделы отчёта.
Так подбирается состав демо: пока здесь не сойдётся, в сервис нести нечего.

    .venv/Scripts/python.exe scripts/demo_offline.py [--docs data/demo] [--limit 120]
"""
from __future__ import annotations

import argparse
import collections
import json
import os
import subprocess
import sys
import tempfile
import time
import types

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(ROOT, "app"))
sys.path.insert(0, os.path.join(ROOT, "docs", "ui"))
sys.path.insert(0, os.path.join(ROOT, "src"))

FILES = ["Штатное расписание 2026.xlsx", "Договоры 2026.xlsx",
         "РКМ ГОЗ Радар-26.xlsx", "Структура цены Грант-26.xlsx"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs", default=os.path.join(ROOT, "data", "demo"))
    ap.add_argument("--limit", type=int, default=120, help="лимит решателя на стадию, с")
    ap.add_argument("--out", default=None, help="куда положить вход и результат")
    ap.add_argument("--settings", default="{}", help="JSON настроек: {\"разрешить дефицит\": \"да\"}")
    args = ap.parse_args()

    os.environ["FOT_SKIP_ORPHANS"] = "1"   # не трогать прогоны сервера
    import build_input
    import extract
    import intake
    import main as app_main
    import reference
    import report
    import rules
    from db import Base, Substitution, session
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    mem = sessionmaker(bind=engine, autoflush=False, future=True)()

    # Правила замещения — из реестра сервиса (общая нормативная база).
    real = session()
    try:
        for s in real.query(Substitution).all():
            mem.add(Substitution(position=s.position, replaced_by=s.replaced_by,
                                 source=s.source))
    finally:
        real.close()
    mem.commit()

    passport_all = {}
    for i, name in enumerate(FILES, 1):
        path = os.path.join(args.docs, name)
        out = extract.extract(path)
        p = out["passport"]
        for line in out["log"]:
            print("  ", name[:28], "|", line)
        for q in out.get("questions") or []:
            print("  ВОПРОС:", q.get("field"), "—", q.get("why"))
        doc = types.SimpleNamespace(id=i, name=name, path=path)
        counts = intake._store_passport(mem, types.SimpleNamespace(year=2026), p, doc)
        print("  ", name[:28], "→", counts)
        if p.get("settings"):
            passport_all["settings"] = p["settings"]

    fake_case = types.SimpleNamespace(year=2026, passport=json.dumps(passport_all, ensure_ascii=False),
                                      plan_settings=None)
    data = app_main._registry_data(mem, fake_case)
    settings = json.loads(args.settings)
    if settings:
        data["settings"] = {**(data.get("settings") or {}), **settings}
    out_dir = args.out or tempfile.mkdtemp(prefix="fot_demo_")
    os.makedirs(out_dir, exist_ok=True)
    src = os.path.join(out_dir, "demo_input.xlsx")
    res = os.path.join(out_dir, "demo_result.xlsx")
    warn = []
    build_input.build(reference.TEMPLATE, src, data, warn)
    for w in warn:
        print("  допущение:", w)
    print("вход:", src)

    t0 = time.time()
    p = subprocess.run([app_main.EXE, "solve", "-i", src, "-o", res, "--time-limit", str(args.limit)],
                       capture_output=True, text=True, cwd=ROOT, timeout=3600)
    sec = time.time() - t0
    print(p.stdout.strip())
    for line in (p.stderr or "").splitlines():
        if line.startswith("[стадия]"):
            print(" ", line)
    if p.returncode != 0:
        print("stderr:", (p.stderr or "").strip()[-2000:])
        print("НЕ СОШЛОСЬ за %.0f с" % sec)
        return 1
    print("решено за %.0f с; результат: %s" % (sec, res))

    # ── что вышло ────────────────────────────────────────────────────
    import openpyxl
    wb = openpyxl.load_workbook(res, data_only=True)
    ws = wb["План выплат"]
    hdr = [c.value for c in ws[1]]
    kinds = collections.Counter()
    by_ctr = collections.defaultdict(lambda: collections.Counter())
    by_emp = collections.defaultdict(lambda: collections.Counter())
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        k, a = d.get("вид выплаты"), d.get("сумма") or 0
        kinds[k] += a
        by_ctr[d.get("договор")][k] += a
        by_emp[d.get("ФИО")][k] += a
    print("\nвиды выплат за год:")
    for k, v in kinds.most_common():
        print("  %-32s %12.0f" % (k, v))
    print("по договорам:")
    for c, cc in sorted(by_ctr.items()):
        print("  %-11s" % c, {k: round(v) for k, v in cc.items()})
    print("по людям:")
    for e, cc in sorted(by_emp.items()):
        print("  %-30s" % e, {k: round(v) for k, v in cc.items()})
    ws = wb["открытые_ставки"]
    hdr = [c.value for c in ws[1]]
    part = collections.defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if str(d.get("основное")).lower() != "да":
            part[d.get("код строки")].append((d.get("договор"), d.get("месяц"), d.get("ставка")))
    print("совместительства (человек: договор, месяцы):")
    for e, rows in sorted(part.items()):
        by = collections.defaultdict(list)
        for c, m, rt in rows:
            by[(c, rt)].append(m)
        print("  %s:" % e, "; ".join("%s ×%s мес %s" % (c, rt, ms) for (c, rt), ms in by.items()))

    rep = report.report(src, res)
    print("\nотчёт: разделы")
    for key in ("договоры", "освоение", "касса", "виды", "регистр", "ставки", "бэп", "п4",
                "трудоемкость", "кто", "незакрыто"):
        v = rep.get(key)
        n = len(v["договоры"]) if isinstance(v, dict) and "договоры" in v else len(v or [])
        print("  %-14s %s" % (key, n))
    for l in rep["трудоемкость"]:
        print("  трудоёмкость %-11s %-26s план %s факт %s людей %s/%s %s"
              % (l["договор"], l["строка"], l["план чел-мес"], l["факт чел-мес"],
                 l["людей макс"], l["людей предел"], l["статус"]))
    for row in rep["п4"]:
        print("  П4:", {k: row[k] for k in list(row)[:6]})
    checks = rules.check(src, res)
    broken = [c for c in checks if c["состояние"] == "нарушено"]
    print("проверок условий: %d, нарушено: %d" % (len(checks), len(broken)))
    for c in broken:
        print("  НАРУШЕНО:", c["правило"], "—", c["факт"])
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
