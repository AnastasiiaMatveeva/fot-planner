# -*- coding: utf-8 -*-
"""Ограничения работы экономиста: числа, а не формулировки.

По каждому условию считается, насколько план к нему подошел: факт, предел,
запас и доля предела. Экономисту нужно не «соблюдено», а «использовано 92 %
предела, узкое место — Иванов, июнь: осталось 13 000 ₽». Тогда видно не
только то, что план допустим, но и где он держится на волоске.

Считается независимо от решателя: по входному файлу и по готовому плану
выплат. Условия «нарушать нельзя» — жесткие: нарушение такого условия значит
ошибку в данных или в сервисе, а не выбор между вариантами.
"""
from __future__ import annotations

import datetime as dt
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
if os.path.join(ROOT, "src") not in sys.path:
    sys.path.insert(0, os.path.join(ROOT, "src"))

#: Виды выплат: как названы в плане и в графе разрешений договора.
KINDS = {
    "оклад": "оклад разрешен",
    "120": "120 разрешена",
    "122": "122 разрешена",
    "124": "124 разрешена",
    "152": "152 разрешена",
    "приказ": "стимулирующая приказом разрешена",
}
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
          "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
RU_MONTH = {m: i + 1 for i, m in enumerate(MONTHS)}
SHORT = ["янв", "фев", "мар", "апр", "май", "июн",
         "июл", "авг", "сен", "окт", "ноя", "дек"]

#: Тип условия. «Нарушать нельзя» — жесткое правило. «С допуском» —
#: отклонение возможно и штрафуется целевой функцией. «Показатель» — не
#: условие, а цифра, на которую смотрят.
HARD, SOFT, INFO = "нарушать нельзя", "с допуском", "показатель"
#: С точностью до рубля: копейки в плане ФОТ ничего не решают.
EPS = 1.0
#: Порог для человеко-месяцев. EPS — рублевый; с ним ставка за месяц (число
#: около единицы) считалась нулем, и закрытая строка выглядела незакрытой.
PM_EPS = 0.001
#: Сколько строк показывать в таблице узких мест.
SHOW = 8


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _yes(v):
    return str(v or "").strip().lower() in ("да", "true", "1", "yes", "+")


def _month(v):
    """Месяц из даты входного файла: число, дата или «01.06.2026»."""
    if v in (None, ""):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.month
    s = str(v).strip()
    if "." in s:
        parts = s.split(".")
        if len(parts) >= 2 and parts[1].isdigit():
            return int(parts[1])
    n = _num(s)
    if n and 1 <= n <= 80000:      # порядковый номер даты Excel
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(n))).month
    return None


def _kind_of(raw):
    s = str(raw or "").strip().lower()
    if s.startswith("оклад"):
        return "оклад"
    for k in ("120", "122", "124", "152"):
        if s.startswith(k):
            return k
    if "приказ" in s:
        return "приказ"
    return None


MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
               "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def _months_of_row(row):
    """План по месяцам из граф «Январь» … «Декабрь»; None — строка годовая."""
    out = {}
    for m, name in enumerate(MONTH_NAMES, start=1):
        v = _num(row.get(name))
        if v is None:
            v = _num(row.get(name.lower()))
        if v:
            out[m] = v
    return out or None


def _rows_of(ws):
    hdr = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        row = {hdr[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)
               if hdr[c - 1]}
        if any(v not in (None, "") for v in row.values()):
            yield row


def _read_input(path):
    from openpyxl import load_workbook
    from fot_planner.position_reference import normalize_position

    wb = load_workbook(path, data_only=True)
    data = {"employees": [], "contracts": {}, "limits": {}, "inflow": {},
            "labor": [], "settings": {}, "year": None}
    for row in _rows_of(wb["сотрудники"]):
        code = row.get("код строки")
        if not code:
            continue
        data["employees"].append({
            "code": str(code), "fio": row.get("фио") or str(code),
            "position": str(row.get("должность") or ""),
            "department": str(row.get("подразделение") or ""),
            "rate": _num(row.get("ставка")) or 1.0,
            "salary": _num(row.get("зарплата")) or 0.0,
            "from": _month(row.get("дата начала")) or 1,
            "to": _month(row.get("дата окончания")) or 12,
        })
    for row in _rows_of(wb["договоры"]):
        code = row.get("код")
        if not code:
            continue
        data["contracts"][str(code)] = {
            "code": str(code), "name": row.get("название") or "",
            "goz": _yes(row.get("ГОЗ")), "fot": _num(row.get("фот")) or 0.0,
            "from": _month(row.get("дата начала")) or 1,
            "to": _month(row.get("дата окончания")) or 12,
            "kinds": {k: _yes(row.get(col)) for k, col in KINDS.items()},
        }
    if "лимиты_по_должностям" in wb.sheetnames:
        for row in _rows_of(wb["лимиты_по_должностям"]):
            pos = row.get("должность")
            if pos:
                data["limits"][normalize_position(str(pos))] = {
                    "оклад": _num(row.get("оклад")), "П2556": _num(row.get("П2556")),
                    "П4": _num(row.get("П4")), "БЭП": _num(row.get("БЭП")),
                }
    ws = wb["фот_по_месяцам"]
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if code:
            data["inflow"][str(code)] = [_num(ws.cell(r, c).value) or 0.0
                                         for c in range(2, 14)]
    if "трудоемкость_по_договорам" in wb.sheetnames:
        for row in _rows_of(wb["трудоемкость_по_договорам"]):
            if row.get("договор"):
                data["labor"].append({
                    "contract": str(row["договор"]),
                    "position": str(row.get("должность") or ""),
                    "person_months": _num(row.get("трудоемкость")),
                    "avg": _num(row.get("средняя стоимость выполнения работ в месяц")),
                    "months": _months_of_row(row),
                })
    if "настройки" in wb.sheetnames:
        for row in _rows_of(wb["настройки"]):
            data["settings"] = dict(row)
            data["year"] = _num(row.get("год"))
            break
    data["norm"] = normalize_position
    return data


def _read_result(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    out = {"plan": [], "labor_control": []}
    for row in _rows_of(wb["План выплат"]):
        code = row.get("код строки")
        month = RU_MONTH.get(str(row.get("месяц") or "").strip())
        amount = _num(row.get("сумма"))
        if not code or not month or not amount:
            continue
        out["plan"].append({
            "emp": str(code), "fio": row.get("ФИО") or str(code),
            "position": str(row.get("должность") or ""),
            "month": month, "contract": str(row.get("договор") or ""),
            "kind": _kind_of(row.get("вид выплаты")), "amount": amount,
        })
    # Дефицит выплат: если экономист разрешил дефицит, решатель пишет, кому и
    # сколько недоплачено; отчёт должен это показать, а не спрятать.
    if "Дефициты" in wb.sheetnames:
        out["deficits"] = []
        for row in _rows_of(wb["Дефициты"]):
            code = row.get("код строки")
            month = RU_MONTH.get(str(row.get("месяц") or "").strip()) or _num(row.get("месяц"))
            if not code or not month:
                continue
            out["deficits"].append({
                "emp": str(code), "fio": row.get("ФИО") or str(code), "month": int(month),
                "due": _num(row.get("требовалось выплатить")) or 0.0,
                "paid": _num(row.get("выплачено")) or 0.0,
                "gap": _num(row.get("дефицит")) or 0.0,
                "why": row.get("причина") or "",
            })
    # Открытые ставки решателя: ставка и признак основного места по месяцам.
    if "открытые_ставки" in wb.sheetnames:
        out["rates"] = []
        for row in _rows_of(wb["открытые_ставки"]):
            code, month = row.get("код строки"), _num(row.get("месяц"))
            if not code or not month:
                continue
            out["rates"].append({
                "emp": str(code), "contract": str(row.get("договор") or ""),
                "month": int(month), "rate": _num(row.get("ставка")) or 0.0,
                "main": _yes(row.get("основное")),
            })
    if "Контроль трудоёмкости" in wb.sheetnames:
        ws = wb["Контроль трудоёмкости"]
        # У договора несколько строк РКМ, и под каждой — своя таблица по
        # месяцам («Показатель»). Раньше эта таблица выключала разбор строк
        # насовсем, и от договора читалась только первая строка: контроль
        # человеко-месяцев проверял «Старшего научного сотрудника» и молчал
        # про инженеров.
        current, in_rules, in_months = None, False, False
        for r in range(1, ws.max_row + 1):
            text = str(ws.cell(r, 1).value or "").strip()
            if text.startswith("ДОГОВОР:"):
                current, in_rules, in_months = (
                    text.replace("ДОГОВОР:", "").strip(), False, False)
                continue
            if text == "Строка / группа":
                in_rules, in_months = True, False
                continue
            if text == "Показатель":
                in_months = True
                continue
            if not text:
                in_months = False
                continue
            # Разбивка решателя: кто и сколько закрыл по каждой строке.
            if text == "Договор" and str(ws.cell(r, 4).value or "").strip() == "Сотрудник":
                in_rules, in_months = False, False
                continue
            ind = str(ws.cell(r, 6).value or "").strip()
            if ind in ("Закрыто чел.-мес.", "Всего денег на строку"):
                vals = [_num(ws.cell(r, 7 + i).value) or 0.0 for i in range(12)]
                out.setdefault("labor_people", []).append({
                    "contract": text, "row": str(ws.cell(r, 2).value or "").strip(),
                    "fio": str(ws.cell(r, 4).value or "").strip(),
                    "position": str(ws.cell(r, 5).value or "").strip(),
                    "kind": "pm" if ind.startswith("Закрыто") else "pay", "months": vals,
                })
                continue
            # Заголовок разбивки («Что сформировало трудоёмкость») стоит в той
            # же графе, что и строки, но чисел под ним нет.
            if in_rules and not in_months and _num(ws.cell(r, 2).value) is not None:
                out["labor_control"].append({
                    "contract": current or "", "row": text,
                    "plan_pm": _num(ws.cell(r, 2).value),
                    "fact_pm": _num(ws.cell(r, 3).value),
                    "plan_sum": _num(ws.cell(r, 5).value),
                    "fact_sum": _num(ws.cell(r, 6).value),
                })
    return out


def _by(plan, *keys):
    out = {}
    for p in plan:
        out.setdefault(tuple(p[k] for k in keys), []).append(p)
    return out


def _sum(rows, kinds=None):
    return sum(p["amount"] for p in rows if kinds is None or p["kind"] in kinds)


def _mo(v):
    if v is None:
        return "—"
    v = float(v)
    if abs(v) < 100 and abs(v - round(v)) > 0.001:
        s = ("%.2f" % abs(v)).rstrip("0").rstrip(".").replace(".", ",")
    else:
        s = "{:,.0f}".format(round(abs(v))).replace(",", " ")
    return ("−" if v < -0.0001 else "") + s


def _mark(out, section):
    for r in out:
        r.setdefault("раздел", section)


# ── сборка условия ───────────────────────────────────────────────────
def _limit_rule(name, meaning, where, unit, rows, kind=HARD, mode="не больше",
                empty="", tolerance=0.0, absolute_tolerance=None):
    """Условие с числовым пределом: факт против предела в каждом случае.

    ``rows`` — [{"объект", "факт", "предел"}]. Считается доля предела, запас
    и самые узкие места: не только «соблюдено», но и насколько близко план
    подошел к границе и что сломается от первой же правки.
    """
    checked = []
    for r in rows:
        limit, fact = r.get("предел"), r.get("факт")
        if limit in (None, 0) or fact is None:
            continue
        allowed = max(EPS if absolute_tolerance is None else absolute_tolerance,
                      abs(limit) * tolerance)
        over = (fact > limit + EPS) if mode == "не больше" else abs(fact - limit) > allowed + 1e-7
        checked.append({"объект": r["объект"], "факт": round(fact, 2),
                        "предел": round(limit, 2), "доля": round(fact / limit, 4),
                        "запас": round(limit - fact, 2), "нарушено": over})
    if not checked:
        return {"правило": name, "смысл": meaning, "где": where, "тип": kind,
                "состояние": "не применялось", "единица": unit, "проверено": 0,
                "факт": empty, "строки": [], "нарушений": 0, "предел_общий": None}
    broken = [c for c in checked if c["нарушено"]]
    tight = sorted(checked, key=lambda c: -c["доля"])
    worst = tight[0]
    if mode == "не больше":
        fact_text = ("использовано %d %% предела; узкое место — %s: запас %s %s"
                     % (round(worst["доля"] * 100), worst["объект"],
                        _mo(worst["запас"]), unit))
    else:
        big = max(checked, key=lambda c: abs(c["запас"]))
        fact_text = ("расхождений %d из %d; наибольшее — %s: %s %s"
                     % (len(broken), len(checked), big["объект"],
                        _mo(big["запас"]), unit))
    # Что показывать в таблице. При равенстве интересны только расхождения:
    # тридцать шесть одинаковых строк «160 000 из 160 000, запас 0» — шум.
    # При пределе интересны самые узкие места, но одинаковые случаи
    # сворачиваются в один со счетчиком.
    candidates = broken if mode == "равно" else broken + [c for c in tight
                                                          if not c["нарушено"]]
    shown, seen = [], {}
    for c in candidates:
        key = (c["факт"], c["предел"])
        if key in seen:
            seen[key]["таких же"] = seen[key].get("таких же", 1) + 1
            continue
        seen[key] = c
        shown.append(c)
    limits = {c["предел"] for c in checked}
    return {"правило": name, "смысл": meaning, "где": where, "тип": kind,
            "состояние": "нарушено" if broken else "соблюдено", "единица": unit,
            "проверено": len(checked), "факт": fact_text,
            "использовано": round(worst["доля"], 4),
            "подпись предела": "нужно" if mode == "равно" else "предел",
            "предел_общий": checked[0]["предел"] if len(limits) == 1 else None,
            "строки": shown[:SHOW], "нарушений": len(broken)}


def _count_rule(name, meaning, where, checked, bad, kind=HARD, empty="", unit=""):
    """Условие без числового предела: сошлось или нет и в скольких случаях."""
    if not checked:
        return {"правило": name, "смысл": meaning, "где": where, "тип": kind,
                "состояние": "не применялось", "единица": unit, "проверено": 0,
                "факт": empty, "строки": [], "нарушений": 0, "предел_общий": None}
    return {"правило": name, "смысл": meaning, "где": where, "тип": kind,
            "состояние": "нарушено" if bad else "соблюдено", "единица": unit,
            "проверено": checked, "нарушений": len(bad), "предел_общий": None,
            "факт": "проверено случаев: %d, нарушений %d" % (checked, len(bad)),
            "строки": bad[:SHOW]}


def check(input_path, result_path):
    """Все условия и их состояние в этом плане, в числах."""
    inp = _read_input(input_path)
    res = _read_result(result_path)
    plan = res["plan"]
    emps = {e["code"]: e for e in inp["employees"]}
    ctr, norm, lim = inp["contracts"], inp["norm"], inp["limits"]
    by_em = _by(plan, "emp", "month")
    out = []

    def base_of(position):
        return (lim.get(norm(position)) or {}).get("оклад")

    def rate_of(rows):
        """Ставка по окладу: оклад, деленный на оклад должности."""
        base = base_of(rows[0]["position"]) if rows else None
        return (_sum(rows, {"оклад"}) / base) if base else None

    # ── зарплата и выплаты ───────────────────────────────────────────
    rows = []
    for e in inp["employees"]:
        for m in range(e["from"], e["to"] + 1):
            if e["salary"]:
                rows.append({"объект": "%s, %s" % (e["fio"], SHORT[m - 1]),
                             "факт": _sum(by_em.get((e["code"], m), [])),
                             "предел": e["salary"]})
    out.append(_limit_rule(
        "Человек получает всю месячную зарплату",
        "Сумма из штатного расписания закрывается целиком — окладом и "
        "надбавками. Недоплата возможна, только если дефицит разрешен в "
        "настройках расчета.",
        "План выплат", "₽", rows, mode="равно"))

    checked, bad = 0, []
    for (code, m, c), rr in _by(plan, "emp", "month", "contract").items():
        oklad, base = _sum(rr, {"оклад"}), base_of(rr[0]["position"])
        if not oklad or not base:
            continue
        checked += 1
        rate = oklad / base
        if abs(rate * 4 - round(rate * 4)) > 0.02:
            bad.append({"объект": "%s, %s, %s" % (rr[0]["fio"], SHORT[m - 1], c),
                        "что": "оклад %s ₽ — это %s ставки, а шаг ставки четверть"
                               % (_mo(oklad), round(rate, 3))})
    out.append(_count_rule(
        "Оклад идет по ставке, а не произвольной суммой",
        "Оклад на договоре — оклад должности, умноженный на открытую там "
        "ставку. Шаг ставки — четверть; уменьшать оклад «для удобства» нельзя, "
        "остаток зарплаты добирается надбавками.",
        "Реестр → Справочник должностей", checked, bad,
        empty="в справочнике нет окладов по этим должностям"))

    checked, bad = 0, []
    for (code, m), rr in by_em.items():
        salary_ctr = {p["contract"] for p in rr if p["kind"] == "оклад"}
        for p in rr:
            if p["kind"] != "122":
                continue
            checked += 1
            if salary_ctr and p["contract"] not in salary_ctr:
                bad.append({"объект": "%s, %s" % (p["fio"], SHORT[m - 1]),
                            "что": "122 на %s, оклад на %s"
                                   % (p["contract"], ", ".join(sorted(salary_ctr)))})
    out.append(_count_rule(
        "Надбавка 122 — с того же договора, что и оклад",
        "За качество платят там же, где сидит оклад этой ставки.",
        "План выплат", checked, bad, empty="122 в этом плане не назначалась"))

    bad = []
    for p in plan:
        c = ctr.get(p["contract"])
        if c and p["kind"] and not c["kinds"].get(p["kind"], False):
            bad.append({"объект": "%s, %s, %s" % (p["fio"], SHORT[p["month"] - 1],
                                                  p["contract"]),
                        "что": "вид «%s» договором запрещен" % p["kind"]})
    out.append(_count_rule(
        "Платим только теми видами, которые разрешены договором",
        "Договор перечисляет, что с него можно платить: оклад, 120, 122, 124, "
        "152, стимулирующую приказом. Остальное с него платить нельзя.",
        "Реестр → Договоры", len(plan), bad))

    checked, bad = 0, []
    for (code, m), rr in by_em.items():
        kinds = {p["kind"] for p in rr}
        if "152" not in kinds:
            continue
        checked += 1
        extra = kinds - {"оклад", "152"}
        if extra:
            bad.append({"объект": "%s, %s" % (rr[0]["fio"], SHORT[m - 1]),
                        "что": "вместе со 152 назначены %s" % ", ".join(sorted(extra))})
    out.append(_count_rule(
        "Режим «оклад плюс 152»",
        "Если строке сотрудника в месяце назначена 152 за дополнительную "
        "работу, других надбавок и приказа в этом месяце у нее нет.",
        "План выплат", checked, bad, empty="152 в этом плане не используется"))
    _mark(out, "Зарплата и выплаты")

    # ── договоры и деньги ────────────────────────────────────────────
    bad, seen = [], set()
    for p in plan:
        c = ctr.get(p["contract"])
        if c and not (c["from"] <= p["month"] <= c["to"]):
            key = "%s, %s" % (p["contract"], SHORT[p["month"] - 1])
            if key not in seen:
                seen.add(key)
                bad.append({"объект": key,
                            "что": "договор действует %s—%s"
                                   % (SHORT[c["from"] - 1], SHORT[c["to"] - 1])})
    out.append(_count_rule(
        "Договор платит только в свои сроки",
        "Ни рубля до начала и после окончания договора.",
        "Реестр → Договоры", len(plan), bad))

    cash_rows, series, worst = [], [], None
    for code in sorted(ctr):
        got = inp["inflow"].get(code) or [0.0] * 12
        paid = [0.0] * 12
        for p in plan:
            if p["contract"] == code:
                paid[p["month"] - 1] += p["amount"]
        if not any(got) and not any(paid):
            continue
        bal, points, low, low_m = 0.0, [], None, 1
        for m in range(12):
            bal += got[m] - paid[m]
            points.append(round(bal, 2))
            if low is None or bal < low:
                low, low_m = bal, m + 1
        series.append({"имя": code, "точки": points})
        # Три графы «факт 0, предел 0, запас 0» ничего не говорят: у кассы
        # интересен сам остаток и месяц, когда он был самым низким.
        cash_rows.append({"объект": "%s, %s" % (code, SHORT[low_m - 1]),
                          "что": ("остаток %s ₽ — самый низкий за год"
                                  % _mo(low)) if low >= -EPS else
                                 ("минус %s ₽ — потрачено больше, чем поступило"
                                  % _mo(-low)),
                          "низ": round(low, 2), "нарушено": low < -EPS})
        if worst is None or low < worst["низ"]:
            worst = cash_rows[-1]
    cash = {
        "правило": "Деньги нельзя потратить раньше, чем они пришли",
        "смысл": "Касса договора считается помесячно, неизрасходованный "
                 "остаток переносится только вперед. Минуса на договоре быть "
                 "не может.",
        "где": "Освоение", "тип": HARD, "единица": "₽", "предел_общий": 0.0,
        "проверено": len(cash_rows),
        "состояние": ("не применялось" if not cash_rows else
                      ("нарушено" if any(r["нарушено"] for r in cash_rows) else "соблюдено")),
        "нарушений": sum(1 for r in cash_rows if r["нарушено"]),
        "факт": ("самый низкий остаток за год — %s: %s ₽" % (worst["объект"], _mo(worst["низ"]))
                 if worst else "поступлений и выплат нет"),
        "строки": sorted(cash_rows, key=lambda r: r["низ"])[:SHOW],
        "график": {"вид": "месяцы", "подпись": "остаток на конец месяца, ₽",
                   "ряды": series[:6], "линия": 0.0},
    }
    out.append(cash)

    rows = [{"объект": code, "факт": sum(p["amount"] for p in plan if p["contract"] == code),
             "предел": c["fot"]} for code, c in sorted(ctr.items()) if c["fot"]]
    out.append(_limit_rule(
        "Выплаты не превышают ФОТ договора",
        "Фонд оплаты труда договора — верхняя граница всех выплат с него за год.",
        "Реестр → Договоры", "₽", rows))

    spend = _limit_rule(
        "ФОТ договоров освоен",
        "Неосвоенный остаток расчет не ломает, но это деньги, которые остались "
        "на договоре к концу года.",
        "Освоение", "₽", rows, kind=INFO)
    if spend["состояние"] != "не применялось":
        left = [r for r in spend["строки"] if r["запас"] > EPS]
        fot = sum(r["предел"] for r in rows)
        paid = sum(r["факт"] for r in rows)
        spend["состояние"] = "внимание" if left else "соблюдено"
        spend["использовано"] = round(paid / fot, 4) if fot else None
        spend["факт"] = ("освоено %s ₽ из %s ₽ — %d %%; не освоено %s ₽"
                         % (_mo(paid), _mo(fot), round(100 * paid / fot) if fot else 0,
                            _mo(fot - paid)))
        spend["строки"] = sorted(spend["строки"], key=lambda r: -r["запас"])[:SHOW]
    out.append(spend)
    _mark(out, "Договоры и деньги")

    # ── лимиты по должностям ─────────────────────────────────────────
    rows = []
    for (code, m, c), rr in _by(plan, "emp", "month", "contract").items():
        p2556 = (lim.get(norm(rr[0]["position"])) or {}).get("П2556")
        rate = rate_of(rr)
        if p2556 and rate:
            rows.append({"объект": "%s, %s, %s" % (rr[0]["fio"], SHORT[m - 1], c),
                         "факт": _sum(rr, {"оклад", "122"}), "предел": p2556 * rate})
    out.append(_limit_rule(
        "П2556: оклад и 122 не выше предела по должности",
        "Приказ 2556 задает максимум штатной части на полную ставку по каждой "
        "должности. На доле ставки предел уменьшается пропорционально.",
        "Реестр → Справочник должностей", "₽", rows,
        empty="в справочнике нет П2556 по этим должностям"))

    rows, series = [], {}
    for (c, m), rr in _by(plan, "contract", "month").items():
        contract = ctr.get(c)
        if not contract or not contract["goz"]:
            continue
        cap, staff, rates = 0.0, 0.0, 0.0
        for (code,), er in _by(rr, "emp").items():
            b = (lim.get(norm(er[0]["position"])) or {}).get("БЭП")
            r = rate_of(er)
            if b and r:
                cap += b * r
            if r:
                rates += r
            staff += _sum(er, {"оклад", "122"})
        if not cap or not rates:
            continue
        rows.append({"объект": "%s, %s" % (c, SHORT[m - 1]),
                     "факт": staff / rates, "предел": cap / rates})
        series.setdefault(c, [None] * 12)[m - 1] = round(staff / rates, 2)
    bep = _limit_rule(
        "БЭП: средняя зарплата по ГОЗ-договору в пределах базовой",
        "На гособоронзаказе ограничена не отдельная зарплата, а средняя по "
        "договору за месяц: сумма оклада и 122, деленная на сумму ставок.",
        "Реестр → Справочник должностей", "₽", rows,
        empty="ГОЗ-договоров в плане нет")
    if series:
        bep["график"] = {"вид": "месяцы",
                         "подпись": "средняя штатная часть на ставку, ₽",
                         "ряды": [{"имя": k, "точки": v} for k, v in list(series.items())[:6]],
                         "линия": min(r["предел"] for r in rows) if rows else None}
    out.append(bep)

    rows = []
    for (code, m), rr in by_em.items():
        if not any(p["kind"] == "124" for p in rr):
            continue
        p4 = (lim.get(norm(rr[0]["position"])) or {}).get("П4")
        rate = sum(filter(None, (rate_of(cr) for (_,), cr in _by(rr, "contract").items())))
        if p4 and rate:
            rows.append({"объект": "%s, %s" % (rr[0]["fio"], SHORT[m - 1]),
                         "факт": _sum(rr, {"оклад", "122", "124"}), "предел": p4 * rate})
    out.append(_limit_rule(
        "П4 при надбавке за интенсивность",
        "Если сотруднику назначена 124, включается предел П4: он смотрит "
        "оклад, 122 и 124 этой строки за месяц целиком, по всем договорам.",
        "Реестр → Справочник должностей", "₽", rows,
        empty="надбавка 124 в этом плане не назначалась"))
    _mark(out, "Лимиты по должностям")

    # ── ставки ───────────────────────────────────────────────────────
    rows, extra = [], []
    for (code, m), rr in by_em.items():
        total = sum(filter(None, (rate_of(cr) for (_,), cr in _by(rr, "contract").items())))
        if not total:
            continue
        staff = emps.get(code, {}).get("rate") or 1.0
        rows.append({"объект": "%s, %s" % (rr[0]["fio"], SHORT[m - 1]),
                     "факт": total, "предел": 1.5})
        if total - staff > 0.01:
            extra.append({"объект": "%s, %s" % (rr[0]["fio"], SHORT[m - 1]),
                          "факт": total - staff, "предел": 0.5})
    out.append(_limit_rule(
        "Суммарная ставка — не больше 1,5",
        "Штатная ставка сохраняется, сверх нее сервис может открыть "
        "совместительство. Обычному сотруднику — до полутора ставок в сумме.",
        "План выплат", "ставки", rows))
    out.append(_limit_rule(
        "Совместительство — не больше 0,5 ставки",
        "Сверх штатной ставки сотруднику можно открыть половину ставки, и "
        "только по должности, которая ему разрешена правилами замещения.",
        "Реестр → Правила замещения", "ставки", extra,
        empty="совместительство в этом плане не открывалось"))
    _mark(out, "Ставки и совместительство")

    # ── трудоемкость ─────────────────────────────────────────────────
    tol = _num(inp["settings"].get("допуск трудоёмкости")) or 0.0
    note = (" Допуск в настройках расчета: %s %%." % _mo(tol * 100)) if tol else ""
    pm_rows, sum_rows = [], []
    for row in res["labor_control"]:
        obj = "%s / %s" % (row["contract"].split("—")[0].strip(), row["row"])
        if row["plan_pm"]:
            pm_rows.append({"объект": obj, "факт": row["fact_pm"] or 0.0,
                            "предел": row["plan_pm"]})
        if row["plan_sum"]:
            sum_rows.append({"объект": obj, "факт": row["fact_sum"] or 0.0,
                             "предел": row["plan_sum"]})
    out.append(_limit_rule(
        "Плановые человеко-месяцы закрыты",
        "По строке расчетно-калькуляционных материалов на договоре должно "
        "набраться столько человеко-месяцев, сколько заложено в РКМ." + note,
        "Реестр → Трудоемкость", "чел.-мес.", pm_rows, kind=SOFT, mode="равно",
        tolerance=tol, absolute_tolerance=0.5,
        empty="трудоемкость по договорам не задана"))
    out.append(_limit_rule(
        "Сумма трудоемкости выбрана",
        "Плановая стоимость строки — человеко-месяцы, умноженные на среднюю "
        "стоимость. Столько же должно быть начислено людям, закрывающим эту "
        "строку." + note,
        "Реестр → Трудоемкость", "₽", sum_rows, kind=SOFT, mode="равно",
        tolerance=tol, absolute_tolerance=0.0,
        empty="трудоемкость по договорам не задана"))

    bad = []
    for lp in inp["labor"]:
        if not lp["position"]:
            continue
        # A labor row may be closed by an employee through an approved
        # substitution (for example, a leading engineer performs a
        # programmer's work).  The detailed attribution sheet is authoritative
        # for this check; requiring an exact position here falsely reports a
        # valid substitution as an uncovered row.
        attributions = [p for p in res.get("labor_people", [])
                        if p.get("contract") == lp["contract"]
                        and p.get("kind") == "pm"
                        and any(float(v or 0) > PM_EPS for v in (p.get("months") or []))]
        who = [p for p in plan if p["contract"] == lp["contract"]
               and norm(p["position"]) == norm(lp["position"])]
        if not who and not attributions:
            bad.append({"объект": "%s / %s" % (lp["contract"], lp["position"]),
                        "что": "на договоре нет выплат по этой должности"})
    out.append(_count_rule(
        "Работу закрывает подходящая должность",
        "Строку трудоемкости может закрыть сотрудник с той же должностью или "
        "тот, кому эта должность разрешена правилами замещения.",
        "Реестр → Правила замещения", len(inp["labor"]), bad,
        empty="трудоемкость по договорам не задана"))
    _mark(out, "Трудоемкость договоров")

    # Дефицит — единственная настройка, смягчающая жесткое условие.
    if _yes(inp["settings"].get("разрешить дефицит")):
        for r in out:
            if r["правило"].startswith("Человек получает"):
                r["тип"] = "смягчено настройкой"
        out.insert(0, {
            "правило": "Дефицит выплат разрешен настройками",
            "смысл": "В настройках расчета разрешено оставить часть зарплаты "
                     "невыплаченной. План при этом сходится, но люди получают меньше.",
            "где": "Настройки расчета", "тип": INFO, "состояние": "внимание",
            "единица": "", "проверено": 0, "нарушений": 0, "строки": [],
            "предел_общий": None,
            "факт": "условие «человек получает всю зарплату» смягчено",
            "раздел": "Настройки расчета"})
    return out


# ── сводка для финансиста ────────────────────────────────────────────
def summary(input_path, result_path):
    """План года цифрами: бюджет, освоение, структура выплат, люди.

    То, с чего финансист начинает разговор о годовом плане: сколько денег
    заложено и сколько разошлось по людям, как выплаты ложатся на месяцы
    рядом с поступлениями, из чего складывается сумма, кто где сидит.
    """
    inp = _read_input(input_path)
    res = _read_result(result_path)
    plan = res["plan"]
    ctr = inp["contracts"]

    paid_m = [0.0] * 12
    got_m = [0.0] * 12
    for p in plan:
        paid_m[p["month"] - 1] += p["amount"]
    for code, row in inp["inflow"].items():
        for m in range(12):
            got_m[m] += row[m] or 0.0

    kinds = {}
    for p in plan:
        kinds[p["kind"] or "прочее"] = kinds.get(p["kind"] or "прочее", 0.0) + p["amount"]

    by_ctr = []
    for code, c in sorted(ctr.items()):
        paid = sum(p["amount"] for p in plan if p["contract"] == code)
        got = sum(inp["inflow"].get(code) or [])
        by_ctr.append({"договор": code, "название": c["name"], "ГОЗ": c["goz"],
                       "ФОТ": round(c["fot"], 2), "поступления": round(got, 2),
                       "выплаты": round(paid, 2),
                       "остаток": round(c["fot"] - paid, 2),
                       "освоение": round(paid / c["fot"], 4) if c["fot"] else None,
                       "месяцы": [round(sum(p["amount"] for p in plan
                                            if p["contract"] == code and p["month"] == m + 1), 2)
                                  for m in range(12)]})

    people = []
    for e in inp["employees"]:
        rows = [p for p in plan if p["emp"] == e["code"]]
        paid = sum(p["amount"] for p in rows)
        months = sorted({p["month"] for p in rows})
        contracts = sorted({p["contract"] for p in rows})
        people.append({"код": e["code"], "фио": e["fio"], "должность": e["position"],
                       "подразделение": e["department"], "ставка": e["rate"],
                       "зарплата": e["salary"], "за год": round(paid, 2),
                       "месяцев": len(months), "договоров": len(contracts),
                       "договоры": contracts})

    total_fot = sum(c["fot"] for c in ctr.values())
    total_paid = sum(paid_m)
    return {
        "год": int(inp["year"] or 0) or None,
        "итого": {
            "ФОТ договоров": round(total_fot, 2),
            "распределено планом": round(total_paid, 2),
            "не распределено": round(total_fot - total_paid, 2),
            "освоение": round(total_paid / total_fot, 4) if total_fot else None,
            "поступления": round(sum(got_m), 2),
            "сотрудников": len(inp["employees"]),
            "договоров": len(ctr),
            "средняя выплата в месяц на человека":
                round(total_paid / max(1, len(inp["employees"])) / 12, 2),
        },
        "месяцы": {"подписи": SHORT, "поступления": [round(v, 2) for v in got_m],
                   "выплаты": [round(v, 2) for v in paid_m],
                   "остаток": [round(sum(got_m[:i + 1]) - sum(paid_m[:i + 1]), 2)
                               for i in range(12)]},
        "виды выплат": [{"вид": k, "сумма": round(v, 2),
                         "доля": round(v / total_paid, 4) if total_paid else 0}
                        for k, v in sorted(kinds.items(), key=lambda kv: -kv[1])],
        "договоры": by_ctr,
        "люди": sorted(people, key=lambda p: -p["за год"]),
    }


def _why_rates(inp, res):
    """Кто закрывает строку трудоёмкости — по разбивке решателя, а не по счёту
    должностей.

    Человеко-месяцы РКМ закрывает только открытая ставка сотрудника на этом
    договоре: надбавки их не создают. Кого решатель отнёс к какой строке, он
    говорит сам — лист «Контроль трудоёмкости» результата. Раньше здесь
    считались люди с совпадающей должностью, и объяснение выходило неверным
    сразу дважды: замещение (главный инженер проекта на строке инженера) не
    попадало в счёт, а ставка человека на договоре приписывалась каждой
    строке этого договора с подходящей должностью. По строке «Инженер»
    (16 чел.-мес., закрывают трое) выходило «1 человек, 8 чел.-мес.» — и
    агент в ленте пересказывал эту неправду.
    """
    out = []
    by_row = {}
    for p in res.get("labor_people") or []:
        if p.get("kind") != "pm":
            continue
        months = [(i + 1, float(v or 0)) for i, v in enumerate(p.get("months") or [])
                  if float(v or 0) > PM_EPS]
        if not months:
            continue
        code = str(p.get("contract") or "").split("—")[0].strip()
        by_row.setdefault((code, str(p.get("row") or "")), []).append((p, months))
    for lp in inp["labor"]:
        if not lp["person_months"]:
            continue
        code, pos = lp["contract"], lp["position"] or ""
        got = by_row.get((code, pos)) or []
        if not got:
            out.append("Договор %s: трудоемкость %s чел.-мес. по должности «%s» "
                       "не закрыта — ставок на этом договоре не открыто."
                       % (code, _mo(lp["person_months"]), pos or "любая"))
            continue
        parts, total = [], 0.0
        for p, months in got:
            amount = sum(v for _m, v in months)
            total += amount
            nums = [m for m, _v in months]
            span = ("%s—%s" % (SHORT[nums[0] - 1], SHORT[nums[-1] - 1])
                    if len(nums) > 1 else SHORT[nums[0] - 1])
            rate = amount / len(nums)
            same = " (замещение)" if norm_position_differs(inp, p, pos) else ""
            parts.append("%s%s — по %s ставки, %s (%s мес. × %s = %s чел.-мес.)"
                         % (p.get("fio") or "?", same, _mo(rate), span,
                            len(nums), _mo(rate), _mo(amount)))
        out.append("Договор %s: трудоемкость %s чел.-мес. по должности «%s», "
                   "закрыто %s чел.-мес. Человеко-месяцы закрывает только "
                   "открытая ставка на этом договоре, и решатель отнёс к этой "
                   "строке: %s."
                   % (code, _mo(lp["person_months"]), pos or "любая",
                      _mo(total), "; ".join(parts)))
    return out


def norm_position_differs(inp, person, position):
    """Человек закрывает строку по правилу замещения, а не своей должностью."""
    norm = inp["norm"]
    have, want = person.get("position") or "", position or ""
    return bool(have and want and norm(have) != norm(want))


def _px(n, one, few, many):
    d, hh = n % 10, n % 100
    if 11 <= hh <= 14:
        return many
    if d == 1:
        return one
    if 2 <= d <= 4:
        return few
    return many


# ── из чего складывается зарплата ────────────────────────────────────
def payroll(input_path, result_path):
    """Состав зарплаты по месяцам: договоры, ставки, виды выплат и смены схемы.

    Матрица «ФИО на месяцы» ничего не говорит: сумма в каждой клетке равна
    месячной зарплате, потому что это жесткое условие. Экономисту нужно то,
    что за этой суммой: с каких договоров она финансируется, на каких
    ставках человек сидит, из каких видов выплат складывается — и как часто
    схема меняется. Смена договора оклада и смена состава выплат — это
    приказы, допники и разговоры с людьми; оптимизатор их штрафует, а
    экономист должен видеть, сколько их получилось.
    """
    inp = _read_input(input_path)
    res = _read_result(result_path)
    plan, lim, norm = res["plan"], inp["limits"], inp["norm"]
    ctr = inp["contracts"]
    order = sorted({p["contract"] for p in plan})

    people = []
    total_switch = total_mix = total_extra = total_opened = total_struct = 0
    for e in inp["employees"]:
        rows = [p for p in plan if p["emp"] == e["code"]]
        if not rows:
            continue
        base = (lim.get(norm(e["position"])) or {}).get("оклад")
        months, prev_ctr, prev_mix, prev_oklad, prev_rate = [], None, None, None, None
        switch = mix_change = extra_months = opened = struct_change = 0
        rates, oklad_share = [], []
        for m in range(1, 13):
            mr = [p for p in rows if p["month"] == m]
            if not mr:
                months.append(None)
                continue
            by_ctr = []
            for code in order:
                cr = [p for p in mr if p["contract"] == code]
                if not cr:
                    continue
                oklad = _sum(cr, {"оклад"})
                by_ctr.append({
                    "код": code,
                    "ставка": round(oklad / base, 2) if base and oklad else None,
                    "сумма": round(_sum(cr), 2),
                    "виды": [{"вид": k, "сумма": round(_sum(cr, {k}), 2)}
                             for k in ("оклад", "120", "122", "124", "152", "приказ")
                             if _sum(cr, {k})],
                })
            salary_ctr = sorted({p["contract"] for p in mr if p["kind"] == "оклад"})
            mix = sorted({p["kind"] for p in mr if p["kind"]})
            rate = sum(c["ставка"] or 0 for c in by_ctr)
            change = []
            if prev_ctr is not None and salary_ctr != prev_ctr:
                # Появился второй договор оклада — это открытая ставка, а не
                # смена: человек остался там, где был, и взял работу еще. А вот
                # когда договор из-под оклада ушел, это перевод: приказ,
                # допник, разговор с человеком. Считаем их по отдельности.
                added = [c for c in salary_ctr if c not in prev_ctr]
                gone = [c for c in prev_ctr if c not in salary_ctr]
                if gone:
                    change.append("оклад переехал: %s → %s"
                                  % (", ".join(prev_ctr), ", ".join(salary_ctr) or "—"))
                    switch += 1
                if added and not gone:
                    change.append("добавился оклад на %s" % ", ".join(added))
                    opened += 1
            if prev_mix is not None and mix != prev_mix:
                change.append("состав: %s → %s" % (", ".join(prev_mix),
                                                   ", ".join(mix)))
                mix_change += 1
            # Набор видов может не меняться, а доли — меняться: открылась
            # половина ставки, оклад вырос, надбавка ужалась. Экономисту это
            # такая же смена структуры, и молчать о ней нельзя.
            oklad = _sum(mr, {"оклад"})
            if prev_oklad is not None and abs(oklad - prev_oklad) > EPS:
                change.append("оклад %s → %s ₽ (ставка %s → %s)"
                              % (_mo(prev_oklad), _mo(oklad),
                                 _mo(prev_rate), _mo(rate)))
            if change:
                struct_change += 1
            if rate > (e["rate"] or 1.0) + 0.01:
                extra_months += 1
            total_m = _sum(mr)
            rates.append(rate)
            if total_m:
                oklad_share.append(oklad / total_m)
            months.append({"м": m, "всего": round(total_m, 2), "ставка": round(rate, 2),
                           "договоры": by_ctr, "виды": mix, "изменилось": change})
            prev_ctr, prev_mix, prev_oklad, prev_rate = salary_ctr, mix, oklad, rate
        year = _sum(rows)
        total_switch += switch
        total_mix += mix_change
        total_extra += extra_months
        total_opened += opened
        total_struct += struct_change
        people.append({
            "код": e["code"], "фио": e["fio"], "должность": e["position"],
            "подразделение": e["department"], "ставка": e["rate"],
            "зарплата": e["salary"], "за год": round(year, 2), "месяцы": months,
            "смен договора оклада": switch, "смен состава выплат": mix_change,
            "открыто ставок": opened, "месяцев со сменой структуры": struct_change,
            "ставка макс": round(max(rates), 2) if rates else None,
            "доля оклада": ({"мин": round(min(oklad_share), 4),
                             "макс": round(max(oklad_share), 4)}
                            if oklad_share else None),
            "месяцев с совместительством": extra_months,
            "договоры": sorted({p["contract"] for p in rows}),
            "доли видов": [{"вид": k, "доля": round(_sum(rows, {k}) / year, 4)}
                           for k in ("оклад", "120", "122", "124", "152", "приказ")
                           if year and _sum(rows, {k})],
        })
    return {
        "договоры": [{"код": c, "название": (ctr.get(c) or {}).get("name", ""),
                      "ГОЗ": bool((ctr.get(c) or {}).get("goz"))} for c in order],
        "почему": _why_rates(inp, res),
        "итого": {"человек": len(people),
                  "смен договора оклада": total_switch,
                  "открыто ставок": total_opened,
                  "смен состава выплат": total_mix,
                  "месяцев со сменой структуры": total_struct,
                  "месяцев с совместительством": total_extra,
                  "за год": round(sum(p["за год"] for p in people), 2)},
        "люди": people,
    }
