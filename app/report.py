# -*- coding: utf-8 -*-
"""План ФОТ на год — данные для отчета в формах экономистов.

Отчет собран по прототипу, согласованному с экономистами: двенадцать
разделов, в каждом только числа и графы. Считается по входному файлу и
готовому плану выплат, без слов решателя — так же, как проверки в rules.py.

Разделы: итоги за год; договоры на конец года; освоение ФОТ по договорам;
касса по месяцам; выплаты по договорам и видам; регистр сотрудников (ставки,
фонды, надбавки); ставки по месяцам; БЭП; П4; трудоемкость; кто закрывает;
незакрытая трудоемкость.
"""
from __future__ import annotations

import datetime as dt

import rules
from rules import SHORT, _by, _month, _num, _rows_of, _sum, _yes

TOL = 0.5          # человеко-месяцы: меньше половины месяца — не расхождение
KIND_NAMES = {"оклад": "оклад", "120": "120 — гостайна", "122": "122 — за качество",
              "124": "124 — интенсивность", "152": "152 — дополнительная работа",
              "приказ": "стимулирующая приказом"}


def _date(v):
    """«dd.mm.yyyy» из даты, строки или порядкового номера Excel."""
    if v in (None, ""):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    if "." in s:
        return s[:10]
    n = _num(s)
    if n and 1 <= n <= 80000:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(n))).strftime("%d.%m.%Y")
    return s


def _r(v, d=2):
    return None if v is None else round(float(v), d)


def _pct(a, b):
    return round(a / b, 4) if b else None


def _extra_input(path):
    """То, что rules._read_input не хранит: даты, счета, предел выплат, категории."""
    from openpyxl import load_workbook
    from fot_planner.position_reference import normalize_position

    wb = load_workbook(path, data_only=True)
    ctr, cat, head, emp = {}, {}, {}, {}
    for row in _rows_of(wb["договоры"]):
        code = row.get("код")
        if not code:
            continue
        deadline = _month(row.get("конечная дата выплат оклада"))
        ctr[str(code)] = {
            "дата начала": _date(row.get("дата начала")),
            "дата окончания": _date(row.get("дата окончания")),
            "срок выплат": _date(row.get("конечная дата выплат оклада"))
                           or _date(row.get("дата окончания")),
            "предел": deadline,
            "счет": row.get("счет") or None,
            "тип": row.get("тип договора") or None,
            "подразделение": row.get("подразделение") or None,
            "признак": ", ".join(
                [x for x in ("ГОЗ" if _yes(row.get("ГОЗ")) else None,
                             "приоритет" if _yes(row.get("Приоритет")) else None) if x]
            ) or (row.get("тип договора") or None),
        }
    if "лимиты_по_должностям" in wb.sheetnames:
        for row in _rows_of(wb["лимиты_по_должностям"]):
            if row.get("должность"):
                cat[normalize_position(str(row["должность"]))] = row.get("категория персонала") or None
    if "трудоемкость_по_договорам" in wb.sheetnames:
        for row in _rows_of(wb["трудоемкость_по_договорам"]):
            if row.get("договор"):
                key = (str(row["договор"]), normalize_position(str(row.get("должность") or "")))
                head[key] = _num(row.get("количество человек"))
    for row in _rows_of(wb["сотрудники"]):
        if row.get("код строки"):
            emp[str(row["код строки"])] = {
                "категория занятости": str(row.get("категория занятости") or "").strip().lower(),
                "тип занятости": str(row.get("тип занятости") or "").strip().lower(),
            }
    return ctr, cat, head, emp


def report(input_path, result_path):
    inp = rules._read_input(input_path)
    res = rules._read_result(result_path)
    extra, cats, heads, emp_extra = _extra_input(input_path)
    plan, ctr, lim, norm = res["plan"], inp["contracts"], inp["limits"], inp["norm"]
    year = int(inp["year"] or 0) or None
    emps = {e["code"]: e for e in inp["employees"]}

    def base_of(position):
        return (lim.get(norm(position)) or {}).get("оклад")

    def rate_of(rows):
        b = base_of(rows[0]["position"]) if rows else None
        return (_sum(rows, {"оклад"}) / b) if b else None

    paid = {c: [0.0] * 12 for c in ctr}
    for p in plan:
        if p["contract"] in paid:
            paid[p["contract"]][p["month"] - 1] += p["amount"]
    got = {c: list(inp["inflow"].get(c) or [0.0] * 12) for c in ctr}
    total_fot = sum(c["fot"] for c in ctr.values())
    total_got = sum(sum(v) for v in got.values())
    total_paid = sum(sum(v) for v in paid.values())

    # ── 1. итоги ──────────────────────────────────────────────────────
    totals = {
        "ФОТ": _r(total_fot), "поступило": _r(total_got), "выплачено": _r(total_paid),
        "остаток ФОТ": _r(total_fot - total_paid),
        "остаток на счетах": _r(total_got - total_paid),
    }

    # ── 2. договоры на конец года ────────────────────────────────────
    contracts = []
    for code, c in sorted(ctr.items()):
        x = extra.get(code, {})
        m_from, m_to = c["from"], c["to"]
        deadline = x.get("предел") or m_to
        pay_m = paid[code]
        band = []
        for m in range(1, 13):
            active = m_from <= m <= m_to
            band.append({"м": m, "с": (2 if pay_m[m - 1] > 0.5 else 1) if active else 0,
                         "предел": m == deadline, "после": active and m > deadline})
        with_pay = [m for m in range(1, 13) if pay_m[m - 1] > 0.5]
        window = [m for m in range(m_from, m_to + 1) if m <= deadline]
        cpaid, cgot = sum(pay_m), sum(got[code])
        status = []
        if any(m > deadline for m in with_pay):
            status.append("после срока")
        if not with_pay:
            status.append("нет выплат")
        else:
            span = range(with_pay[0], with_pay[-1] + 1)
            if any(pay_m[m - 1] <= 0.5 for m in span):
                status.append("с перерывом")
            if with_pay[-1] < deadline and c["fot"] - cpaid <= 1:
                status.append("раньше срока")
        if c["fot"] - cpaid > 1:
            status.append("не освоен")
        contracts.append({
            "код": code, "название": c["name"], "ГОЗ": c["goz"],
            "признак": x.get("признак"), "счет": x.get("счет"), "тип": x.get("тип"),
            "выплаты": ", ".join(KIND_NAMES[k].split(" — ")[0] for k in
                                 ("оклад", "120", "122", "124", "152", "приказ") if c["kinds"].get(k)),
            "дата начала": x.get("дата начала"), "дата окончания": x.get("дата окончания"),
            "срок выплат": x.get("срок выплат"), "с": m_from, "по": m_to, "предел": deadline,
            "подразделение": x.get("подразделение"),
            "месяцы": band, "месяцев с выплатами": len(with_pay), "месяцев в окне": len(window),
            "ФОТ": _r(c["fot"]), "поступило": _r(cgot), "выплачено": _r(cpaid),
            "остаток": _r(c["fot"] - cpaid), "освоение": _pct(cpaid, c["fot"]),
            "статус": ", ".join(status) or "в срок",
        })

    # ── 3. освоение ФОТ по договорам ─────────────────────────────────
    usage = []
    org_bal, org_plan, org_fact = [0.0] * 13, [0.0] * 12, [0.0] * 12
    for c in contracts:
        code, fot = c["код"], c["ФОТ"] or 0.0
        n = c["по"] - c["с"] + 1
        even = fot / n if n else 0.0
        bal, bals, plans, facts = None, [], [], []
        for m in range(1, 13):
            active = c["с"] <= m <= c["по"]
            if m == c["с"]:
                bal = fot
            bals.append(_r(bal) if active else None)
            plans.append(_r(even) if active else None)
            facts.append(_r(paid[code][m - 1]) if active else None)
            if active:
                bal -= paid[code][m - 1]
                org_bal[m - 1] += bals[-1] or 0.0
                org_plan[m - 1] += even
                org_fact[m - 1] += paid[code][m - 1]
        bals.append(_r(bal))
        org_bal[12] += bal or 0.0
        usage.append({"код": code, "название": c["название"], "признак": c["признак"],
                      "счет": c["счет"], "остатки": bals, "план": plans, "факт": facts})
    usage_total = {"остатки": [_r(v) for v in org_bal], "план": [_r(v) for v in org_plan],
                   "факт": [_r(v) for v in org_fact]}

    # ── 4. касса по месяцам ──────────────────────────────────────────
    cash = []
    for c in contracts:
        code = c["код"]
        bal, rows = 0.0, {"начало": [], "поступление": [], "доступно": [], "выплаты": [],
                          "конец": [], "освоено": []}
        cum = 0.0
        for m in range(1, 13):
            active = c["с"] <= m <= c["по"]
            if not active:
                for k in rows:
                    rows[k].append(None)
                continue
            inflow, pay = got[code][m - 1], paid[code][m - 1]
            rows["начало"].append(_r(bal))
            rows["поступление"].append(_r(inflow))
            rows["доступно"].append(_r(bal + inflow))
            rows["выплаты"].append("запрет" if m > c["предел"] else _r(pay))
            bal = bal + inflow - pay
            rows["конец"].append(_r(bal))
            cum += pay
            rows["освоено"].append(_pct(cum, c["ФОТ"]))
        cash.append({"код": code, "строки": rows, "год": {
            "поступление": _r(sum(got[code])), "выплаты": _r(sum(paid[code])),
            "конец": _r(bal)}})

    # ── 5. выплаты по договорам и видам ──────────────────────────────
    kinds = []
    for code, c in sorted(ctr.items()):
        crows = [p for p in plan if p["contract"] == code]
        cpaid = _sum(crows)
        for k in ("оклад", "120", "122", "124", "152", "приказ"):
            s = _sum(crows, {k})
            allowed = c["kinds"].get(k)
            if not s and not allowed:
                continue
            months = len({p["month"] for p in crows if p["kind"] == k})
            kinds.append({"договор": code, "вид": KIND_NAMES[k], "сумма": _r(s),
                          "доля договора": _pct(s, cpaid), "доля фонда": _pct(s, total_paid),
                          "месяцев": months, "разрешен": bool(allowed)})

    # ── 6–7. регистр и ставки по месяцам ─────────────────────────────
    # Основное место решатель выбирает на каждый месяц отдельно: ставка на
    # нем равна штатной, остальные ставки того же месяца — внутреннее
    # совместительство не больше 0,5. В сентябре человек может целиком
    # перейти с одного договора на другой — и основное место переедет с ним.
    # Внешний совместитель основного места здесь не имеет, его предел 0,5.
    register, rates_tbl = [], []
    for e in inp["employees"]:
        erows = [p for p in plan if p["emp"] == e["code"]]
        if not erows:
            continue
        base = base_of(e["position"])
        limit_full = (lim.get(norm(e["position"])) or {}).get("П2556")
        cat = cats.get(norm(e["position"]))
        ex = emp_extra.get(e["code"], {})
        cat_emp = ex.get("категория занятости")
        outer_part = "совмест" in (ex.get("тип занятости") or "")
        by_c = _by(erows, "contract")
        person_rates, fonds = {}, {}
        for (code,), crows in sorted(by_c.items()):
            rate_m, fond_m, nadb_m = [None] * 12, [None] * 12, [None] * 12
            for m in range(1, 13):
                mr = [p for p in crows if p["month"] == m]
                if not mr:
                    continue
                okl = _sum(mr, {"оклад"})
                rate_m[m - 1] = _r(okl / base, 2) if base and okl else (0.0 if not okl else None)
                fond_m[m - 1] = _r(okl)
                nadb_m[m - 1] = _r(_sum(mr) - okl)
            person_rates[code] = rate_m
            fonds[code] = (fond_m, nadb_m)
        # Основное место и ставки — как их видит решатель (лист
        # «открытые_ставки»). Без листа (старый результат) — договор с
        # наибольшей окладной ставкой в месяце.
        main_m = [None] * 12
        solver_rates = [r for r in (res.get("rates") or []) if r["emp"] == e["code"]]
        if solver_rates:
            for r in solver_rates:
                if r["contract"] in person_rates:
                    person_rates[r["contract"]][r["month"] - 1] = _r(r["rate"], 2)
                    if r["main"]:
                        main_m[r["month"] - 1] = r["contract"]
        else:
            for i in range(12):
                cand = [(rm[i], code) for code, rm in person_rates.items() if rm[i]]
                if cand:
                    main_m[i] = max(cand)[1]
        for code in sorted(person_rates):
            rate_m, (fond_m, nadb_m) = person_rates[code], fonds[code]
            crows = by_c[(code,)]
            periods, cur = [], None
            for m in range(1, 13):
                if fond_m[m - 1] is None:
                    cur = None
                    continue
                is_main = main_m[m - 1] == code
                key = (fond_m[m - 1], nadb_m[m - 1], rate_m[m - 1], is_main)
                if cur and cur["ключ"] == key and cur["по"] == m - 1:
                    cur["по"], cur["месяцев"] = m, cur["месяцев"] + 1
                else:
                    cur = {"ключ": key, "с": m, "по": m, "месяцев": 1}
                    periods.append(cur)
            acc = extra.get(code, {}).get("счет")
            nadb_codes = sorted({p["kind"] for p in crows if p["kind"] and p["kind"] != "оклад"})
            months_main = sum(1 for i in range(12) if main_m[i] == code)
            months_on = sum(1 for i in range(12) if rate_m[i])
            how = ("надбавка с договора" if not months_on
                   else ("по совместительству" if outer_part or not months_main
                         else ("основное место работы" if months_main == months_on
                               else "основное / совместительство")))
            register.append({
                "табельный": e["code"], "отдел": e["department"] or None, "фио": e["fio"],
                "категория персонала": cat, "должность": e["position"], "занятость": how,
                "тип занятости": ("внешний совместитель" if outer_part else "основной"),
                "лицевой счет оклада": acc, "договор": code,
                "лицевой счет надбавки": acc if nadb_codes else None,
                "код надбавки": ", ".join(nadb_codes) or None,
                "периоды": [{"с": p["с"], "по": p["по"], "месяцев": p["месяцев"],
                             "ставка": p["ключ"][2], "фонд зп": p["ключ"][0],
                             "фонд надбавок": p["ключ"][1],
                             "занятость": ("основное" if p["ключ"][3] and not outer_part
                                           else ("совместительство" if p["ключ"][2] else None)),
                             "итого": _r((p["ключ"][0] + p["ключ"][1]) * p["месяцев"]),
                             "лимит": _r(limit_full * p["ключ"][2]) if limit_full and p["ключ"][2] else None,
                             "запас": (_r(limit_full * p["ключ"][2] - p["ключ"][0] - p["ключ"][1])
                                       if limit_full and p["ключ"][2] else None)}
                            for p in periods],
            })
        total_limit = (0.5 if outer_part or cat_emp == "студент"
                       else (0.75 if cat_emp == "аспирант" else 1.5))
        total_m = [None] * 12
        for code, rm in person_rates.items():
            for i, v in enumerate(rm):
                if v is not None:
                    total_m[i] = (total_m[i] or 0.0) + v
        contracts_r = []
        for code, rm in sorted(person_rates.items()):
            if not any(rm):
                continue    # договор только с надбавкой: ставки на нем нет
            mains = [False] * 12 if outer_part else [main_m[i] == code for i in range(12)]
            months_main = sum(mains)
            months_on = sum(1 for v in rm if v)
            part_vals = [v for i, v in enumerate(rm) if v and not mains[i]]
            contracts_r.append({
                "код": code,
                "занятость": ("основное" if months_main and months_main == months_on
                              else ("совместительство" if not months_main else "смешанная")),
                "основное": mains,
                "предел": None if not part_vals else 0.5,
                "месяцы": rm, "макс": max([v for v in rm if v is not None] or [None]),
                "макс совм": max(part_vals) if part_vals else None,
            })
        rates_tbl.append({
            "табельный": e["code"], "фио": e["fio"], "штатная": e["rate"],
            "категория занятости": cat_emp or None,
            "тип занятости": ("внешний совместитель" if outer_part else None),
            "договоры": contracts_r,
            "всего": [_r(v) for v in total_m], "предел": total_limit,
            "макс": max([v for v in total_m if v is not None] or [None]),
        })

    # ── 6.1 выплаты по месяцам: сотрудник → договор → вид выплаты ────
    # План выплат как он есть: каждая сумма в своём месяце. Регистр (6.2)
    # сворачивает месяцы в периоды, и помесячная картина в нём не видна.
    monthly, org_m = [], [0.0] * 12
    for e in inp["employees"]:
        erows = [p for p in plan if p["emp"] == e["code"]]
        if not erows:
            continue
        lines = []
        # Порядок видов — как в зарплате: оклад, потом надбавки по коду, приказ.
        order = {k: i for i, k in enumerate(("оклад", "120", "122", "124", "152", "приказ"))}
        groups = sorted(_by(erows, "contract", "kind").items(),
                        key=lambda kv: (kv[0][0], order.get(kv[0][1], 99), str(kv[0][1])))
        for (code, kind), rr in groups:
            vals = [None] * 12
            for p in rr:
                if p["amount"]:
                    vals[p["month"] - 1] = (vals[p["month"] - 1] or 0.0) + p["amount"]
            if not any(vals):
                continue
            lines.append({"договор": code, "вид": kind,
                          "месяцы": [_r(v) if v is not None else None for v in vals],
                          "год": _r(sum(v or 0.0 for v in vals))})
        tot = [None] * 12
        for p in erows:
            if p["amount"]:
                tot[p["month"] - 1] = (tot[p["month"] - 1] or 0.0) + p["amount"]
                org_m[p["month"] - 1] += p["amount"]
        monthly.append({"табельный": e["code"], "фио": e["fio"], "строки": lines,
                        "итого": [_r(v) if v is not None else None for v in tot],
                        "год": _r(sum(v or 0.0 for v in tot))})

    # ── 8. БЭП ───────────────────────────────────────────────────────
    bep = []
    for (code, m), rr in sorted(_by(plan, "contract", "month").items()):
        c = ctr.get(code)
        if not c or not c["goz"]:
            continue
        # БЭП свой у должности: предел договора — сумма «БЭП × ставка» по
        # людям, предел средней — то же, делённое на сумму ставок.
        cap, staff, rates = 0.0, 0.0, 0.0
        for (ecode,), er in _by(rr, "emp").items():
            b = (lim.get(norm(er[0]["position"])) or {}).get("БЭП")
            r = rate_of(er)
            if b and r:
                cap += b * r
            if r:
                rates += r
            staff += _sum(er, {"оклад", "122"})
        if not cap or not rates:
            continue
        avg, limit = staff / rates, cap / rates
        bep.append({"договор": code, "месяц": m, "сумма": _r(staff), "ставок": _r(rates),
                    "средняя": _r(avg), "БЭП": _r(limit), "запас": _r(limit - avg),
                    "отклонение": _pct(avg - limit, limit)})

    # ── 9. П4 ────────────────────────────────────────────────────────
    p4 = []
    for (ecode, m), rr in sorted(_by(plan, "emp", "month").items()):
        if not any(p["kind"] == "124" for p in rr):
            continue
        limit = (lim.get(norm(rr[0]["position"])) or {}).get("П4")
        rate = sum(filter(None, (rate_of(cr) for (_,), cr in _by(rr, "contract").items())))
        if not limit or not rate:
            continue
        total = _sum(rr, {"оклад", "122", "124"})
        p4.append({"табельный": ecode, "фио": rr[0]["fio"], "месяц": m,
                   "оклад": _r(_sum(rr, {"оклад"})), "122": _r(_sum(rr, {"122"})),
                   "124": _r(_sum(rr, {"124"})), "итого": _r(total), "ставка": _r(rate),
                   "предел": _r(limit * rate), "запас": _r(limit * rate - total),
                   "отклонение": _pct(total - limit * rate, limit * rate)})

    # ── 10–12. трудоемкость ──────────────────────────────────────────
    # Допуск трудоёмкости из настроек расчёта: внутри него расхождение не
    # считается отклонением — ни по чел.-мес., ни по сумме, ни по средней.
    tol = _num((inp.get("settings") or {}).get("допуск трудоёмкости"))
    tol = 0.05 if tol is None else float(tol)
    # Кто и сколько закрыл по строке — из разбивки решателя (лист «Контроль
    # трудоёмкости»): ведущий инженер может закрывать строку инженеров, по
    # совпадению должности этого не увидеть. Без разбивки (старый результат)
    # — по совпадению должности, как раньше.
    labor, who, gaps = [], [], []
    code_by_fio = {p["fio"]: p["emp"] for p in plan}
    attributed = {}
    for a in res.get("labor_people") or []:
        key = (a["contract"], norm(a["row"]) if a["row"] else None)
        attributed.setdefault(key, {}).setdefault(a["fio"], {"position": a["position"]})[a["kind"]] = a["months"]

    def people_by_position(code, want):
        crows = [p for p in plan if p["contract"] == code
                 and (want is None or norm(p["position"]) == want)]
        out = []
        for (ecode,), er in sorted(_by(crows, "emp").items()):
            rate_m, pay_m = [None] * 12, [None] * 12
            for m in range(1, 13):
                mr = [p for p in er if p["month"] == m]
                if mr:
                    rate_m[m - 1], pay_m[m - 1] = _r(rate_of(mr) or 0.0), _r(_sum(mr))
            out.append((ecode, er[0]["fio"], er[0]["position"], rate_m, pay_m))
        return out

    def people_by_solver(attr):
        out = []
        for fio, a in sorted(attr.items(), key=lambda kv: code_by_fio.get(kv[0], kv[0])):
            pm, pay = a.get("pm") or [0.0] * 12, a.get("pay") or [0.0] * 12
            if not any(pm) and not any(pay):
                continue
            rate_m = [(_r(v) if v else None) for v in pm]
            pay_m = [(_r(pay[i]) if pm[i] or pay[i] else None) for i in range(12)]
            out.append((code_by_fio.get(fio, ""), fio, a["position"], rate_m, pay_m))
        return out

    for lp in inp["labor"]:
        if not lp["person_months"]:
            continue
        code, want = lp["contract"], norm(lp["position"]) if lp["position"] else None
        people, pm_m, sum_m, heads_m = [], [0.0] * 12, [0.0] * 12, [0] * 12
        attr = attributed.get((code, want))
        found = people_by_solver(attr) if attr is not None else people_by_position(code, want)
        for ecode, fio, position, rate_m, pay_m in found:
            for i in range(12):
                pm_m[i] += rate_m[i] or 0.0
                sum_m[i] += pay_m[i] or 0.0
                if rate_m[i]:
                    heads_m[i] += 1
            people.append({"табельный": ecode, "фио": fio, "должность": position,
                           "ставка": rate_m, "начислено": pay_m,
                           "средняя": [(_r(pay_m[i] / rate_m[i]) if rate_m[i] else None)
                                       for i in range(12)],
                           "ставка год": _r(sum(v or 0 for v in rate_m)),
                           "начислено год": _r(sum(v or 0 for v in pay_m))})
        plan_pm, avg = lp["person_months"], lp["avg"] or 0.0
        plan_sum = plan_pm * avg
        fact_pm, fact_sum = sum(pm_m), sum(sum_m)
        fact_avg = fact_sum / fact_pm if fact_pm else None
        pm_ok = abs(fact_pm - plan_pm) <= max(TOL, tol * plan_pm)
        sum_ok = not plan_sum or abs(fact_sum - plan_sum) <= tol * plan_sum
        avg_ok = not avg or fact_avg is None or abs(fact_avg - avg) <= tol * avg
        status = ("сходится" if pm_ok and sum_ok and avg_ok
                  else ("не закрыто" if fact_pm < plan_pm - max(TOL, tol * plan_pm)
                        else ("отклонение по сумме" if not sum_ok else "отклонение по средней")))
        headcount = heads.get((code, want or ""))
        labor.append({"договор": code, "строка": lp["position"] or "любая должность",
                      "план чел-мес": _r(plan_pm), "факт чел-мес": _r(fact_pm),
                      "д чел-мес": _r(fact_pm - plan_pm), "план сумма": _r(plan_sum),
                      "факт сумма": _r(fact_sum), "д сумма": _r(fact_sum - plan_sum),
                      "средняя план": _r(avg) if avg else None, "средняя факт": _r(fact_avg),
                      "д средней": _r(fact_avg - avg) if avg and fact_avg is not None else None,
                      "людей предел": headcount, "людей макс": max(heads_m) if heads_m else 0,
                      "допуск": tol, "вне допуска": {"чел-мес": not pm_ok, "сумма": not sum_ok,
                                                      "средняя": not avg_ok},
                      "статус": status})
        who.append({"договор": code, "строка": lp["position"] or "любая должность",
                    "люди": people, "итого ставка": [_r(v) for v in pm_m],
                    "итого начислено": [_r(v) for v in sum_m],
                    "итого ставка год": _r(fact_pm), "итого начислено год": _r(fact_sum)})
        if fact_pm < plan_pm - TOL:
            months_active = [m for m in range(ctr[code]["from"], ctr[code]["to"] + 1)] if code in ctr else []
            gap = plan_pm - fact_pm
            n = len(months_active) or 12
            gaps.append({"договор": code, "должность": lp["position"] or "любая",
                         "нужно": _r(plan_pm), "закрыто": _r(fact_pm), "не закрыто": _r(gap),
                         "месяцев": n, "ставок в месяц": _r(gap / n)})

    # Настройки расчёта, от которых зависит результат: экономист должен видеть
    # их рядом с планом, а не искать в файле входа.
    st = inp.get("settings") or {}
    settings = [{"имя": k, "значение": st.get(k)} for k in
                ("допуск трудоёмкости", "макс договоров оклада в год", "разрешить дефицит")
                if st.get(k) is not None]
    deficits = [{"табельный": d["emp"], "фио": d["fio"], "месяц": d["month"],
                 "положено": _r(d["due"]), "выплачено": _r(d["paid"]), "дефицит": _r(d["gap"]),
                 "причина": d["why"]} for d in (res.get("deficits") or []) if d["gap"] > 0.5]
    return {
        "год": year, "итоги": totals, "договоры": contracts,
        "настройки": settings, "дефицит": deficits,
        "освоение": {"договоры": usage, "итого": usage_total},
        "касса": cash, "виды": kinds, "регистр": register, "ставки": rates_tbl,
        "помесячно": {"сотрудники": monthly, "итого": [_r(v) for v in org_m],
                      "год": _r(sum(org_m))},
        "бэп": bep, "п4": p4, "трудоемкость": labor, "кто": who, "незакрыто": gaps,
        "месяцы": SHORT,
    }
