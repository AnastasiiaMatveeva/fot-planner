# -*- coding: utf-8 -*-
"""Выгрузка плана в Excel в тех же разделах, что на вкладке «План ФОТ».

Решатель пишет свой файл результата (план выплат, касса, контроль
трудоёмкости, цели, открытые ставки). Экономист же смотрит отчёт из
пятнадцати разделов, который сервис считает по входу и результату
(report.py, rules.py). Скачивание отдавало файл решателя, и в нём не было
половины таблиц с экрана. Здесь к файлу решателя дописываются листы по
разделам отчёта — один к одному с интерфейсом: те же графы, те же числа.

Технические листы решателя остаются в конце книги: они основание расчёта.
"""
from __future__ import annotations

import json
import os
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import report as report_mod
import rules as rules_mod

SHORT = report_mod.SHORT
_HEAD_FILL = PatternFill(fill_type="solid", fgColor="E8EEF6")
_THIN = Side(style="thin", color="BFC7D1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOLD = Font(name="Arial", size=10, bold=True)
_FONT = Font(name="Arial", size=10)
_MONEY = "#,##0.00"
_RATE = "0.00"
_PCT = "0.0%"


def _yn(v):
    return "да" if v else "нет"


def _pct(v):
    return None if v is None else float(v)


def _sheet(wb, title, headers, rows, *, formats=None, widths=None, freeze=1):
    """Лист-таблица: заголовок, строки, автофильтр, закреплённые графы.

    ``formats`` — {индекс графы: формат числа}; ``widths`` — ширины граф.
    Значение ячейки — как есть: числа числами, чтобы в Excel считались.
    """
    title = title[:31]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.border = _BOLD, _HEAD_FILL, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    formats = formats or {}
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = _FONT, _BORDER
            fmt = formats.get(c - 1)
            if fmt and isinstance(v, (int, float)):
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")
    for c in range(1, len(headers) + 1):
        w = (widths or {}).get(c - 1)
        if w is None:
            longest = max([len(str(headers[c - 1]))] + [len(str(r[c - 1])) for r in rows[:200]
                                                         if c - 1 < len(r) and r[c - 1] is not None])
            w = min(max(9, longest + 2), 48)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = ws.cell(2, freeze + 1)
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)
    ws.sheet_view.showGridLines = False
    return ws


def _months_head(prefix=""):
    return [prefix + m for m in SHORT]


def build(input_path, result_path, out_path, goals=None, versions=None):
    """Собрать книгу: листы разделов отчёта + листы решателя из result_path."""
    rep = report_mod.report(input_path, result_path)
    checks = rules_mod.check(input_path, result_path)
    shutil.copyfile(result_path, out_path)
    wb = load_workbook(out_path)
    solver_sheets = list(wb.sheetnames)
    money = {}

    # 1. Итоги
    t = rep["итоги"]
    rows = [[k, v] for k, v in t.items()]
    _sheet(wb, "1 Итоги", ["Показатель", "Сумма, ₽"], rows, formats={1: _MONEY}, widths={0: 28, 1: 18})

    # 2. Договоры
    heads = ["Шифр", "Название", "Признак", "Счет", "Тип", "Подразделение", "Виды выплат",
             "Дата начала", "Дата окончания", "Срок выплат", "ФОТ, ₽", "Поступило, ₽", "Выплачено, ₽",
             "Остаток, ₽", "Освоение", "Месяцев с выплатами", "Месяцев в окне", "Статус"]
    rows = [[c["код"], c["название"], c["признак"], c["счет"], c["тип"], c["подразделение"],
             c["выплаты"], c["дата начала"], c["дата окончания"], c["срок выплат"], c["ФОТ"],
             c["поступило"], c["выплачено"], c["остаток"], _pct(c["освоение"]),
             c["месяцев с выплатами"], c["месяцев в окне"], c["статус"]] for c in rep["договоры"]]
    _sheet(wb, "2 Договоры", heads, rows,
           formats={10: _MONEY, 11: _MONEY, 12: _MONEY, 13: _MONEY, 14: _PCT})

    # 3. Освоение ФОТ по договорам: остатки, план, факт по месяцам
    heads = ["Шифр", "Название", "Показатель"] + _months_head() + ["На конец года"]
    rows = []
    for u in rep["освоение"]["договоры"]:
        rows.append([u["код"], u["название"], "остаток ФОТ на начало месяца"] + u["остатки"][:12] + [u["остатки"][12]])
        rows.append([u["код"], u["название"], "план равномерного освоения"] + u["план"] + [None])
        rows.append([u["код"], u["название"], "выплачено"] + u["факт"] + [None])
    tot = rep["освоение"]["итого"]
    rows.append(["Итого", "", "остаток ФОТ на начало месяца"] + tot["остатки"][:12] + [tot["остатки"][12]])
    rows.append(["Итого", "", "план равномерного освоения"] + tot["план"] + [None])
    rows.append(["Итого", "", "выплачено"] + tot["факт"] + [None])
    _sheet(wb, "3 Освоение", heads, rows, formats={i: _MONEY for i in range(3, 16)}, freeze=3)

    # 4. Касса
    heads = ["Шифр", "Показатель"] + _months_head() + ["За год"]
    rows = []
    names = {"начало": "остаток на начало", "поступление": "поступление", "доступно": "доступно",
             "выплаты": "выплаты", "конец": "остаток на конец", "освоено": "освоено нарастающим, доля"}
    for c in rep["касса"]:
        for key, label in names.items():
            year = c["год"].get(key) if key in ("поступление", "выплаты", "конец") else None
            vals = [v if v != "запрет" else "запрет" for v in c["строки"][key]]
            rows.append([c["код"], label] + vals + [year])
    _sheet(wb, "4 Касса", heads, rows, formats={i: _MONEY for i in range(2, 15)}, freeze=2)

    # 5. Виды выплат
    heads = ["Договор", "Вид выплаты", "Сумма за год, ₽", "Доля договора", "Доля фонда", "Месяцев", "Разрешен договором"]
    rows = [[k["договор"], k["вид"], k["сумма"], _pct(k["доля договора"]), _pct(k["доля фонда"]),
             k["месяцев"], _yn(k["разрешен"])] for k in rep["виды"]]
    _sheet(wb, "5 Виды выплат", heads, rows, formats={2: _MONEY, 3: _PCT, 4: _PCT})

    # 6.1 Выплаты по месяцам
    heads = ["Табельный", "ФИО", "Договор", "Вид выплаты"] + _months_head() + ["За год, ₽"]
    rows = []
    pm = rep["помесячно"]
    for p in pm["сотрудники"]:
        for l in p["строки"]:
            rows.append([p["табельный"], p["фио"], l["договор"], l["вид"]] + l["месяцы"] + [l["год"]])
        rows.append([p["табельный"], p["фио"], "Итого по сотруднику", ""] + p["итого"] + [p["год"]])
    rows.append(["Итого по организации", "", "", ""] + pm["итого"] + [pm["год"]])
    _sheet(wb, "6.1 Выплаты по месяцам", heads, rows, formats={i: _MONEY for i in range(4, 17)}, freeze=4)

    # 6.2 Регистр по периодам
    heads = ["Табельный", "ФИО", "Подразделение", "Категория персонала", "Должность", "Занятость",
             "Тип занятости", "Договор", "Лицевой счет оклада", "Лицевой счет надбавки", "Код надбавки",
             "С (месяц)", "По (месяц)", "Месяцев", "Ставка", "Занятость в периоде", "Фонд оклада в месяц, ₽",
             "Фонд надбавок в месяц, ₽", "Итого за период, ₽", "Лимит П2556 на ставку, ₽", "Запас, ₽"]
    rows = []
    for s in rep["регистр"]:
        for p in s["периоды"]:
            rows.append([s["табельный"], s["фио"], s["отдел"], s["категория персонала"], s["должность"],
                         s["занятость"], s["тип занятости"], s["договор"], s["лицевой счет оклада"],
                         s["лицевой счет надбавки"], s["код надбавки"], SHORT[p["с"] - 1], SHORT[p["по"] - 1],
                         p["месяцев"], p["ставка"], p["занятость"], p["фонд зп"], p["фонд надбавок"],
                         p["итого"], p["лимит"], p["запас"]])
    _sheet(wb, "6.2 Регистр", heads, rows,
           formats={14: _RATE, 16: _MONEY, 17: _MONEY, 18: _MONEY, 19: _MONEY, 20: _MONEY}, freeze=2)

    # 6.3 ШР на дату, детализация назначений (форма кадров)
    heads = ["Таб.№", "Назначение", "Фамилия И.О., уч. ст., уч. зван.", "Код подр.", "Подразделение",
             "Должность", "Категория персонала", "Код пар-ра", "Название параметра", "Ставка",
             "Номинальное значение параметра", "Значение параметра по ставке", "Сумма в руб.",
             "Начало действия", "Окончание действия", "Код шифра затрат", "Шифр затрат", "Лицевой счет",
             "УИ ПНИЭР", "Номер приказа ввода", "Дата приказа ввода", "Номер приказа закрытия",
             "Дата приказа закрытия"]
    rows = [[r["таб"], r["назначение"], r["фио"], r["код_подр"], r["подразделение"], r["должность"],
             r["категория"], r["код"], r["параметр"], r["ставка"], r["номинал"], r["по_ставке"], r["сумма"],
             r["начало"], r["окончание"], r["код_шифра"], r["шифр"], r["счет"], None, None, None, None, None]
            for r in rep.get("шр") or []]
    _sheet(wb, "6.3 ШР детализация", heads, rows,
           formats={9: _RATE, 10: _MONEY, 11: _MONEY, 12: _MONEY}, freeze=3)

    # 7.1 Сотрудники, 7.2 Ставки по месяцам
    pinfo = {}
    for s in rep["регистр"]:
        pinfo.setdefault(s["табельный"], s)
    heads = ["Табельный", "ФИО", "Должность", "Подразделение", "Категория персонала", "Тип занятости",
             "Категория занятости", "Штатная ставка", "Предел ставки"]
    rows = []
    for p in rep["ставки"]:
        s = pinfo.get(p["табельный"], {})
        rows.append([p["табельный"], p["фио"], s.get("должность"), s.get("отдел"), s.get("категория персонала"),
                     p["тип занятости"] or s.get("тип занятости"), p["категория занятости"], p["штатная"], p["предел"]])
    _sheet(wb, "7.1 Сотрудники", heads, rows, formats={7: _RATE, 8: _RATE})
    heads = ["Табельный", "ФИО", "Договор", "Занятость"] + _months_head() + ["Макс в плане"]
    rows = []
    for p in rep["ставки"]:
        for c in p["договоры"]:
            for label, is_main in (("основное", True), ("совместительство", False)):
                vals = [v if (v and bool(c["основное"][i]) == is_main) else None for i, v in enumerate(c["месяцы"])]
                if not any(vals):
                    continue
                rows.append([p["табельный"], p["фио"], c["код"], label] + vals + [max(v for v in vals if v)])
        rows.append([p["табельный"], p["фио"], "Всего", ""] + p["всего"] + [p["макс"]])
    _sheet(wb, "7.2 Ставки по месяцам", heads, rows, formats={i: _RATE for i in range(4, 17)}, freeze=4)

    # 8. БЭП
    heads = ["Договор", "Месяц", "Сумма оклада и 122, ₽", "Сумма ставок", "Средняя на ставку, ₽", "БЭП, ₽",
             "Запас, ₽", "Отклонение от БЭП"]
    rows = [[b["договор"], SHORT[b["месяц"] - 1], b["сумма"], b["ставок"], b["средняя"], b["БЭП"], b["запас"],
             _pct(b["отклонение"])] for b in rep["бэп"]]
    _sheet(wb, "8 БЭП", heads, rows, formats={2: _MONEY, 3: _RATE, 4: _MONEY, 5: _MONEY, 6: _MONEY, 7: _PCT})

    # 9. П4
    heads = ["Табельный", "ФИО", "Месяц", "Оклад, ₽", "122, ₽", "124, ₽", "Итого по П4, ₽", "Суммарная ставка",
             "Предел П4 на ставку, ₽", "Запас, ₽", "Отклонение от П4"]
    rows = [[p["табельный"], p["фио"], SHORT[p["месяц"] - 1], p["оклад"], p["122"], p["124"], p["итого"],
             p["ставка"], p["предел"], p["запас"], _pct(p["отклонение"])] for p in rep["п4"]]
    _sheet(wb, "9 П4", heads, rows,
           formats={3: _MONEY, 4: _MONEY, 5: _MONEY, 6: _MONEY, 7: _RATE, 8: _MONEY, 9: _MONEY, 10: _PCT})

    # 10. Трудоёмкость
    heads = ["Договор", "Строка РКМ", "План, чел.-мес.", "Факт, чел.-мес.", "Δ чел.-мес.", "План, ₽", "Факт, ₽",
             "Δ ₽", "Средняя план, ₽", "Средняя факт, ₽", "Δ средней", "Людей предел", "Людей макс в месяце",
             "Допуск", "Статус"] + _months_head("план ")
    rows = [[l["договор"], l["строка"], l["план чел-мес"], l["факт чел-мес"], l["д чел-мес"], l["план сумма"],
             l["факт сумма"], l["д сумма"], l["средняя план"], l["средняя факт"], l["д средней"],
             l["людей предел"], l["людей макс"], l["допуск"], l["статус"]]
            + (l.get("план по месяцам") or [None] * 12) for l in rep["трудоемкость"]]
    _sheet(wb, "10 Трудоёмкость", heads, rows,
           formats={2: _RATE, 3: _RATE, 4: _RATE, 5: _MONEY, 6: _MONEY, 7: _MONEY, 8: _MONEY, 9: _MONEY,
                    10: _MONEY, 13: _PCT, **{i: _RATE for i in range(15, 27)}}, freeze=2)

    # 11. Исполнители
    heads = ["Договор", "Строка РКМ", "Табельный", "ФИО", "Должность", "Показатель"] + _months_head() + ["За год"]
    rows = []
    for w in rep["кто"]:
        for p in w["люди"]:
            rows.append([w["договор"], w["строка"], p["табельный"], p["фио"], p["должность"], "закрыто ставкой"]
                        + p["ставка"] + [p["ставка год"]])
            rows.append([w["договор"], w["строка"], p["табельный"], p["фио"], p["должность"], "начислено, ₽"]
                        + p["начислено"] + [p["начислено год"]])
        rows.append([w["договор"], w["строка"], "", "Итого по строке", "", "ставка"]
                    + w["итого ставка"] + [w["итого ставка год"]])
        rows.append([w["договор"], w["строка"], "", "Итого по строке", "", "начислено, ₽"]
                    + w["итого начислено"] + [w["итого начислено год"]])
    _sheet(wb, "11 Исполнители", heads, rows, formats={i: "#,##0.00" for i in range(6, 19)}, freeze=2)

    # 12. Нехватка
    heads = ["Договор", "Должность", "Нужно, чел.-мес.", "Закрыто", "Не закрыто", "Месяцев", "Ставок в месяц"]
    rows = [[g["договор"], g["должность"], g["нужно"], g["закрыто"], g["не закрыто"], g["месяцев"],
             g["ставок в месяц"]] for g in rep["незакрыто"]]
    _sheet(wb, "12 Нехватка", heads, rows, formats={2: _RATE, 3: _RATE, 4: _RATE, 6: _RATE})

    # 13. Расчёт: версии, цели, настройки
    rows = []
    for v in versions or []:
        rows.append(["версия", v.get("номер"), v.get("статус"), v.get("секунд"), v.get("дата"), None, None])
    for g in goals or []:
        rows.append(["цель", g.get("цель"), g.get("приоритет"), g.get("значение"), g.get("единица"),
                     g.get("вес"), g.get("вклад")])
    for st in rep["настройки"]:
        rows.append(["настройка", st["имя"], st["значение"], None, None, None, None])
    try:
        import reference
        for field, x in reference.read_sources().items():
            rows.append(["источник справочника", field, x.get("документ"), x.get("основание"),
                         x.get("действует с") or x.get("загружен"), None, None])
    except Exception:  # noqa: BLE001 — источники не обязательны для выгрузки
        pass
    _sheet(wb, "13 Расчёт", ["Что", "Имя", "Значение / приоритет", "Значение", "Единица / дата", "Вес", "Вклад"],
           rows, widths={0: 12, 1: 48})

    # 14. Дефицит
    heads = ["Табельный", "ФИО", "Месяц", "Положено, ₽", "Выплачено, ₽", "Дефицит, ₽", "Причина"]
    rows = [[d["табельный"], d["фио"], SHORT[d["месяц"] - 1], d["положено"], d["выплачено"], d["дефицит"],
             d["причина"]] for d in rep["дефицит"]]
    _sheet(wb, "14 Дефицит", heads, rows, formats={3: _MONEY, 4: _MONEY, 5: _MONEY})

    # 15. Контроль расчёта
    heads = ["Раздел", "Правило", "Тип", "Состояние", "Факт", "Проверено", "Нарушений", "Использовано предела",
             "Где смотреть", "Смысл"]
    rows = [[c["раздел"], c["правило"], c["тип"], c["состояние"], c["факт"], c["проверено"], c["нарушений"],
             _pct(c.get("использовано")), c["где"], c["смысл"]] for c in checks]
    ws = _sheet(wb, "15 Контроль расчёта", heads, rows, formats={7: _PCT}, widths={1: 44, 4: 60, 9: 70})
    for r in range(2, ws.max_row + 1):
        for c in (5, 10):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # Листы решателя — в конец, как основание расчёта.
    for name in solver_sheets:
        wb.move_sheet(name, offset=len(wb.sheetnames))
    wb.active = 0
    wb.save(out_path)
    return out_path
