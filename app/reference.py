# -*- coding: utf-8 -*-
"""Справочник должностей: чтение, сверка, запись.

Хранится не в базе, а листом «лимиты_по_должностям» входного файла — того
самого, который читает решатель. Так правка попадает в расчет без
промежуточных копий и без риска, что интерфейс и модель разойдутся.

Должности сопоставляются справочником синонимов самого сервиса: в документах
пишут «Вед. инженер», «Инженер I категории», «МНС», и второй справочник для
этого заводить незачем.
"""
from __future__ import annotations

import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(ROOT, "src"))

TEMPLATE = os.path.abspath(os.environ.get("FOT_REFERENCE_TEMPLATE") or
                           os.path.join(ROOT, "docs", "ui", "demo_input.xlsx"))
SHEET = "лимиты_по_должностям"
COLS = {"оклад": 6, "П2556": 7, "П4": 8, "БЭП": 9}
#: Справка П4 даёт среднюю зарплату по категории с отпускными; предел П4 —
#: без них: средняя делится на 1,086866 (доп. зарплата 8,6866 %). Правило
#: экономистов («справочная информация_ЗП.xlsx», 09.09.2026):
#: 162 648,30 / 1,086866 = 149 648,90.
P4_VACATION_COEF = 1.086866
FIELD_KEY = {"оклад": "sal", "П2556": "p2556", "П4": "p4", "БЭП": "bep"}

#: Категория персонала по её названию в документе: справка П4 даёт среднюю
#: зарплату «научным работникам» и «научно-техническому персоналу», а не
#: должностям; в справочнике категория стоит у каждой должности.
_CATEGORY_WORDS = (("научно-техн", "НТП"), ("нтп", "НТП"),
                   ("научн", "НР"), ("нр", "НР"),
                   ("административ", "АУП"), ("ауп", "АУП"),
                   ("профессорско", "ППС"), ("ппс", "ППС"),
                   ("прочий", "ПП"), ("прочие", "ПП"))


def category_of(text):
    """«Научно-технический персонал» → «НТП», «Научные работники» → «НР»."""
    t = " ".join(str(text or "").lower().replace("ё", "е").split())
    if not t:
        return None
    for word, code in _CATEGORY_WORDS:
        if word in t:
            return code
    return None


def _is_everyone(text):
    """«Основные исполнители ГОЗ», «все работники» — величина на всех."""
    t = str(text or "").lower()
    return any(w in t for w in ("исполнител", "все ", "всех", "работники по гоз", "гоз"))


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    s = "".join(c for c in s if c.isdigit() or c in ".-")
    try:
        return float(s) if s else None
    except ValueError:
        return None


def read_rows():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    out = []
    for r in range(2, ws.max_row + 1):
        pos = ws.cell(r, 1).value
        if not pos:
            continue
        out.append({"pos": str(pos), "cat": ws.cell(r, 2).value or "",
                    "page": ws.cell(r, 3).value or "",
                    "group": ws.cell(r, 4).value,
                    "level": ws.cell(r, 5).value,
                    "sal": _num(ws.cell(r, 6).value),
                    "p2556": _num(ws.cell(r, 7).value),
                    "p4": _num(ws.cell(r, 8).value),
                    "bep": _num(ws.cell(r, 9).value),
                    "note": ws.cell(r, 10).value or ""})
    return out


_INDEX = {"for": None, "map": {}}


def _index(rows):
    names = tuple(r["pos"] for r in rows)
    if _INDEX["for"] == names:
        return _INDEX["map"]
    from fot_planner.position_reference import (
        default_position_synonyms, normalize_position,
    )
    idx = {normalize_position(r["pos"]): r for r in rows}
    for raw, canonical in default_position_synonyms().items():
        target = idx.get(normalize_position(canonical))
        if target is not None:
            idx.setdefault(normalize_position(raw), target)
    _INDEX["for"], _INDEX["map"] = names, idx
    return idx


#: Хвосты перечислений: должностей за ними нет.
_TAILS = ("и пр", "пр", "и др", "др", "и т д", "и т п", "прочие", "прочее")
#: Окончания творительного падежа: «заведующий лабораторией, сектором» —
#: это дополнения к должности, а не отдельные должности.
_INSTR = ("ой", "ей", "ом", "ем", "ью", "ами", "ями")
#: Должности, которые управляют дополнением: за ними в перечислении идут
#: «центра, службы, отдела» — это одна должность, а не три.
_GOVERNS = ("директор", "начальник", "заведующий", "заведующая", "руководитель",
            "заместитель", "зам")


def _split_outside(text):
    """Разбить перечисление по «,», «;» и переводам строк, не трогая скобки."""
    out, buf, depth = [], [], 0
    for ch in text:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        if ch in ",;\n" and depth == 0:
            out.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    out.append("".join(buf))
    return [p.strip(" .;,") for p in out if p.strip(" .;,")]


def _singular(text):
    """Заголовок группы — в единственное число: «Главные специалисты».

    В приказе группы названы во множественном числе, в справочнике должности
    — в единственном. Правило грубое, по окончаниям, и работает как запасной
    ключ поиска: не сошлось — ищем по исходному написанию.
    """
    out = []
    for w in text.split():
        low = w.lower()
        if len(low) > 4 and low.endswith(("ые", "ие")):
            out.append(w[:-2] + ("ый" if low.endswith("ые") else "ий"))
        elif len(low) > 5 and low.endswith("ли"):
            out.append(w[:-2] + "ль")
        elif len(low) > 4 and low.endswith(("ы", "и")) and not low.endswith("ии"):
            out.append(w[:-1])
        else:
            out.append(w)
    return " ".join(out)


def _expand_group(text):
    """Строку приказа — в должности справочника.

    Разбираем три вида групп:

    * «Ведущие специалисты ... наименование "ведущий": программист, инженер»
      — производное наименование ставится приставкой к каждой должности;
    * «Специалисты, по которым может устанавливаться II или I
      внутридолжностная категория: инженер» — в справочнике это отдельные
      строки «инженер 1 категории» и «инженер 2 категории»;
    * «Директор, начальник центра, службы, отдела» — одинокое слово в конце
      перечисления относится к предыдущему главному слову.

    Двоеточие раскрываем только там, где перед ним условие, а не должность:
    в «Главные специалисты: в отделах, лабораториях» после двоеточия места
    работы, и должность стоит слева.
    """
    head, sep, tail = text.partition(":")
    hl = head.lower()
    rule = any(w in hl for w in ("по котор", "наименование", "категори"))
    ruled = bool(sep and rule and tail.strip())
    body = tail if ruled else head
    prefix = "Ведущий" if (ruled and "производное" in hl and "ведущ" in hl) else ""
    cats = ("1 категории", "2 категории") if (ruled and "внутридолжностн" in hl) else ()

    parts, last_head = [], ""
    for item in _split_outside(body):
        if item.lower().strip(" .") in _TAILS:
            continue
        words = item.split()
        if len(words) == 1 and last_head.lower() in _GOVERNS:
            item = last_head + " " + item
        elif len(words) > 1:
            last_head = words[0]
        parts.append(item)

    out = []
    for item in parts:
        base = [item]
        # «Заведующие (лабораторией, сектором)» — скобки раскрываем только
        # при творительном падеже: в «Начальник (директор, заведующий)»
        # перечислены синонимы одной должности, их склеивать нельзя.
        m = re.match(r"^([^(]+)\(([^)]*)\)\s*$", item)
        if m:
            inner = [x for x in _split_outside(m.group(2))
                     if x.lower().strip(" .") not in _TAILS]
            inner = [x for x in inner if x.split()[0].lower().endswith(_INSTR)]
            if inner:
                base = [m.group(1).strip() + " " + x for x in inner]
        for name in base:
            if prefix:
                out.append(prefix + " " + name.lower())
            elif cats:
                out.extend(name + " " + c for c in cats)
            else:
                out.append(name)
    return out


def _names(raw):
    """Одна строка документа — в список должностей.

    В положении об оплате труда оклад задан не должности, а квалификационной
    группе: «Аналитик; архитектор; аудитор; бухгалтер» — одна сумма на всех.
    В приказе № 2556 группа названа условием: «Ведущие специалисты, по
    которым может устанавливаться ... "ведущий": программист, инженер».
    Разворачиваем и то и другое, иначе ни одна должность не сопоставится.
    """
    text = " ".join(str(raw or "").split())
    if not text:
        return []
    names = _expand_group(text) or [text]
    out = []
    for n in names:
        for v in (n, _singular(n)):
            v = v.strip(" .;,")
            if v and v not in out:
                out.append(v)
    return out


def _key(value):
    """Число величины — в вид, по которому его находят в тексте документа."""
    f = float(value)
    return "%d" % f if f == int(f) else "%s" % round(f, 4)


def compare(rows):
    """Величины из документа против действующего справочника.

    Возвращает расхождения, названия должностей, которых нет в справочнике,
    графы, которые документ заполняет, и то, что документ дал справочнику:
    величины и строки, в которых они стояли. По ним карточка документа
    подсвечивает в тексте то, что ушло в расчёт, — величина без строки
    подсветила бы и соседнюю графу с тем же числом.
    """
    from fot_planner.position_reference import normalize_position
    current = read_rows()
    idx = _index(current)
    changes, unknown, seen, values, lines, miss = [], [], set(), {}, {}, {}
    for row in rows or []:
        field = row.get("field")
        value = row.get("value")
        key = FIELD_KEY.get(field)
        if not key or value is None:
            continue
        # Число из документа и число в справочнике — разные: у П4 в справке
        # стоит средняя зарплата, а пределом становится она же без отпускных.
        # Подсвечивать в тексте надо то, что там написано, иначе маска ищет
        # величину, которой в документе нет.
        raw = value
        if field == "П4":
            value = round(float(value) / P4_VACATION_COEF, 2)
        names = _names(row.get("pos"))
        hits = []
        for pos in names:
            hit = idx.get(normalize_position(pos))
            if hit is not None and hit not in hits:
                hits.append(hit)
        # Величина не должности, а категории персонала или всех исполнителей:
        # раскладываем по должностям справочника с этой категорией.
        if not hits:
            cat = category_of(row.get("pos"))
            if cat:
                hits = [r for r in current if str(r.get("cat") or "").strip().upper() == cat]
            elif field == "БЭП" and _is_everyone(row.get("pos")):
                hits = list(current)
        matched = False
        for hit in hits:
            matched = True
            seen.add(field)
            values[_key(value)] = field
            values[_key(raw)] = field
            src = " ".join(str(row.get("pos") or "").split())
            if src:
                lines.setdefault(src, set()).add(_key(value))
                lines[src].add(_key(raw))
            old = hit[key]
            # Копейки, потерянные распознаванием скана («280 024» вместо
            # «280 023,62»), — не расхождение: величина та же с точностью до рубля.
            differs = old is None or abs(float(old) - float(value)) > 1.0
            if differs and not any(
                    c["pos"] == hit["pos"] and c["field"] == field for c in changes):
                changes.append({"pos": hit["pos"], "field": field,
                                "old": old, "new": float(value)})
        if not matched:
            label = names[0] if len(names) == 1 else str(row.get("pos"))[:80]
            if label not in unknown:
                unknown.append(label)
            src = " ".join(str(row.get("pos") or "").split())
            miss.setdefault(src, set()).add(_key(value))
            miss[src].add(_key(raw))
    given = {"значения": values,
             "строки": [{"строка": k, "значения": sorted(v)}
                        for k, v in lines.items()],
             "без должности": [{"строка": k, "значения": sorted(v)}
                               for k, v in miss.items()]}
    return changes, unknown[:20], sorted(seen), given


def apply(edits):
    """Записать величины в лист. edits: [{pos, field, value}]."""
    if not edits:
        return []
    from fot_planner.position_reference import normalize_position
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    at = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v:
            at[normalize_position(v)] = r
    applied = []
    for e in edits:
        r = at.get(normalize_position(e.get("pos", "")))
        col = COLS.get(e.get("field"))
        if not r or not col:
            continue
        new = _num(e.get("value") if "value" in e else e.get("new"))
        old = _num(ws.cell(r, col).value)
        if old == new:
            continue
        ws.cell(r, col).value = new
        applied.append({"pos": ws.cell(r, 1).value, "field": e["field"],
                        "old": old, "new": new})
    if applied:
        wb.save(TEMPLATE)
    return applied


#: Лист источников: какая величина справочника из какого документа.
#: Величины П2556, П4 и БЭП приходят приказом, справкой и письмом; связь
#: с документом — часть основания расчёта, без неё цифра «ничья».
SOURCES_SHEET = "источники_справочника"
SOURCE_FIELDS = ("оклад", "П2556", "П4", "БЭП")
_SOURCE_COLS = ("величина", "документ", "id документа", "загружен", "основание",
                "действует с", "записано")


def _sources_sheet(wb):
    if SOURCES_SHEET in wb.sheetnames:
        return wb[SOURCES_SHEET]
    ws = wb.create_sheet(SOURCES_SHEET)
    for c, name in enumerate(_SOURCE_COLS, start=1):
        ws.cell(1, c).value = name
    return ws


def read_sources():
    """{величина: {документ, id, загружен, основание, действует с, записано}}."""
    wb = openpyxl.load_workbook(TEMPLATE)
    if SOURCES_SHEET not in wb.sheetnames:
        return {}
    ws = wb[SOURCES_SHEET]
    out = {}
    for r in range(2, ws.max_row + 1):
        field = ws.cell(r, 1).value
        if not field:
            continue
        out[str(field)] = {"документ": ws.cell(r, 2).value,
                           "document_id": ws.cell(r, 3).value,
                           "загружен": ws.cell(r, 4).value,
                           "основание": ws.cell(r, 5).value,
                           "действует с": ws.cell(r, 6).value,
                           "записано": ws.cell(r, 7).value}
    return out


def set_source(field, doc_name, doc_id, uploaded=None, basis=None, effective=None,
               when=None):
    """Записать (или заменить) документ-источник величины справочника."""
    import datetime as _dt
    field = str(field or "").strip()
    if field not in SOURCE_FIELDS:
        return False
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = _sources_sheet(wb)
    row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "") == field:
            row = r
            break
    if row is None:
        row = ws.max_row + 1
    values = (field, doc_name, doc_id, uploaded, basis, effective,
              when or _dt.datetime.now().strftime("%d.%m.%Y %H:%M"))
    for c, v in enumerate(values, start=1):
        ws.cell(row, c).value = v
    wb.save(TEMPLATE)
    return True


def forget_source(doc_id):
    """Снять ссылку на удалённый документ-источник.

    Величина остаётся в справочнике: по ней уже посчитаны планы, и стирать её
    вместе с документом нельзя. Но основание пропадает, и об этом надо
    сказать вслух — иначе в реестре стоит ссылка на документ, которого нет.
    Возвращает список величин, оставшихся без основания.
    """
    if not doc_id:
        return []
    wb = openpyxl.load_workbook(TEMPLATE)
    if SOURCES_SHEET not in wb.sheetnames:
        return []
    ws = wb[SOURCES_SHEET]
    lost = []
    for r in range(2, ws.max_row + 1):
        field = ws.cell(r, 1).value
        if not field or str(ws.cell(r, 3).value or "") != str(doc_id):
            continue
        lost.append(str(field))
        for c in range(2, len(_SOURCE_COLS) + 1):
            ws.cell(r, c).value = None
        ws.cell(r, 5).value = "документ удалён"
    if lost:
        wb.save(TEMPLATE)
    return lost


def rows_by_anchors(path):
    """Запасной разбор документа: таблица с ожидаемыми заголовками."""
    wb = openpyxl.load_workbook(path, data_only=True)
    want = {"оклад": "оклад", "п2556": "П2556", "п4": "П4"}
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 12) + 1):
            hdr = {}
            for c in range(1, ws.max_column + 1):
                t = str(ws.cell(r, c).value or "").strip().lower()
                if t.startswith("должност"):
                    hdr["pos"] = c
                for k, field in want.items():
                    if t.replace(" ", "").startswith(k):
                        hdr[field] = c
            if "pos" in hdr and len(hdr) > 1:
                rows = []
                for rr in range(r + 1, ws.max_row + 1):
                    pos = ws.cell(rr, hdr["pos"]).value
                    if not pos:
                        continue
                    for field, col in hdr.items():
                        if field == "pos":
                            continue
                        v = _num(ws.cell(rr, col).value)
                        if v is not None:
                            rows.append({"pos": str(pos), "field": field, "value": v})
                return rows
    return []


def upsert_position(pos, cat=None, sal=None, p2556=None, p4=None, bep=None,
                    note=None):
    """Завести должность в справочнике или дописать ей недостающие величины.

    ``apply`` умеет только править числа у должности, которая в листе уже
    есть: строки, которой нет, он молча пропускает. Для документа, где названа
    новая должность, этого мало — ее надо завести.

    Страницу, номер группы и уровень не выдумываем: решатель без них обходится,
    такая должность попадает в группу «без окладной группы». Соврать здесь
    хуже, чем оставить пусто, — по этим полям строятся группы
    взаимозаменяемости.

    Возвращает ("заведена"|"дополнена"|"без изменений", что именно изменилось).
    """
    from fot_planner.position_reference import normalize_position

    name = str(pos or "").strip()
    if not name:
        return "без изменений", []

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    at = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v:
            at[normalize_position(v)] = r

    values = {2: cat, 6: sal, 7: p2556, 8: p4, 9: bep, 10: note}
    row = at.get(normalize_position(name))
    what = "дополнена"
    if row is None:
        row = ws.max_row + 1
        ws.cell(row, 1).value = name
        what = "заведена"

    changed = []
    titles = {2: "категория", 6: "оклад", 7: "П2556", 8: "П4", 9: "БЭП",
              10: "примечание"}
    for col, value in values.items():
        if value in (None, ""):
            continue
        old = ws.cell(row, col).value
        new = _num(value) if col in (6, 7, 8, 9) else str(value).strip()
        if old == new or (old not in (None, "") and _num(old) == new
                          and col in (6, 7, 8, 9)):
            continue
        ws.cell(row, col).value = new
        changed.append((titles[col], old, new))

    if what == "заведена" or changed:
        wb.save(TEMPLATE)
        return what, changed
    return "без изменений", []
